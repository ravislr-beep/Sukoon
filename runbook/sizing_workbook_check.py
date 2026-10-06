"""Recalculate the sizing workbooks in LibreOffice and check them.

1. Every formula cell must evaluate without an error value.
2. The worked example must match an independent Python model of the same published rules.
"""
import math
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

import sizing_wb_ref as R
import sizing_workbook as W

ERRORS = ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "Err:")


def recalc(path):
    tmp = Path(tempfile.mkdtemp())
    ods = tmp / "x.ods"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "ods", "--outdir", str(tmp), str(path)],
                   check=True, capture_output=True)
    (tmp / (path.stem + ".ods")).rename(ods)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "o"), str(ods)],
                   check=True, capture_output=True)
    out = tmp / "o" / "x.xlsx"
    return out


def scan_errors(formulas_wb, values_wb):
    bad = []
    for ws in formulas_wb.worksheets:
        vs = values_wb[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    v = vs[c.coordinate].value
                    if isinstance(v, str) and v.startswith(ERRORS):
                        bad.append((ws.title, c.coordinate, v, c.value[:90]))
    return bad


# ---------------------------------------------------------------- independent model of the worked example

CF = {t: dict(zip(["in", "out", "part", "compact", "conn", "req", "maxu", "msgmax", "partops"], R.CF_TABLE[t]))
      for t in R.CF_TYPES}


def kafka_model():
    envs = {code: dict(zip(["mparts", "qps", "topics", "ppt", "compact", "msgs", "bytes", "min", "fan", "mout", "day",
                            "mgib", "ret", "comp", "conn", "req", "msgmax"], v)) for code, v in W.KD_EXAMPLE.items()}
    clusters = {}
    for code, _d, land, k, _s in W.EXAMPLE_ENVS:
        e = envs[code]
        parts = e["mparts"] if e["mparts"] is not None else (e["qps"] + e["topics"]) * e["ppt"]
        mbps_in = e["min"] if e["min"] is not None else e["msgs"] * e["bytes"] / 1e6
        mbps_out = e["mout"] if e["mout"] is not None else mbps_in * e["fan"]
        gib = e["mgib"] if e["mgib"] is not None else e["day"] * e["bytes"] / 2 ** 30
        c = clusters.setdefault(k, dict(part=0, compact=0, inn=0, out=0, conn=0, req=0, gib=0, gibr=0, store=0,
                                        msg=0, rank=0, n=0))
        c["part"] += parts
        c["compact"] += e["compact"] or 0
        c["inn"] += mbps_in
        c["out"] += mbps_out
        c["conn"] += e["conn"]
        c["req"] += e["req"]
        c["gib"] += gib
        c["gibr"] += gib * e["fan"]
        c["store"] += gib * e["ret"] / 24 / e["comp"]
        c["msg"] = max(c["msg"], e["msgmax"])
        c["rank"] = max(c["rank"], R.LANDSCAPES.index(land) + 1)
        c["n"] += 1
    res = {}
    for i, k in enumerate(W.KIDS):
        if k not in clusters:
            continue
        c = clusters[k]
        sla = float(W.KC_EXAMPLE["sla"][i].rstrip("%")) / 100
        priv, byok, mz = W.KC_EXAMPLE["priv"][i], W.KC_EXAMPLE["byok"][i], W.KC_EXAMPLE["mz"][i]
        growth, util = W.KC_EXAMPLE["growth"][i], W.KC_EXAMPLE["util"][i]
        f = (1 + growth) / util
        p = dict(part=math.ceil(c["part"] * f - 1e-9), compact=math.ceil(c["compact"] * f - 1e-9), inn=c["inn"] * f,
                 out=c["out"] * f, conn=math.ceil(c["conn"] * f - 1e-9), req=math.ceil(c["req"] * f - 1e-9))
        elig = {"Basic": sla <= 0.995 and priv != "Yes" and byok != "Yes",
                "Standard": sla <= 0.9999 and priv != "Yes" and byok != "Yes",
                "Enterprise": sla <= 0.9999, "Dedicated": sla <= 0.9999}
        slamin = {"Basic": 1, "Standard": 2 if sla > 0.999 else 1, "Enterprise": 2 if sla > 0.999 else 1,
                  "Dedicated": 2 if (sla > 0.9995 or mz == "Yes") else 1}
        maxu = {"Basic": 50, "Standard": 10, "Enterprise": 32, "Dedicated": 24}
        units = {}
        for t in R.CF_TYPES:
            L = CF[t]
            units[t] = max(math.ceil(p["inn"] / L["in"]), math.ceil(p["out"] / L["out"]),
                           math.ceil(p["part"] / L["part"]), math.ceil(p["compact"] / L["compact"]),
                           math.ceil(p["conn"] / L["conn"]), math.ceil(p["req"] / L["req"]), slamin[t])
        typ = next((t for t in R.CF_TYPES[:3] if elig[t] and units[t] <= maxu[t]), "Dedicated")
        post = c["store"] * (1 + growth) * 3
        brokers = 3
        prod = c["rank"] >= 3
        rpb = math.ceil(p["part"] * 3 / brokers)
        cpu_p = 2 if rpb <= 1000 else (8 if rpb <= 2000 else 16)
        ram_p = 2 if rpb <= 300 else (8 if rpb <= 1000 else (32 if rpb <= 2000 else 64))
        res[k] = dict(type=typ, units=units[typ], p_part=p["part"], b_in=c["gib"] * (1 + growth) * 730 / 24,
                      b_out=c["gibr"] * (1 + growth) * 730 / 24, b_post=post,
                      k_cpu=max(4 if prod else 2, cpu_p), k_ram=max(16 if prod else 8, ram_p),
                      k_disk=max(200 if prod else 100, math.ceil(post / brokers / 0.7)))
    return res


def search_model():
    ps = {key: dict(n=n, cpu=cpu, ram=ram, disk=disk) for key, n, cpu, ram, disk, _t in R.PS_ROWS}
    svcs = {}
    for code, _d, land, _k, s in W.EXAMPLE_ENVS:
        gb, docs, idx, pri, grow, yrs, tmp = W.SD_EXAMPLE[code]
        v = svcs.setdefault(s, dict(fgb=0, tmp=0, docs=0, prim=0, rank=0))
        v["fgb"] += gb * (1 + grow) ** yrs
        v["tmp"] += tmp
        v["docs"] += docs * (1 + grow) ** yrs
        v["prim"] += idx * pri
        v["rank"] = max(v["rank"], R.LANDSCAPES.index(land) + 1)
    res = {}
    for i, s in enumerate(W.SIDS):
        if s not in svcs:
            continue
        v = svcs[s]
        cfg = {k: W.SS_EXAMPLE[k][i] for k in W.SS_EXAMPLE}
        grp = "Prod" if v["rank"] >= 3 else "Test"
        prof = "Large" if v["fgb"] >= 2000 else "Default"
        data = ps[f"{prof}|Data|{grp}"]
        cpu = cfg["cpu"] or data["cpu"]
        ram = cfg["ram"] or data["ram"]
        disk = cfg["disk"] or data["disk"]
        heap = ram / 2
        src = v["fgb"] + v["tmp"]
        rep = cfg["rep"]
        need = src * (1 + rep) * (1 + cfg["idxovh"]) / (1 - cfg["linux"]) / (1 - cfg["ovh"])
        shards = v["prim"] * (1 + rep)
        n = max(max(data["n"], 3) if cfg["ha"] == "Yes" else data["n"], math.ceil(need / (disk * 0.85)),
                math.ceil(shards / 1000), math.ceil(shards / (25 * heap)), rep + 1)
        if cfg["ha"] == "Yes" and cfg["zones"] > 1:
            n = math.ceil(n / cfg["zones"]) * cfg["zones"]
        mn = 3 if (grp == "Prod" or cfg["ha"] == "Yes") else ps[f"{prof}|Master|{grp}"]["n"]
        res[s] = dict(prof=prof, dn=n, cpu=cpu, ram=ram, disk=disk, need=need, mn=mn,
                      mram=ps[f"{prof}|Master|Prod"]["ram"],
                      use=src * (1 + rep) * (1 + cfg["idxovh"]) / (n * disk * (1 - cfg["linux"]) * (1 - cfg["ovh"])))
    return res


def compare(values_wb, ks, sz):
    ws_k, ws_s = values_wb[W.S_KS], values_wb[W.S_SZ]
    diffs, n = [], 0
    for k, exp in kafka_model().items():
        col = W.get_column_letter(W.COL0 + W.KIDS.index(k))
        for key, want in exp.items():
            got = ws_k[f"{col}{ks.rows[key]}"].value
            n += 1
            ok = got == want if isinstance(want, str) else abs((got or 0) - want) <= max(1e-6, abs(want) * 1e-9)
            if not ok:
                diffs.append(("Kafka", k, key, want, got))
    for s, exp in search_model().items():
        col = W.get_column_letter(W.COL0 + W.SIDS.index(s))
        for key, want in exp.items():
            got = ws_s[f"{col}{sz.rows[key]}"].value
            n += 1
            ok = got == want if isinstance(want, str) else abs((got or 0) - want) <= max(1e-6, abs(want) * 1e-9)
            if not ok:
                diffs.append(("Search", s, key, want, got))
    return n, diffs


def main():
    ok = True
    for example, name in [(False, "Pega_26_Kafka_OpenSearch_Sizing_Template.xlsx"),
                          (True, "Pega_26_Kafka_OpenSearch_Sizing_Worked_Example.xlsx")]:
        path = W.OUT / name
        _wb, ks, sz, _kc, _ss = W.build(example)
        calc = recalc(path)
        f_wb = load_workbook(path)
        v_wb = load_workbook(calc, data_only=True)
        bad = scan_errors(f_wb, v_wb)
        nform = sum(1 for ws in f_wb.worksheets for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("="))
        print(f"{name}: {nform} formulas, {len(bad)} errors")
        for b in bad[:40]:
            print("   ", b)
        ok &= not bad
        if example:
            n, diffs = compare(v_wb, ks, sz)
            print(f"  model comparison: {n} values, {len(diffs)} differences")
            for d in diffs:
                print("   ", d)
            ok &= not diffs
        shutil.copy(calc, W.OUT.parent / ("/tmp/recalc_" + name))
    return ok


if __name__ == "__main__" and "--scenarios" not in sys.argv:
    sys.exit(0 if main() else 1)


# ------------------------------------------------------------------------- scenario tests (hand-worked)

def scenarios():
    """Change inputs in the worked example, recalculate, and compare with hand-worked results."""
    from openpyxl.utils import get_column_letter as L
    _wb, ks, sz, kc, ss = W.build(True)
    kd = {k: L(i + 1) for i, (k, *_r) in enumerate(W.KD_COLS)}
    sd = {k: L(i + 1) for i, (k, *_r) in enumerate(W.SD_COLS)}
    col = lambda ident: L(W.COL0 + (W.KIDS + W.SIDS).index(ident) % 6)
    cases = [
        ("K3 public networking -> Standard sized by partitions/requests/SLA",
         [(W.S_KC, f"{col('K3')}{kc.rows['priv']}", "No")],
         [(W.S_KS, "K3", "type", "Standard"), (W.S_KS, "K3", "units", 7)]),
        ("K1 public networking but Standard too small -> Enterprise",
         [(W.S_KC, f"{col('K1')}{kc.rows['priv']}", "No")],
         [(W.S_KS, "K1", "type", "Enterprise"), (W.S_KS, "K1", "units", 2)]),
        ("K2 forced Standard with private networking -> FAIL",
         [(W.S_KC, f"{col('K2')}{kc.rows['pref']}", "Standard")],
         [(W.S_KS, "K2", "type", "Standard"), (W.S_KS, "K2", "c_pref", "FAIL*")]),
        ("K2 50,000 measured partitions in PERF -> Enterprise 30 eCKU, fast-scaling note",
         [(W.S_KD, f"{kd['mparts']}{W.DEM_ROW0 + 3}", 50000)],
         [(W.S_KS, "K2", "type", "Enterprise"), (W.S_KS, "K2", "units", 30), (W.S_KS, "K2", "c_fast", "INFO*")]),
        ("K2 70,000 measured partitions in PERF -> Dedicated above purchase limit",
         [(W.S_KD, f"{kd['mparts']}{W.DEM_ROW0 + 3}", 70000)],
         [(W.S_KS, "K2", "type", "Dedicated"), (W.S_KS, "K2", "units", 28), (W.S_KS, "K2", "c_fit", "WARN*")]),
        ("K1 self-managed keys -> Enterprise with key note",
         [(W.S_KC, f"{col('K1')}{kc.rows['byok']}", "Yes")],
         [(W.S_KS, "K1", "type", "Enterprise"), (W.S_KS, "K1", "c_byok", "INFO*")]),
        ("New sandbox on K4 at 99.5 %, public -> Basic 3 eCKU; on S4 without HA -> 1 data node",
         [(W.S_SETUP, f"B{W.ENV_ROW0 + 6}", "SBX"), (W.S_SETUP, f"D{W.ENV_ROW0 + 6}", "Development"),
          (W.S_SETUP, f"E{W.ENV_ROW0 + 6}", "Yes"), (W.S_SETUP, f"F{W.ENV_ROW0 + 6}", "K4"),
          (W.S_SETUP, f"G{W.ENV_ROW0 + 6}", "S4"),
          (W.S_KC, f"{col('K4')}{kc.rows['sla']}", "99.5%"), (W.S_KC, f"{col('K4')}{kc.rows['priv']}", "No"),
          (W.S_KD, f"{kd['status']}{W.DEM_ROW0 + 6}", "Estimated"), (W.S_KD, f"{kd['qps']}{W.DEM_ROW0 + 6}", 5),
          (W.S_KD, f"{kd['topics']}{W.DEM_ROW0 + 6}", 1), (W.S_KD, f"{kd['msgs']}{W.DEM_ROW0 + 6}", 100),
          (W.S_KD, f"{kd['bytes']}{W.DEM_ROW0 + 6}", 1000), (W.S_KD, f"{kd['day']}{W.DEM_ROW0 + 6}", 1e6),
          (W.S_KD, f"{kd['conn']}{W.DEM_ROW0 + 6}", 10), (W.S_KD, f"{kd['req']}{W.DEM_ROW0 + 6}", 20),
          (W.S_KD, f"{kd['msgmax']}{W.DEM_ROW0 + 6}", 100000),
          (W.S_SS, f"{col('S4')}{ss.rows['ha']}", "No"), (W.S_SS, f"{col('S4')}{ss.rows['rep']}", 0),
          (W.S_SS, f"{col('S4')}{ss.rows['zones']}", 1),
          (W.S_SD, f"{sd['status']}{W.DEM_ROW0 + 6}", "Estimated"), (W.S_SD, f"{sd['gb']}{W.DEM_ROW0 + 6}", 2),
          (W.S_SD, f"{sd['docs']}{W.DEM_ROW0 + 6}", 10000), (W.S_SD, f"{sd['idx']}{W.DEM_ROW0 + 6}", 60),
          (W.S_SD, f"{sd['grow']}{W.DEM_ROW0 + 6}", 0), (W.S_SD, f"{sd['yrs']}{W.DEM_ROW0 + 6}", 1)],
         [(W.S_KS, "K4", "type", "Basic"), (W.S_KS, "K4", "units", 3), (W.S_SZ, "S4", "dn", 1),
          (W.S_SZ, "S4", "mn", 0), (W.S_SZ, "S4", "srs", 1), (W.S_SZ, "S4", "c_rep", "WARN*")]),
        ("PROD search 1,500 GB -> Large profile, storage sets 36 nodes, alternative 9 x 981 GB",
         [(W.S_SD, f"{sd['gb']}{W.DEM_ROW0 + 5}", 1500)],
         [(W.S_SZ, "S3", "prof", "Large"), (W.S_SZ, "S3", "dn", 36), (W.S_SZ, "S3", "driver", "Storage"),
          (W.S_SZ, "S3", "alt_n", 9), (W.S_SZ, "S3", "alt", 981), (W.S_SZ, "S3", "c_t4", "INFO*")]),
    ]
    src = W.OUT / "Pega_26_Kafka_OpenSearch_Sizing_Worked_Example.xlsx"
    fails = 0
    for name, edits, expects in cases:
        wb = load_workbook(src)
        for sheet, cell, val in edits:
            wb[sheet][cell] = val
        tmp = Path(tempfile.mkdtemp()) / "scenario.xlsx"
        wb.save(tmp)
        v = load_workbook(recalc(tmp), data_only=True)
        res = []
        for sheet, ident, key, want in expects:
            rows = ks.rows if sheet == W.S_KS else sz.rows
            got = v[sheet][f"{col(ident)}{rows[key]}"].value
            ok = (str(got).startswith(want[:-1]) if isinstance(want, str) and want.endswith("*")
                  else (got == want if isinstance(want, str) else abs((got or 0) - want) < 1e-9))
            res.append(ok)
            if not ok:
                print(f"    {ident}.{key}: expected {want}, got {got}")
        bad = scan_errors(load_workbook(tmp), v)
        status = "PASS" if all(res) and not bad else "FAIL"
        fails += status == "FAIL"
        print(f"  scenario {status}: {name}" + (f" ({len(bad)} formula errors)" if bad else ""))
    return fails == 0


if __name__ == "__main__" and "--scenarios" in sys.argv:
    sys.exit(0 if scenarios() else 1)
