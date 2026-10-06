"""Kafka (Confluent Cloud) and OpenSearch sizing workbook for Pega Platform '26 on Azure AKS.

Writes two files to out/:
  Pega_26_Kafka_OpenSearch_Sizing_Template.xlsx       blank inputs for the customer
  Pega_26_Kafka_OpenSearch_Sizing_Worked_Example.xlsx the same workbook with illustrative inputs

All results are Excel formulas, so the customer and the providers can follow and change them.
"""
import math
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import sizing_wb_ref as R

OUT = Path(__file__).parent / "out"
VERSION = "1.0"
DATE = "6 October 2026"

NAVY, WHITE = "1F3864", "FFFFFF"
FILL = {k: PatternFill("solid", fgColor=v) for k, v in dict(
    input="FFF2CC", calc="F2F2F2", key="C6EFCE", ref="DDEBF7", warn="FFEB9C", fail="FFC7CE",
    head=NAVY, section="D9E1F2", white="FFFFFF", note="FBFBFB").items()}
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_BASE = Font(name="Calibri", size=10)
F_BOLD = Font(name="Calibri", size=10, bold=True)
F_HEAD = Font(name="Calibri", size=10, bold=True, color=WHITE)
F_TITLE = Font(name="Calibri", size=16, bold=True, color=NAVY)
F_SUB = Font(name="Calibri", size=11, italic=True, color="404040")
F_LINK = Font(name="Calibri", size=10, color="0563C1", underline="single")
F_KEY = Font(name="Calibri", size=10, bold=True, color="006100")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

S_SETUP, S_KC, S_KD, S_SS, S_SD = "1 Setup", "2 Kafka Clusters", "3 Kafka Demand", "4 Search Services", "5 Search Demand"
S_KS, S_SZ, S_CR, S_OR, S_CK = "6 Kafka Sizing", "7 Search Sizing", "8 Confluent Request", "9 Search Request", "10 Checks"
S_REF, S_TAB, S_REFS, S_LOG, S_COVER, S_GUIDE = "Ref_Data", "Ref_Tables", "References", "Change Log", "Cover", "Guide"

KIDS = [f"K{i}" for i in range(1, 7)]
SIDS = [f"S{i}" for i in range(1, 7)]
N_ENV = 8
ENV_ROW0 = 14          # first environment row on the Setup sheet
DEM_ROW0 = 7           # first environment row on the demand sheets
ID_ROW = 4             # row holding K1..K6 / S1..S6 on the vertical sheets
COL0 = 6               # column F holds the first cluster or service

COMMENT_AUTHOR = "Sizing workbook"


def q(sheet):
    return f"'{sheet}'"


def note(cell, text, w=320, h=None):
    c = Comment(text, COMMENT_AUTHOR)
    c.width = w
    c.height = h or max(80, 16 * (len(text) // 45 + 2))
    cell.comment = c


def style(cell, kind="calc", fmt=None, bold=False, align=None):
    cell.font = F_KEY if kind == "key" else (F_BOLD if bold else F_BASE)
    if kind in FILL:
        cell.fill = FILL[kind]
    cell.border = BOX
    cell.alignment = align or WRAP
    if fmt:
        cell.number_format = fmt
    if kind == "input":
        cell.protection = Protection(locked=False)


def header_row(ws, row, values, height=30):
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, v)
        c.fill, c.font, c.border, c.alignment = FILL["head"], F_HEAD, BOX, CENTER
    ws.row_dimensions[row].height = height


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws["A2"] = sub
    ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def fit_height(ws, row, texts_and_widths, base=13.5):
    lines = 1
    for text, width in texts_and_widths:
        if text and isinstance(text, str) and not text.startswith("="):
            lines = max(lines, sum(math.ceil(max(1, len(p)) / max(1.0, width * 1.15)) for p in text.split("\n")))
    ws.row_dimensions[row].height = max(base, base * lines + 2)


def protect(ws):
    p = ws.protection
    p.sheet = True
    p.formatColumns = False
    p.formatRows = False
    p.formatCells = False
    p.autoFilter = False
    p.sort = False


def define(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def status_format(ws, rng):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({rng.split(":")[0].replace("$", "")},4)="FAIL"'],
                                                   fill=FILL["fail"], font=Font(color="9C0006", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({rng.split(":")[0].replace("$", "")},4)="WARN"'],
                                                   fill=FILL["warn"], font=Font(color="9C5700", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({rng.split(":")[0].replace("$", "")},4)="PASS"'],
                                                   fill=FILL["key"], font=Font(color="006100", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({rng.split(":")[0].replace("$", "")},4)="INFO"'],
                                                   fill=FILL["ref"], font=Font(color="1F3864")))


class Ctx(dict):
    """format_map context: {c} column, {id} the ID cell, {key} a row on this sheet, {X.key} a row on another sheet."""

    def __init__(self, col, rows, ext):
        super().__init__()
        self.col, self.rows, self.ext = col, rows, ext

    def __missing__(self, key):
        if key == "c":
            return self.col
        if key == "id":
            return f"{self.col}${ID_ROW}"
        if key in self.rows:
            return f"{self.col}{self.rows[key]}"
        if key in self.ext:
            sheet, row = self.ext[key]
            return f"{q(sheet)}!{self.col}{row}"
        raise KeyError(key)


class Grid:
    """A sheet laid out as rows of items with one column per cluster or service (F to K)."""

    def __init__(self, ws, ids, ext=None, first_row=5, prefix="", how_label="How it is worked out"):
        self.ws, self.ids, self.ext = ws, ids, ext or {}
        self.cols = [get_column_letter(COL0 + i) for i in range(len(ids))]
        self.rows, self.r, self.n, self.prefix = {}, first_row, 0, prefix
        header_row(ws, ID_ROW, ["Step", "Item", "Unit", how_label, "Ref"] + ids)
        note(ws.cell(ID_ROW, 5), "Source IDs. P = Pegasystems, C = Confluent, O = OpenSearch Project, A = AWS OpenSearch "
                                 "Service guidance, W = workbook rule. Full links are on the References sheet.")
        widths(ws, [7, 34, 11, 58, 9] + [19] * len(ids))
        ws.freeze_panes = ws.cell(ID_ROW + 1, COL0)

    def section(self, text):
        ws = self.ws
        for col in range(1, COL0 + len(self.ids)):
            c = ws.cell(self.r, col)
            c.fill, c.border = FILL["section"], BOX
        ws.cell(self.r, 1, text).font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        ws.row_dimensions[self.r].height = 16
        self.r += 1

    def add(self, key, label, unit, how, ref, formula, kind="calc", fmt=None, comment=None, values=None):
        ws, r = self.ws, self.r
        self.n += 1
        self.rows[key] = r
        ws.cell(r, 1, f"{self.prefix}{self.n:02d}")
        ws.cell(r, 2, label)
        ws.cell(r, 3, unit)
        ws.cell(r, 4, how)
        ws.cell(r, 5, ref)
        for col in range(1, 6):
            style(ws.cell(r, col), "white", bold=(col == 2))
        if comment:
            note(ws.cell(r, 2), comment)
        for i, c in enumerate(self.cols):
            cell = ws.cell(r, COL0 + i)
            if values is not None:
                v = values[i]
            elif callable(formula):
                v = formula(c)
            elif isinstance(formula, str) and formula.startswith("="):
                v = formula.format_map(Ctx(c, self.rows, self.ext))
            else:
                v = formula
            cell.value = v
            style(cell, kind, fmt, align=Alignment(horizontal="center", vertical="top", wrap_text=True))
        fit_height(ws, r, [(label, 34), (how, 58)])
        if isinstance(formula, str) and formula.startswith("=") and kind in ("white", "key") and fmt is None:
            longest = max((len(t) for t in re.findall(r'"([^"]*)"', formula)), default=0)
            if longest > 16:
                lines = math.ceil((longest + 16) / 16.0)
                ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 13.5 * lines + 2)
        self.r += 1
        return r

    def grey_unused(self, envs_key="envs", last_row=None):
        """Grey out the columns of clusters or services that have no environments; this rule runs first."""
        ws, c0, c1 = self.ws, self.cols[0], self.cols[-1]
        last = last_row or self.r
        rule = FormulaRule(formula=[f"{c0}${self.rows[envs_key]}=0"], font=Font(color="A6A6A6"),
                           fill=PatternFill("solid", fgColor="F7F7F7", bgColor="F7F7F7"), stopIfTrue=True)
        for rules in ws.conditional_formatting._cf_rules.values():
            for x in rules:
                x.priority += 1
        rule.priority = 1
        ws.conditional_formatting.add(f"{c0}{ID_ROW + 1}:{c1}{last}", rule)

    def ref(self, key):
        return self.rows[key]


# ----------------------------------------------------------------------------------------------- inputs

EXAMPLE_ENVS = [
    # code, description, landscape, Kafka cluster, search service
    ("DEV", "Development", "Development", "K1", "S1"),
    ("SIT", "System integration test", "Testing", "K1", "S1"),
    ("UAT", "User acceptance test", "Testing", "K1", "S1"),
    ("PERF", "Performance test", "Stage", "K2", "S2"),
    ("PREPROD", "Pre-production", "Stage", "K2", "S2"),
    ("PROD", "Production", "Production", "K3", "S3"),
]

KC_DEFAULT = {
    "name": ["cc-np1", "cc-np2", "cc-prd", "", "", ""],
    "region": ["", "", "", "", "", ""],
    "sla": ["99.9%", "99.99%", "99.99%", "99.99%", "99.99%", "99.99%"],
    "priv": ["Yes"] * 6, "byok": ["No"] * 6, "mz": ["Yes"] * 6, "pref": ["Auto"] * 6,
    "bill": ["Invoice"] * 6, "growth": [0.0] * 6, "util": [0.70] * 6,
}
SS_DEFAULT = {
    "name": ["os-np1", "os-np2", "os-prd", "", "", ""],
    "region": [""] * 6,
    "ha": ["Yes"] * 6, "zones": [3] * 6, "rep": [1] * 6, "prof": ["Auto"] * 6,
    "ovh": [0.20] * 6, "linux": [0.05] * 6, "idxovh": [0.10] * 6,
    "cpu": [None] * 6, "ram": [None] * 6, "disk": [None] * 6,
    "ver": ["OpenSearch 2.15"] * 6, "priv": ["Yes"] * 6, "cmk": ["No"] * 6,
}

# Worked example: illustrative figures only, marked "Example" in every row.
KD_EXAMPLE = {
    # env: measured parts, QPs, other topics, parts/topic, compacted, msg/s, bytes, MBps, fan-out, read MBps,
    #      msgs/day, GiB/day, retention h, compression, connections, requests/s, largest message
    "DEV": (None, 120, 30, 6, 0, 1000, 1000, None, 2, None, 20e6, None, 168, 1, 600, 1000, 1000000),
    "SIT": (None, 120, 30, 6, 0, 1500, 1000, None, 2, None, 30e6, None, 168, 1, 600, 1200, 1000000),
    "UAT": (None, 120, 30, 6, 0, 2000, 1000, None, 2, None, 40e6, None, 168, 1, 800, 1500, 1000000),
    "PERF": (None, 120, 30, 6, 0, 8000, 1000, None, 2, None, 150e6, None, 72, 1, 1500, 4000, 2000000),
    "PREPROD": (None, 120, 30, 6, 0, 4000, 1000, None, 2, None, 80e6, None, 168, 1, 1500, 3000, 2000000),
    "PROD": (None, 120, 30, 6, 0, 8000, 1000, None, 2, None, 300e6, None, 168, 1, 1800, 4000, 2000000),
}
SD_EXAMPLE = {
    # env: searchable GB, documents, indexes, primaries per index, growth per year, years, temporary GB
    "DEV": (5, 100000, 60, 1, 0.20, 3, 0),
    "SIT": (10, 200000, 60, 1, 0.20, 3, 0),
    "UAT": (15, 300000, 60, 1, 0.20, 3, 0),
    "PERF": (40, 800000, 60, 1, 0.20, 3, 0),
    "PREPROD": (40, 800000, 60, 1, 0.20, 3, 0),
    "PROD": (40, 2000000, 60, 1, 0.20, 3, 0),
}
KC_EXAMPLE = dict(KC_DEFAULT, growth=[0.20] * 6)
SS_EXAMPLE = dict(SS_DEFAULT, cpu=[4, None, None, None, None, None], ram=[16, None, None, None, None, None])


def build_setup(wb, example):
    ws = wb.create_sheet(S_SETUP)
    title(ws, "1  Setup: programme details and environments",
          "Enter the programme details and the environments. Assign each environment to a Kafka cluster (K1 to K6) "
          "and a search service (S1 to S6).")
    widths(ws, [6, 26, 30, 16, 11, 14, 14, 44])
    fields = [
        ("Customer", "", "Customer or organisation name."),
        ("Programme", "Pega 8.8 to Pega Platform '26 (26.1.1) clone and upgrade on Azure AKS", "Programme name."),
        ("Prepared by", "", "Name and role of the person who completed the inputs."),
        ("Date of inputs", "", "Date the inputs were collected. Measured figures age quickly; record the date."),
        ("Azure region", "", "Azure region for AKS, Confluent Cloud and the search service. Keep them in the same region."),
        ("Target Pega version", "Pega Platform '26 (26.1.1)", "Fixed for this workbook. The references are for Pega '26."),
        ("Source system", "Pega 8.8 with embedded Kafka (Stream service) and embedded search",
         "Where the measured inputs come from."),
        ("Workbook status", "Draft", "Draft, For provider quote or Final."),
    ]
    for i, (label, value, tip) in enumerate(fields):
        r = 4 + i
        ws.cell(r, 2, label)
        style(ws.cell(r, 2), "white", bold=True)
        c = ws.cell(r, 3, value)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        style(c, "input")
        for col in range(4, 7):
            style(ws.cell(r, col), "input")
        note(ws.cell(r, 2), tip)
    dv = DataValidation(type="list", formula1='"Draft,For provider quote,Final"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C11")

    header_row(ws, ENV_ROW0 - 1, ["#", "Environment code", "Description", "Pega landscape", "In scope",
                                  "Kafka cluster", "Search service", "Notes"], height=32)
    tips = {
        2: "Short code used everywhere else in the workbook, for example PROD.",
        4: "Pega's sizing tables use Production, Stage, Testing and Development [P1][P2]. Production and Stage use the "
           "production rows of the Pega tables; Testing and Development use the development rows. Use Stage for a "
           "production-like test environment such as PERF or PREPROD.",
        5: "Only environments marked Yes are counted.",
        6: "Kafka cluster ID (K1 to K6). Several Pega environments can share a cluster; set a different "
           "stream.streamNamePattern prefix per environment [P7].",
        7: "Search service ID (S1 to S6). SRS can support more than one Pega deployment and isolates data with a unique "
           "CUSTOMER_DEPLOYMENT_ID. Do not share the service with non-Pega software [P2].",
    }
    for col, tip in tips.items():
        note(ws.cell(ENV_ROW0 - 1, col), tip)
    for i in range(N_ENV):
        r = ENV_ROW0 + i
        ws.cell(r, 1, f"E{i + 1}")
        style(ws.cell(r, 1), "calc", align=CENTER)
        vals = list(EXAMPLE_ENVS[i]) if i < len(EXAMPLE_ENVS) else ["", "", "", "", ""]
        row_vals = [vals[0], vals[1], vals[2], "Yes" if vals[0] else "", vals[3], vals[4], ""]
        for j, v in enumerate(row_vals):
            c = ws.cell(r, 2 + j, v if v != "" else None)
            style(c, "input", align=CENTER if j in (2, 3, 4, 5) else WRAP)
    rng = lambda col: f"{get_column_letter(col)}{ENV_ROW0}:{get_column_letter(col)}{ENV_ROW0 + N_ENV - 1}"
    for col, f1 in [(4, "=LANDSCAPES"), (5, '"Yes,No"'), (6, "=KAFKA_IDS"), (7, "=SEARCH_IDS")]:
        dv = DataValidation(type="list", formula1=f1, allow_blank=True)
        dv.error, dv.errorTitle = "Choose a value from the list.", "Invalid entry"
        ws.add_data_validation(dv)
        dv.add(rng(col))
    r = ENV_ROW0 + N_ENV + 1
    ws.cell(r, 2, "Environment assignment used in the delivered template follows runbook v2.3 Section 5 (option B): "
                  "DEV, SIT and UAT share K1 and S1; PERF and PREPROD share K2 and S2; PROD has K3 and S3. "
                  "Change it if the design changes.").font = F_SUB
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 30
    ws.cell(r, 2).alignment = WRAP
    for name, col in [("ENV_CODE", 2), ("ENV_LANDSCAPE", 4), ("ENV_SCOPE", 5), ("ENV_KAFKA", 6), ("ENV_SEARCH", 7)]:
        define(wb, name, f"{q(S_SETUP)}!${get_column_letter(col)}${ENV_ROW0}:${get_column_letter(col)}${ENV_ROW0 + N_ENV - 1}")
    ws.freeze_panes = "A4"
    return ws


def vertical_inputs(wb, sheet, heading, sub, ids, items, defaults, prefix):
    ws = wb.create_sheet(sheet)
    title(ws, heading, sub)
    g = Grid(ws, ids, prefix=prefix, how_label="Guidance")
    for item in items:
        if item[0] == "section":
            g.section(item[1])
            continue
        key, label, unit, how, ref, dv_list, fmt, comment = item
        vals = defaults.get(key, [None] * len(ids))
        g.add(key, label, unit, how, ref, None, kind="input", fmt=fmt, comment=comment, values=vals)
        if dv_list:
            dv = DataValidation(type="list", formula1=dv_list, allow_blank=True)
            dv.error, dv.errorTitle = "Choose a value from the list.", "Invalid entry"
            ws.add_data_validation(dv)
            r = g.ref(key)
            dv.add(f"{g.cols[0]}{r}:{g.cols[-1]}{r}")
    return ws, g


KC_ITEMS = [
    ("section", "A. Identification"),
    ("name", "Cluster name", "text", "Name the customer will use with Confluent, for example cc-prd.", "-", None, None, None),
    ("region", "Azure region", "text", "Leave blank to use the region on the Setup sheet.", "-", None, None, None),
    ("section", "B. Service requirements (customer decisions)"),
    ("sla", "Required availability SLA", "%",
     "99.5 % is only offered on Basic. 99.9 % needs 1 eCKU on Standard or Enterprise; 99.99 % needs 2 eCKU. "
     "Dedicated gives 99.95 % single-zone or 99.99 % multi-zone (2 CKU or more).", "C1", '"99.5%,99.9%,99.95%,99.99%"',
     None, "Choose the availability the business needs for this cluster. The workbook adds the minimum units Confluent "
           "requires for that SLA."),
    ("priv", "Private networking required", "Yes/No",
     "Yes if AKS must reach Kafka over Azure Private Link. Confluent positions Standard for public networking and "
     "Enterprise for private networking.", "C1", '"Yes,No"', None, None),
    ("byok", "Self-managed encryption keys (BYOK)", "Yes/No",
     "On Azure, self-managed keys are available on Dedicated and Enterprise only, and are fixed at cluster creation.",
     "C4", '"Yes,No"', None, None),
    ("mz", "Multi-zone required (Dedicated)", "Yes/No",
     "Applies only if Dedicated is chosen. Multi-zone needs at least 2 CKU and cannot be changed after creation.",
     "C1", '"Yes,No"', None, None),
    ("pref", "Preferred cluster type", "list",
     "Auto lets the workbook choose the first suitable type in the order Basic, Standard, Enterprise, Dedicated. "
     "Choose a type to force it; the checks show whether it meets the requirements.", "C1",
     '"Auto,Basic,Standard,Enterprise,Dedicated"', None, None),
    ("bill", "Confluent billing method", "list",
     "Sets the Dedicated purchase limit: 4 CKU with credit card, 24 CKU with invoice or marketplace billing. "
     "Higher limits are available by request.", "C1", '"Invoice,Credit card"', None, None),
    ("section", "C. Planning allowance (customer decisions)"),
    ("growth", "Growth allowance over the planning period", "%",
     "Expected growth in partitions, throughput, connections and stored data. Pega publishes no growth figure; "
     "this is a business input.", "W", None, "0%", "Enter the growth you expect over the life of this sizing, for "
                                                   "example 20 % for new applications and case types."),
    ("util", "Target maximum utilisation", "%",
     "Planned capacity = demand x (1 + growth) / target. 70 % matches the average CPU level at which Pega advises "
     "a bigger node for its stream service. Change it if your policy differs.", "P5", None, "0%",
     "A lower value buys more headroom and more units. Do not set above 100 %."),
    ("notes", "Notes", "text", "Decisions or assumptions for this cluster.", "-", None, None, None),
]

SS_ITEMS = [
    ("section", "A. Identification"),
    ("name", "Service name", "text", "Name the customer will use with the search provider, for example os-prd.", "-",
     None, None, None),
    ("region", "Azure region", "text", "Leave blank to use the region on the Setup sheet.", "-", None, None, None),
    ("section", "B. Service requirements (customer decisions)"),
    ("ha", "High availability required", "Yes/No",
     "Yes applies Pega's rule of at least three cluster manager nodes, three data nodes and three SRS pods.",
     "P2", '"Yes,No"', None, None),
    ("zones", "Availability zones", "count",
     "With high availability, data nodes are rounded up to a multiple of the zone count, as OpenSearch advises for "
     "three zones.", "O1", '"1,2,3"', "0", None),
    ("rep", "Replicas per primary shard", "count",
     "OpenSearch default is 1. Each replica is a full extra copy and is placed on a different node from its "
     "primary, so replicas + 1 data nodes are needed. For a single-node test service, as in Pega's test rows, "
     "set 0.", "O2, O5", '"0,1,2"', "0", "SRS may set its own index settings. Confirm with GET _cat/indices after "
                                         "the first SRS build in DEV."),
    ("prof", "Pega sizing profile", "list",
     "Auto picks Large when searchable data reaches the 2 TB of Pega's large example, otherwise Default. Choose "
     "Minimum, Default or Large to override.", "P2", '"Auto,Minimum,Default,Large"', None, None),
    ("section", "C. Provider storage overheads (replace with the provider's figures)"),
    ("ovh", "Provider storage overhead", "%",
     "Space the provider reserves on each node. 20 % is the AWS OpenSearch Service figure (capped at 20 GiB per "
     "instance). Ask your provider for theirs.", "A1", None, "0%", None),
    ("linux", "Operating system reserved space", "%", "Linux reserves 5 % of the file system by default.", "A1",
     None, "0%", None),
    ("idxovh", "Indexing overhead", "%", "Index size on disk relative to the source data; often about 10 %.", "A1",
     None, "0%", "Replace with a measured figure: compare pri.store.size from _cat/indices with the source data."),
    ("section", "D. Data node size overrides (leave blank to use the Pega profile size)"),
    ("cpu", "Data node vCPU (override)", "vCPU", "Leave blank to use the Pega profile size.", "P2", None, "0", None),
    ("ram", "Data node RAM (override)", "GB", "Leave blank to use the Pega profile size. Java heap is taken as half "
                                             "of RAM.", "P2, O4", None, "0", None),
    ("disk", "Data node disk (override)", "GB", "Leave blank to use the Pega profile size. A larger disk can reduce "
                                               "the node count when storage drives it.", "P2", None, "#,##0", None),
    ("section", "E. Service options"),
    ("ver", "Engine version offered by the provider", "list",
     "Pega supports OpenSearch 1.3, 2.15 and 2.19 with the search-n-reporting-service-os image; 2.15 is Pega's best "
     "practice.", "P2", '"OpenSearch 1.3,OpenSearch 2.15,OpenSearch 2.19,Other"', None, None),
    ("priv", "Private connectivity required", "Yes/No", "For the provider request only.", "-", '"Yes,No"', None, None),
    ("cmk", "Customer-managed encryption key", "Yes/No", "For the provider request only.", "-", '"Yes,No"', None, None),
    ("notes", "Notes", "text", "Decisions or assumptions for this service.", "-", None, None, None),
]

# Kafka demand columns: key, header, unit, kind, default, formula, fmt, comment, width
KD_COLS = [
    ("env", "Environment", "", "calc", None, "=IF({S}!B{sr}=\"\",\"\",{S}!B{sr})", None, None, 13),
    ("cluster", "Kafka cluster", "", "calc", None, "=IF(AND({S}!E{sr}=\"Yes\",{S}!F{sr}<>\"\"),{S}!F{sr},\"\")", None,
     "Taken from the Setup sheet. Blank when the environment is out of scope.", 8),
    ("land", "Pega landscape", "", "calc", None, "=IF({S}!D{sr}=\"\",\"\",{S}!D{sr})", None, None, 12),
    ("rank", "Landscape rank", "1-4", "calc", None, "=IFERROR(MATCH({land},LANDSCAPES,0),0)", "0",
     "1 Development, 2 Testing, 3 Stage, 4 Production. Used to find the most demanding landscape on a cluster.", 8),
    ("status", "Input status", "list", "input", None, None, None,
     "Measured: taken from the running system. Estimated: worked out from design figures. Example: illustrative "
     "only, must be replaced before the provider quote.", 11),
    ("mparts", "Measured leader partitions", "count", "input", None, None, "#,##0",
     "Total partitions of all Pega topics (leaders only, before replication). Take it from the Stream service "
     "landing page in Pega 8.8, or from the Confluent topic list after the first build in DEV. When filled in, "
     "this figure is used instead of the build-up to the right. A count taken on 8.8 does not include topics that "
     "Pega '25 and later add, for example messaging that moved from Hazelcast to Kafka [P8]; measure again in DEV "
     "on Pega '26.", 12),
    ("qps", "Queue processors", "count", "input", None, None, "#,##0",
     "Number of queue processor rules, including Pega's own. Each gets a Kafka topic with six partitions by "
     "default [P3].", 11),
    ("topics", "Other Pega topics", "count", "input", None, None, "#,##0",
     "Stream data sets and other Pega topics, such as the system pulse topic (6 partitions [P6]) and the "
     "messaging topics that replace Hazelcast from Pega '25 [P8].", 10),
    ("ppt", "Partitions per topic", "count", "input", 6, None, "0",
     "Default 6 from Pega 8.7 [P3][P4]. Change it only if the DSS "
     "prconfig/dsm/services/stream/pyTopicPartitionsCount/default is set to another value.", 10),
    ("parts", "Leader partitions used", "count", "calc", None, "=IF({mparts}<>\"\",{mparts},({qps}+{topics})*{ppt})",
     "#,##0", "Measured figure if given, otherwise (queue processors + other topics) x partitions per topic.", 11),
    ("compact", "Of which compacted", "count", "input", None, None, "#,##0",
     "Partitions of topics with cleanup.policy=compact, if any. Confluent limits these separately [C1]. Leave blank "
     "if none.", 10),
    ("msgs", "Peak messages written per second", "msg/s", "input", None, None, "#,##0",
     "Highest sustained rate across all Pega topics in the busiest hour.", 12),
    ("bytes", "Average message size", "bytes", "input", None, None, "#,##0", "Average record size written.", 10),
    ("min", "Measured peak write", "MBps", "input", None, None, "0.00",
     "Optional. Confluent received bytes at peak. Overrides messages x size.", 10),
    ("in", "Peak write used", "MBps", "calc", None, "=IF({min}<>\"\",{min},{msgs}*{bytes}/W_BYTES_MB)", "0.00",
     "Measured figure if given, otherwise messages per second x average size / 1,000,000.", 10),
    ("fan", "Consumer fan-out", "reads per write", "input", 1, None, "0.0",
     "How many times each message is read. 1 when one consumer group reads each topic.", 10),
    ("mout", "Measured peak read", "MBps", "input", None, None, "0.00",
     "Optional. Confluent sent bytes at peak. Overrides write x fan-out.", 10),
    ("out", "Peak read used", "MBps", "calc", None, "=IF({mout}<>\"\",{mout},{in}*{fan})", "0.00", None, 10),
    ("day", "Messages written per day", "count", "input", None, None, "#,##0", "Total messages written in a day.", 13),
    ("mgib", "Measured GiB written per day", "GiB", "input", None, None, "#,##0.0", "Optional. Overrides messages x size.", 10),
    ("gib", "GiB written per day used", "GiB", "calc", None, "=IF({mgib}<>\"\",{mgib},{day}*{bytes}/W_BYTES_GIB)",
     "#,##0.0", None, 10),
    ("gibr", "GiB read per day", "GiB", "calc", None, "=IF(AND({mout}<>\"\",{min}<>\"\",{min}>0),{gib}*{mout}/{min},{gib}*{fan})",
     "#,##0.0", "GiB written x fan-out, or x the measured read/write ratio when both peaks are measured.", 10),
    ("ret", "Retention", "hours", "input", 168, None, "#,##0",
     "How long Kafka keeps messages. Confluent's topic default is 168 hours (retention.ms 604,800,000) [C3]. Pega "
     "does not publish a retention figure; this is a design input.", 9),
    ("comp", "Compression ratio", "x", "input", 1, None, "0.0",
     "Use 1 unless producer compression is set; Confluent keeps the producer's compression (compression.type = "
     "producer) [C3]. Pega notes compression.type can decrease disk space [P1].", 9),
    ("store", "Stored GiB (before replication)", "GiB", "calc", None,
     "=IF({ret}=\"\",0,{gib}*{ret}/24/IF({comp}>0,{comp},1))", "#,##0.0",
     "GiB per day x retention days / compression ratio. Confluent replicates three times [C2].", 11),
    ("conn", "Peak client connections", "count", "input", None, None, "#,##0",
     "Highest number of client connections from all Pega pods. Measure it in DEV from the Confluent cluster "
     "metrics. Pega publishes no figure, so leave blank rather than guess; the checks will show it as not assessed.", 11),
    ("req", "Peak requests per second", "req/s", "input", None, None, "#,##0",
     "Highest Kafka requests per second from the Confluent cluster metrics. Pega publishes no figure.", 11),
    ("msgmax", "Largest message", "bytes", "input", None, None, "#,##0",
     "Largest single record. Pega's default limit is 5,000,000 bytes [P1].", 11),
    ("notes", "Notes", "", "input", None, None, None, None, 24),
]

SD_COLS = [
    ("env", "Environment", "", "calc", None, "=IF({S}!B{sr}=\"\",\"\",{S}!B{sr})", None, None, 13),
    ("svc", "Search service", "", "calc", None, "=IF(AND({S}!E{sr}=\"Yes\",{S}!G{sr}<>\"\"),{S}!G{sr},\"\")", None,
     "Taken from the Setup sheet. Blank when the environment is out of scope.", 8),
    ("land", "Pega landscape", "", "calc", None, "=IF({S}!D{sr}=\"\",\"\",{S}!D{sr})", None, None, 12),
    ("rank", "Landscape rank", "1-4", "calc", None, "=IFERROR(MATCH({land},LANDSCAPES,0),0)", "0", None, 8),
    ("status", "Input status", "list", "input", None, None, None,
     "Measured, Estimated or Example. Example values must be replaced before the provider quote.", 11),
    ("gb", "Searchable data today", "GB", "input", None, None, "#,##0.0",
     "Primary index size, not counting replicas. After the first SRS build in DEV, use GET _cat/indices?v and add "
     "up pri.store.size [A1]. Before that, use the size of the current 8.8 search indexes.", 12),
    ("docs", "Documents today", "count", "input", None, None, "#,##0",
     "Number of indexed documents (docs.count in _cat/indices).", 12),
    ("idx", "Indexes", "count", "input", None, None, "#,##0", "Number of indexes SRS creates for this environment.", 9),
    ("pri", "Primary shards per index", "count", "input", 1, None, "0",
     "OpenSearch default is 1 [O2]. Confirm with _cat/indices after the first SRS build.", 9),
    ("grow", "Annual growth", "%", "input", None, None, "0%", "Expected growth in searchable data per year.", 9),
    ("yrs", "Planning period", "years", "input", None, None, "0", "Years the sizing must last.", 8),
    ("tmp", "Temporary extra data", "GB", "input", None, None, "#,##0",
     "Optional. Extra source data held at the same time, for example a second copy while a full re-index runs. "
     "Leave blank if none.", 10),
    ("fgb", "Searchable data at end of period", "GB", "calc", None, "={gb}*(1+{grow})^{yrs}", "#,##0.0",
     "Today's data x (1 + growth) ^ years.", 12),
    ("fdocs", "Documents at end of period", "count", "calc", None, "={docs}*(1+{grow})^{yrs}", "#,##0", None, 12),
    ("prim", "Primary shards", "count", "calc", None, "={idx}*{pri}", "#,##0", "Indexes x primary shards per index.", 9),
    ("notes", "Notes", "", "input", None, None, None, None, 24),
]


def build_demand(wb, sheet, heading, sub, cols, prefix, example_rows, groups):
    ws = wb.create_sheet(sheet)
    title(ws, heading, sub)
    letters = {k: get_column_letter(i + 1) for i, (k, *_r) in enumerate(cols)}
    for start, end, text in groups:
        a, b = letters[start], letters[end]
        ws.merge_cells(f"{a}5:{b}5")
        c = ws[f"{a}5"]
        c.value, c.fill, c.font, c.alignment = text, FILL["section"], Font(name="Calibri", size=10, bold=True, color=NAVY), CENTER
    header_row(ws, 6, [f"{h}\n({u})" if u else h for _, h, u, *_r in cols], height=58)
    for i, (key, h, u, kind, default, formula, fmt, comment, w) in enumerate(cols):
        col = i + 1
        ws.column_dimensions[get_column_letter(col)].width = w
        if comment:
            note(ws.cell(6, col), comment)
        for e in range(N_ENV):
            r = DEM_ROW0 + e
            cell = ws.cell(r, col)
            if formula:
                ctx = {k: f"{v}{r}" for k, v in letters.items()}
                ctx.update(S=q(S_SETUP), sr=ENV_ROW0 + e)
                cell.value = formula.format(**ctx)
                style(cell, "calc", fmt, align=CENTER)
            else:
                v = default
                code = EXAMPLE_ENVS[e][0] if e < len(EXAMPLE_ENVS) else None
                if example_rows and code in example_rows:
                    v = example_rows[code].get(key, v)
                cell.value = v
                style(cell, "input", fmt, align=WRAP if key == "notes" else CENTER)
        define(wb, f"{prefix}_{key.upper()}", f"{q(sheet)}!${letters[key]}${DEM_ROW0}:${letters[key]}${DEM_ROW0 + N_ENV - 1}")
    dv = DataValidation(type="list", formula1='"Measured,Estimated,Example"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{letters['status']}{DEM_ROW0}:{letters['status']}{DEM_ROW0 + N_ENV - 1}")
    ws.freeze_panes = ws.cell(DEM_ROW0, 3)
    for e in range(N_ENV):
        ws.row_dimensions[DEM_ROW0 + e].height = 18
    r = DEM_ROW0 + N_ENV + 1
    ws.cell(r, 1, "Yellow cells are inputs; grey cells are worked out. Hover over a column heading for guidance and "
                  "the source. Where a measured figure is given it replaces the estimate beside it.").font = F_SUB
    return ws, letters


def kd_example():
    keys = ["mparts", "qps", "topics", "ppt", "compact", "msgs", "bytes", "min", "fan", "mout", "day", "mgib", "ret",
            "comp", "conn", "req", "msgmax"]
    out = {}
    for env, vals in KD_EXAMPLE.items():
        d = dict(zip(keys, vals))
        d["status"] = "Example"
        d["notes"] = "Illustrative figures for the worked example. Replace with measured values."
        out[env] = d
    return out


def sd_example():
    keys = ["gb", "docs", "idx", "pri", "grow", "yrs", "tmp"]
    out = {}
    for env, vals in SD_EXAMPLE.items():
        d = dict(zip(keys, vals))
        d["status"] = "Example"
        d["notes"] = "Illustrative figures for the worked example. Replace with measured values."
        out[env] = d
    return out


# --------------------------------------------------------------------------------------- reference data

def build_ref(wb):
    ws = wb.create_sheet(S_REF)
    title(ws, "Ref_Data: published figures used by the formulas",
          "Every value below has a named range used in the formulas, the source ID and the quoted wording. "
          f"Checked {R.CHECKED}. Do not edit unless the source changes; record any change in the Change Log.")
    widths(ws, [22, 16, 52, 14, 14, 8, 96])
    header_row(ws, 4, ["Name (named range)", "Group", "Parameter", "Value", "Unit", "Ref", "Quoted evidence"])
    r = 5
    for name, group, param, value, unit, ref, quote in R.SCALARS:
        vals = [name, group, param, value, unit, ref, quote]
        for j, v in enumerate(vals):
            c = ws.cell(r, j + 1, v)
            style(c, "ref" if j == 3 else "white", "0%" if unit == "fraction" else ("#,##0" if isinstance(v, (int, float)) and j == 3 and v >= 1000 else None),
                  align=CENTER if j in (3, 4, 5) else WRAP)
        define(wb, name, f"{q(S_REF)}!$D${r}")
        fit_height(ws, r, [(param, 52), (quote, 96)])
        r += 1

    ws.freeze_panes = "A5"
    protect(ws)
    return ws


def build_ref_tables(wb):
    ws = wb.create_sheet(S_TAB)
    title(ws, "Ref_Tables: published tables used by the formulas",
          f"Lookup tables with named ranges. Checked {R.CHECKED}. Sources in square brackets are on the References "
          "sheet.")
    widths(ws, [26, 13, 13, 15, 14, 14, 14, 12, 16, 16, 62])
    r = 4
    ws.cell(r, 1, "Confluent Cloud limits per eCKU or CKU, and cluster-level limits [C1][C3]").font = F_BOLD
    r += 1
    header_row(ws, r, ["Cluster type"] + R.CF_COLUMNS + ["Notes"], height=72)
    hdr = r
    r += 1
    first = r
    for t in R.CF_TYPES:
        style(ws.cell(r, 1, t), "white", bold=True)
        for j, v in enumerate(R.CF_TABLE[t]):
            style(ws.cell(r, 2 + j, v), "ref", "#,##0", align=CENTER)
        style(ws.cell(r, 2 + len(R.CF_COLUMNS), R.CF_NOTES[t]), "white")
        fit_height(ws, r, [(R.CF_NOTES[t], 62)])
        r += 1
    define(wb, "CF_TYPES", f"{q(S_TAB)}!$A${first}:$A${r - 1}")
    for j, name in enumerate(R.CF_NAMES):
        col = get_column_letter(2 + j)
        define(wb, name, f"{q(S_TAB)}!${col}${first}:${col}${r - 1}")
    note(ws.cell(hdr, 2), "Per eCKU (Basic, Standard, Enterprise) or per CKU (Dedicated). Units needed for a "
                          "dimension = planned demand / this figure, rounded up [C1].")
    note(ws.cell(hdr, 8), "Basic 50, Standard 10, Enterprise 32 eCKU (also the Azure Private Link limit for "
                          "Enterprise). Dedicated shows the 24-CKU invoice purchase limit; Azure supports up to "
                          "100 CKU by request [C1].")
    note(ws.cell(hdr, 9), "Topic max.message.bytes maximum: 8,388,608 on Basic and Standard; 20,971,520 on "
                          "Dedicated and Enterprise [C3].")
    ws.cell(r, 1, "Freight clusters are not assessed: Confluent describes them for workloads that tolerate relaxed "
                  "latency [C1], and Azure self-managed keys are not offered on Freight [C4].").font = F_SUB
    r += 2

    ws.cell(r, 1, "Pega search sizing tables (Tables 2, 3 and 4) [P2]").font = F_BOLD
    r += 1
    header_row(ws, r, ["Key (profile|role|group)", "Instances", "CPU per instance", "RAM GB per instance",
                       "Storage GB per instance", "Pega table"], height=32)
    r += 1
    first = r
    for key, n, cpu, ram, disk, tbl in R.PS_ROWS:
        style(ws.cell(r, 1, key), "white")
        for j, v in enumerate([n, cpu, ram, disk]):
            style(ws.cell(r, 2 + j, v), "ref", "#,##0", align=CENTER)
        style(ws.cell(r, 6, tbl), "white", align=CENTER)
        r += 1
    for name, col in [("PS_KEY", "A"), ("PS_N", "B"), ("PS_CPU", "C"), ("PS_RAM", "D"), ("PS_DISK", "E")]:
        define(wb, name, f"{q(S_TAB)}!${col}${first}:${col}${r - 1}")
    for text in ["Group Prod = Pega's 'Production, Stage' rows; Test = 'Testing, Development' rows. SRS figures are "
                 "for pods on AKS, not for the search provider. 0 means N/A in Pega's table.", R.PEGA_TABLE4_NOTE]:
        c = ws.cell(r, 1, text)
        c.font, c.alignment = F_SUB, WRAP
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=11)
        fit_height(ws, r, [(text, 200)])
        r += 1
    r += 1

    ws.cell(r, 1, "Maximum partitions per broker by machine type [P5]").font = F_BOLD
    r += 1
    header_row(ws, r, ["CPU cores", "Memory GiB", "Maximum partitions per broker"], height=44)
    r += 1
    first = r
    for cpu, ram, mx in R.PM_ROWS:
        for j, v in enumerate([cpu, ram, mx]):
            style(ws.cell(r, 1 + j, v), "ref", "#,##0", align=CENTER)
        r += 1
    for name, col in [("PM_CPU", "A"), ("PM_RAM", "B"), ("PM_MAX", "C")]:
        define(wb, name, f"{q(S_TAB)}!${col}${first}:${col}${r - 1}")
    ws.cell(r, 1, "Pega lists the last row as 16+ cores and 64+ GiB. The page is for the embedded stream service; "
                  "it is used here only for the self-managed reference figures.").font = F_SUB
    r += 2

    ws.cell(r, 1, "Lists").font = F_BOLD
    r += 1
    header_row(ws, r, ["Pega landscapes", "Kafka cluster IDs", "Search service IDs", "Supported OpenSearch"])
    r += 1
    first = r
    for i in range(6):
        vals = [R.LANDSCAPES[i] if i < 4 else None, KIDS[i], SIDS[i], R.OS_VERSIONS[i] if i < 4 else None]
        for j, v in enumerate(vals):
            if v is not None:
                style(ws.cell(r + i, 1 + j, v), "ref", align=CENTER)
    define(wb, "LANDSCAPES", f"{q(S_TAB)}!$A${first}:$A${first + 3}")
    define(wb, "KAFKA_IDS", f"{q(S_TAB)}!$B${first}:$B${first + 5}")
    define(wb, "SEARCH_IDS", f"{q(S_TAB)}!$C${first}:$C${first + 5}")
    protect(ws)
    return ws


def build_references(wb):
    ws = wb.create_sheet(S_REFS)
    title(ws, "References", f"Sources for every figure and rule in the workbook. Links checked {R.CHECKED}.")
    widths(ws, [7, 20, 52, 70, 70])
    header_row(ws, 4, ["ID", "Publisher", "Title", "Link", "Used for"])
    r = 5
    for rid, (pub, ttl, url, use) in R.REFERENCES.items():
        vals = [rid, pub, ttl, url, use]
        for j, v in enumerate(vals):
            c = ws.cell(r, j + 1, v)
            style(c, "white", bold=(j == 0), align=CENTER if j == 0 else WRAP)
        ws.cell(r, 4).hyperlink = url
        ws.cell(r, 4).font = F_LINK
        fit_height(ws, r, [(ttl, 52), (url, 70), (use, 70)])
        r += 1
    style(ws.cell(r, 1, "W"), "white", bold=True, align=CENTER)
    style(ws.cell(r, 2, "This workbook"), "white")
    style(ws.cell(r, 3, "Workbook rule or customer decision"), "white")
    style(ws.cell(r, 4, "-"), "white")
    style(ws.cell(r, 5, "A rule the workbook applies that no vendor publishes, such as the order in which cluster "
                        "types are tried or the units used for conversions. Each one is labelled W where it is used."),
          "white")
    fit_height(ws, r, [(ws.cell(r, 5).value, 70)])
    r += 2
    ws.cell(r, 1, "AWS guidance (A1 to A3) is for Amazon OpenSearch Service. It is used because Pega does not publish "
                  "a storage or shard formula; ask your managed OpenSearch provider for its own overhead figures "
                  "and replace them on the Search Services sheet.").font = F_SUB
    protect(ws)


# ----------------------------------------------------------------------------------------- Kafka sizing

def env_list(dem_sheet, id_col, env_col, ident):
    parts = [f'IF({q(dem_sheet)}!${id_col}${DEM_ROW0 + e}={ident},{q(dem_sheet)}!${env_col}${DEM_ROW0 + e}&" ","")'
             for e in range(N_ENV)]
    return "=TRIM(" + "&".join(parts) + ")"


def build_kafka_sizing(wb, kc, kd_letters):
    ws = wb.create_sheet(S_KS)
    title(ws, "6  Kafka sizing: Confluent Cloud cluster type, units and pricing quantities",
          "One column per Kafka cluster. Every figure is a formula; the 'How it is worked out' column and the Ref "
          "column explain each step. Green rows are the results.")
    ext = {f"KC_{k}": (S_KC, r) for k, r in kc.rows.items()}
    g = Grid(ws, KIDS, ext=ext, prefix="K.")
    g.section("A. Environments on this cluster")
    g.add("envs", "Environments assigned", "count", "In-scope environments assigned to this cluster on the Setup sheet.",
          "-", "=COUNTIF(KD_CLUSTER,{id})", fmt="0")
    g.add("envlist", "Environment codes", "text", "", "-",
          lambda c: env_list(S_KD, kd_letters["cluster"], kd_letters["env"], f"{c}${ID_ROW}"))
    g.add("name", "Cluster name", "text", "From the Kafka Clusters sheet.", "-",
          "=IF({KC_name}=\"\",{id},{KC_name})")
    g.add("rank", "Most demanding landscape (rank)", "1-4", "1 Development, 2 Testing, 3 Stage, 4 Production.", "-",
          "=SUMPRODUCT(MAX((KD_CLUSTER={id})*KD_RANK))", fmt="0")
    g.add("land", "Most demanding landscape", "text", "", "P1",
          "=IF({rank}=0,\"-\",INDEX(LANDSCAPES,{rank}))")
    g.add("profile", "Pega Kafka sizing row used", "text",
          "Production or Stage use Pega's Production row; Testing or Development use the Development row (W).",
          "P1, W", "=IF({rank}=0,\"-\",IF({rank}>=3,\"Production\",\"Development\"))")

    g.section("B. Demand: sum of the environments on this cluster (from the Kafka Demand sheet)")
    g.add("d_part", "Leader partitions", "count",
          "Sum of leader partitions. Confluent counts partitions before replication [C1].", "C1, P3",
          "=SUMIF(KD_CLUSTER,{id},KD_PARTS)", fmt="#,##0")
    g.add("d_compact", "Compacted partitions", "count", "Sum of compacted partitions.", "C1",
          "=SUMIF(KD_CLUSTER,{id},KD_COMPACT)", fmt="#,##0")
    g.add("d_in", "Peak write (ingress)", "MBps",
          "Sum of environment peaks. This is cautious because the peaks may not happen at the same time.", "-",
          "=SUMIF(KD_CLUSTER,{id},KD_IN)", fmt="0.00")
    g.add("d_out", "Peak read (egress)", "MBps", "Sum of environment peaks.", "-",
          "=SUMIF(KD_CLUSTER,{id},KD_OUT)", fmt="0.00")
    g.add("d_conn", "Peak client connections", "count", "Sum of environment peaks.", "-",
          "=SUMIF(KD_CLUSTER,{id},KD_CONN)", fmt="#,##0")
    g.add("d_req", "Peak requests per second", "req/s", "Sum of environment peaks.", "-",
          "=SUMIF(KD_CLUSTER,{id},KD_REQ)", fmt="#,##0")
    g.add("d_msg", "Largest message", "bytes", "Largest of the environment figures.", "-",
          "=SUMPRODUCT(MAX((KD_CLUSTER={id})*KD_MSGMAX))", fmt="#,##0")
    g.add("d_gib", "GiB written per day", "GiB", "Sum of environments.", "-",
          "=SUMIF(KD_CLUSTER,{id},KD_GIB)", fmt="#,##0.0")
    g.add("d_gibr", "GiB read per day", "GiB", "Sum of environments.", "-",
          "=SUMIF(KD_CLUSTER,{id},KD_GIBR)", fmt="#,##0.0")
    g.add("d_store", "Stored GiB before replication", "GiB", "Sum of environments (GiB per day x retention).", "-",
          "=SUMIF(KD_CLUSTER,{id},KD_STORE)", fmt="#,##0.0")
    g.add("n_ex", "Environments with Example inputs", "count", "", "-",
          "=SUMPRODUCT((KD_CLUSTER={id})*(KD_STATUS=\"Example\"))", fmt="0")
    g.add("n_nostat", "Environments without an input status", "count", "", "-",
          "=SUMPRODUCT((KD_CLUSTER={id})*(KD_STATUS=\"\"))", fmt="0")
    g.add("n_noconn", "Environments without a connections figure", "count", "", "-",
          "=SUMPRODUCT((KD_CLUSTER={id})*(KD_CONN=\"\"))", fmt="0")
    g.add("n_noreq", "Environments without a requests figure", "count", "", "-",
          "=SUMPRODUCT((KD_CLUSTER={id})*(KD_REQ=\"\"))", fmt="0")

    g.add("ready", "Demand entered", "1/0", "1 when the cluster has environments and a partition or throughput "
                                            "figure; results are shown only then.", "-",
          "=IF(AND({envs}>0,OR({d_part}>0,{d_in}>0)),1,0)", fmt="0")

    g.section("C. Planning allowance")
    g.add("growth", "Growth allowance", "%", "From the Kafka Clusters sheet.", "W",
          "=IF({KC_growth}=\"\",0,{KC_growth})", fmt="0%")
    g.add("util", "Target maximum utilisation", "%", "From the Kafka Clusters sheet. 70 % matches Pega's CPU action point.",
          "P5", "=IF({KC_util}=\"\",P_CPU_ACT,{KC_util})", fmt="0%")
    g.add("factor", "Planning factor", "x", "(1 + growth) / target utilisation.", "W",
          "=IF({util}>0,(1+{growth})/{util},0)", fmt="0.00")
    g.add("p_part", "Planned leader partitions", "count", "Leader partitions x planning factor, rounded up.", "W",
          "=ROUNDUP({d_part}*{factor},0)", fmt="#,##0")
    g.add("p_compact", "Planned compacted partitions", "count", "Compacted partitions x planning factor, rounded up.",
          "W", "=ROUNDUP({d_compact}*{factor},0)", fmt="#,##0")
    g.add("p_in", "Planned ingress", "MBps", "Peak write x planning factor.", "W", "={d_in}*{factor}", fmt="0.00")
    g.add("p_out", "Planned egress", "MBps", "Peak read x planning factor.", "W", "={d_out}*{factor}", fmt="0.00")
    g.add("p_conn", "Planned client connections", "count", "Connections x planning factor, rounded up.", "W",
          "=ROUNDUP({d_conn}*{factor},0)", fmt="#,##0")
    g.add("p_req", "Planned requests per second", "req/s", "Requests x planning factor, rounded up.", "W",
          "=ROUNDUP({d_req}*{factor},0)", fmt="#,##0")

    g.section("D. Requirements from the Kafka Clusters sheet")
    g.add("sla", "Required SLA", "%", "Converted from the list choice.", "C1",
          "=IFERROR(VALUE(SUBSTITUTE({KC_sla},\"%\",\"\"))/100,0)", fmt="0.00%")
    g.add("priv", "Private networking", "Yes/No", "", "C1", "=IF({KC_priv}=\"\",\"No\",{KC_priv})")
    g.add("byok", "Self-managed keys", "Yes/No", "", "C4", "=IF({KC_byok}=\"\",\"No\",{KC_byok})")
    g.add("mz", "Multi-zone (Dedicated)", "Yes/No", "", "C1", "=IF({KC_mz}=\"\",\"Yes\",{KC_mz})")
    g.add("pref", "Preferred type", "list", "", "-", "=IF({KC_pref}=\"\",\"Auto\",{KC_pref})")
    g.add("bill", "Billing method", "list", "", "C1", "=IF({KC_bill}=\"\",\"Invoice\",{KC_bill})")

    g.section("E. Units each cluster type would need (per-unit limits are on Ref_Tables)")
    rules = {
        "Basic": ("=IF(AND({sla}<=0.995,{priv}<>\"Yes\",{byok}<>\"Yes\"),\"Yes\",\"No\")",
                  "Only for a 99.5 % SLA, public networking and no self-managed keys.", "C1, C4", "1"),
        "Standard": ("=IF(AND({sla}<=0.9999,{priv}<>\"Yes\",{byok}<>\"Yes\"),\"Yes\",\"No\")",
                     "Public networking only; no self-managed keys on Azure.", "C1, C4",
                     "IF({sla}>0.999,2,1)"),
        "Enterprise": ("=IF({sla}<=0.9999,\"Yes\",\"No\")", "Public or private networking; self-managed keys "
                                                            "available on Azure.", "C1, C4", "IF({sla}>0.999,2,1)"),
        "Dedicated": ("=IF({sla}<=0.9999,\"Yes\",\"No\")", "Public or private networking; self-managed keys "
                                                           "available on Azure.", "C1, C4",
                      "IF(OR({sla}>0.9995,{mz}=\"Yes\"),2,1)"),
    }
    for i, t in enumerate(R.CF_TYPES, 1):
        k = t[0]
        elig, how, ref, slamin = rules[t]
        g.add(f"e_{k}", f"{t}: eligible", "Yes/No", how, ref, elig)
        g.add(f"m_{k}", f"{t}: SLA minimum units", "units",
              {"Basic": "1 eCKU.", "Standard": "2 eCKU for 99.99 %, otherwise 1.",
               "Enterprise": "2 eCKU for 99.99 %, otherwise 1.",
               "Dedicated": "2 CKU for 99.99 % or multi-zone, otherwise 1."}[t], "C1", "=" + slamin, fmt="0")
        g.add(f"u_{k}", f"{t}: units needed", "units",
              "Largest of: each planned dimension / the per-unit limit, rounded up, and the SLA minimum.", "C1",
              "=IF({envs}=0,0,MAX(ROUNDUP({p_in}/INDEX(CF_IN,%d),0),ROUNDUP({p_out}/INDEX(CF_OUT,%d),0),"
              "ROUNDUP({p_part}/INDEX(CF_PART,%d),0),ROUNDUP({p_compact}/INDEX(CF_COMPACT,%d),0),"
              "ROUNDUP({p_conn}/INDEX(CF_CONN,%d),0),ROUNDUP({p_req}/INDEX(CF_REQ,%d),0),{m_%s}))" % (i, i, i, i, i, i, k),
              fmt="0")
        if t == "Dedicated":
            g.add(f"x_{k}", f"{t}: maximum units", "units", "Purchase limit: 4 CKU with credit card, 24 with invoice.",
                  "C1", "=IF({bill}=\"Credit card\",CF_DED_CARD,CF_DED_INVOICE)", fmt="0")
        else:
            g.add(f"x_{k}", f"{t}: maximum units", "units", "Cluster maximum for the type.", "C1",
                  f"=INDEX(CF_MAXU,{i})", fmt="0")
        g.add(f"f_{k}", f"{t}: fits within maximum", "Yes/No", "", "C1",
              "=IF({u_%s}<={x_%s},\"Yes\",\"No\")" % (k, k))

    g.section("F. Recommended Confluent Cloud cluster")
    g.add("auto", "Type chosen by the workbook rule", "text",
          "First type, in the order Basic, Standard, Enterprise, Dedicated, that is eligible and fits within its "
          "maximum (W). The order follows Confluent's 'best for' guidance for each type.", "C1, W",
          "=IF({envs}=0,\"Not used\",IF(AND({e_B}=\"Yes\",{f_B}=\"Yes\"),\"Basic\",IF(AND({e_S}=\"Yes\",{f_S}=\"Yes\"),"
          "\"Standard\",IF(AND({e_E}=\"Yes\",{f_E}=\"Yes\"),\"Enterprise\",\"Dedicated\"))))")
    g.add("type", "Recommended cluster type", "text", "The preferred type if one is set, otherwise the rule above.",
          "C1", "=IF({envs}=0,\"Not used\",IF({ready}=0,\"Inputs needed\",IF({pref}=\"Auto\",{auto},{pref})))",
          kind="key")
    g.add("idx", "Type position (1 Basic to 4 Dedicated)", "1-4", "", "-",
          "=IFERROR(MATCH({type},CF_TYPES,0),0)", fmt="0")
    g.add("unit", "Capacity unit", "text", "eCKU for elastic types; CKU for Dedicated.", "C1",
          "=IF({idx}=0,\"-\",IF({type}=\"Dedicated\",\"CKU\",\"eCKU\"))")
    g.add("units", "Recommended units", "units", "Units needed for the recommended type.", "C1",
          "=IF({idx}=0,0,CHOOSE({idx},{u_B},{u_S},{u_E},{u_D}))", kind="key", fmt="0")
    g.add("maxu", "Maximum units for the type", "units", "", "C1",
          "=IF({idx}=0,0,CHOOSE({idx},{x_B},{x_S},{x_E},{x_D}))", fmt="0")
    g.add("r_part", "Units for partitions", "units", "Planned leader partitions / partitions per unit.", "C1",
          "=IF({idx}=0,0,ROUNDUP({p_part}/INDEX(CF_PART,{idx}),0))", fmt="0")
    g.add("r_in", "Units for ingress", "units", "Planned ingress / ingress per unit.", "C1",
          "=IF({idx}=0,0,ROUNDUP({p_in}/INDEX(CF_IN,{idx}),0))", fmt="0")
    g.add("r_out", "Units for egress", "units", "Planned egress / egress per unit.", "C1",
          "=IF({idx}=0,0,ROUNDUP({p_out}/INDEX(CF_OUT,{idx}),0))", fmt="0")
    g.add("r_compact", "Units for compacted partitions", "units", "", "C1",
          "=IF({idx}=0,0,ROUNDUP({p_compact}/INDEX(CF_COMPACT,{idx}),0))", fmt="0")
    g.add("r_conn", "Units for client connections", "units", "", "C1",
          "=IF({idx}=0,0,ROUNDUP({p_conn}/INDEX(CF_CONN,{idx}),0))", fmt="0")
    g.add("r_req", "Units for requests", "units", "", "C1",
          "=IF({idx}=0,0,ROUNDUP({p_req}/INDEX(CF_REQ,{idx}),0))", fmt="0")
    g.add("r_sla", "Units for the SLA", "units", "", "C1",
          "=IF({idx}=0,0,CHOOSE({idx},{m_B},{m_S},{m_E},{m_D}))", fmt="0")
    g.add("driver", "Dimension that sets the size", "text",
          "The first dimension, in the order listed above, whose units equal the recommendation.", "C1",
          "=IF({idx}=0,\"-\",IF({units}={r_part},\"Partitions\",IF({units}={r_in},\"Ingress\",IF({units}={r_out},"
          "\"Egress\",IF({units}={r_compact},\"Compacted partitions\",IF({units}={r_conn},\"Client connections\","
          "IF({units}={r_req},\"Requests\",\"SLA minimum\")))))))", kind="key")
    g.add("ut_part", "Partition use at today's demand", "%", "Leader partitions / (units x partitions per unit).",
          "C1", "=IF({units}=0,0,{d_part}/({units}*INDEX(CF_PART,{idx})))", fmt="0%")
    g.add("ut_in", "Ingress use at today's peak", "%", "", "C1",
          "=IF({units}=0,0,{d_in}/({units}*INDEX(CF_IN,{idx})))", fmt="0%")
    g.add("ut_out", "Egress use at today's peak", "%", "", "C1",
          "=IF({units}=0,0,{d_out}/({units}*INDEX(CF_OUT,{idx})))", fmt="0%")
    g.add("ut_conn", "Connection use at today's peak", "%", "", "C1",
          "=IF({units}=0,0,{d_conn}/({units}*INDEX(CF_CONN,{idx})))", fmt="0%")
    g.add("ut_req", "Request use at today's peak", "%", "", "C1",
          "=IF({units}=0,0,{d_req}/({units}*INDEX(CF_REQ,{idx})))", fmt="0%")
    g.add("pmin", "Time to create today's partitions", "minutes",
          "Leader partitions / partition creations allowed per 5 minutes x 5. Matters when topics are created "
          "during cutover. Uses today's partitions.", "C1", "=IF({idx}=0,0,{d_part}/INDEX(CF_PARTOPS,{idx})*5)", fmt="#,##0")

    g.section("G. Quantities for the Confluent price quote (per month, with the growth allowance)")
    g.add("b_uh", "Capacity unit-hours (upper bound)", "unit-hours",
          "Recommended units x 730 hours. Elastic types are billed on the highest units used in each hour, so the "
          "bill can be lower.", "C2, W", "={units}*W_HOURS", fmt="#,##0", kind="key")
    g.add("b_in", "Ingress", "GiB per month", "GiB written per day x (1 + growth) x 730 / 24.", "C2, W",
          "={d_gib}*(1+{growth})*W_HOURS/24", fmt="#,##0", kind="key")
    g.add("b_out", "Egress", "GiB per month", "GiB read per day x (1 + growth) x 730 / 24.", "C2, W",
          "={d_gibr}*(1+{growth})*W_HOURS/24", fmt="#,##0", kind="key")
    g.add("b_pre", "Average stored data before replication", "GiB", "Stored GiB x (1 + growth).", "C2",
          "={d_store}*(1+{growth})", fmt="#,##0")
    g.add("b_post", "Billed storage after replication", "GiB", "Before-replication figure x 3. Confluent bills "
                                                               "storage after replication.", "C2",
          "={b_pre}*CF_STORAGE_X", fmt="#,##0", kind="key")
    g.add("b_gbh", "Storage", "GiB-hours per month", "Billed storage x 730.", "C2, W", "={b_post}*W_HOURS",
          fmt="#,##0", kind="key")

    g.section("H. Pega reference sizing for an equivalent self-managed cluster (for comparison only)")
    g.add("k_n", "Brokers", "count", "Pega Table 1 instance count.", "P1",
          "=IF({envs}=0,0,IF({profile}=\"Production\",P_KAFKA_PRD_N,P_KAFKA_DEV_N))", fmt="0")
    g.add("k_cpu_t", "CPU per broker from Table 1", "cores", "", "P1",
          "=IF({envs}=0,0,IF({profile}=\"Production\",P_KAFKA_PRD_CPU,P_KAFKA_DEV_CPU))", fmt="0")
    g.add("k_ram_t", "RAM per broker from Table 1", "GB", "", "P1",
          "=IF({envs}=0,0,IF({profile}=\"Production\",P_KAFKA_PRD_RAM,P_KAFKA_DEV_RAM))", fmt="0")
    g.add("k_disk_t", "Storage per broker from Table 1", "GB", "", "P1",
          "=IF({envs}=0,0,IF({profile}=\"Production\",P_KAFKA_PRD_DISK,P_KAFKA_DEV_DISK))", fmt="#,##0")
    g.add("k_rpb", "Partition replicas per broker", "count",
          "Planned leader partitions x replication factor 3 / brokers. Pega's Helm chart allows a replication "
          "factor of at most 3 and Confluent uses 3.", "P7, C3",
          "=IF({k_n}=0,0,ROUNDUP({p_part}*CF_RF/{k_n},0))", fmt="#,##0")
    g.add("k_cpu_p", "CPU for that partition count", "cores",
          "Smallest machine in Pega's partitions-per-broker table that supports the count.", "P5",
          "=IF({k_n}=0,0,IF({k_rpb}<=INDEX(PM_MAX,1),INDEX(PM_CPU,1),IF({k_rpb}<=INDEX(PM_MAX,2),INDEX(PM_CPU,2),"
          "IF({k_rpb}<=INDEX(PM_MAX,4),INDEX(PM_CPU,4),INDEX(PM_CPU,5)))))", fmt="0")
    g.add("k_ram_p", "RAM for that partition count", "GiB", "", "P5",
          "=IF({k_n}=0,0,IF({k_rpb}<=INDEX(PM_MAX,1),INDEX(PM_RAM,1),IF({k_rpb}<=INDEX(PM_MAX,2),INDEX(PM_RAM,2),"
          "IF({k_rpb}<=INDEX(PM_MAX,4),INDEX(PM_RAM,4),INDEX(PM_RAM,5)))))", fmt="0")
    g.add("k_cpu", "Reference CPU per broker", "cores", "Larger of the two CPU figures.", "P1, P5",
          "=MAX({k_cpu_t},{k_cpu_p})", fmt="0", kind="key")
    g.add("k_ram", "Reference RAM per broker", "GB", "Larger of the two RAM figures.", "P1, P5",
          "=MAX({k_ram_t},{k_ram_p})", fmt="0", kind="key")
    g.add("k_disk_n", "Storage needed per broker", "GB",
          "Billed storage after replication / brokers / (1 - 30 %), so that 30 % stays free.", "P5",
          "=IF({k_n}=0,0,{b_post}/{k_n}/(1-P_FREE_DISK))", fmt="#,##0")
    g.add("k_disk", "Reference storage per broker", "GB", "Larger of Table 1 storage and the storage needed.",
          "P1, P5", "=MAX({k_disk_t},ROUNDUP({k_disk_n},0))", fmt="#,##0", kind="key")

    g.section("I. Checks (PASS, INFO, WARN, FAIL)")
    chk_first = g.r
    g.add("c_in", "Inputs complete", "status", "Partitions and throughput present; every environment has an input "
                                               "status.", "-",
          "=IF({envs}=0,\"-\",IF({d_part}=0,\"FAIL: no partition figure\",IF({d_in}=0,\"FAIL: no throughput figure\","
          "IF({n_nostat}>0,\"WARN: input status missing for \"&{n_nostat}&\" environment(s)\",\"PASS\"))))", kind="white")
    g.add("c_ex", "No example values", "status", "Example values must be replaced before asking for a quote.", "-",
          "=IF({envs}=0,\"-\",IF({n_ex}>0,\"WARN: \"&{n_ex}&\" environment(s) use example values\",\"PASS\"))",
          kind="white")
    g.add("c_conn", "Connections assessed", "status", "Pega publishes no connection figure, so it must be measured.",
          "-", "=IF({envs}=0,\"-\",IF({n_noconn}>0,\"WARN: not measured for \"&{n_noconn}&\" environment(s)\","
               "\"PASS\"))", kind="white")
    g.add("c_req", "Requests assessed", "status", "Pega publishes no request-rate figure, so it must be measured.",
          "-", "=IF({envs}=0,\"-\",IF({n_noreq}>0,\"WARN: not measured for \"&{n_noreq}&\" environment(s)\","
               "\"PASS\"))", kind="white")
    g.add("c_sla", "SLA selected", "status", "", "C1",
          "=IF({envs}=0,\"-\",IF({sla}=0,\"FAIL: choose the required SLA\",\"PASS\"))", kind="white")
    g.add("c_pref", "Preferred type meets the requirements", "status",
          "A forced type must be eligible for the SLA, networking and key requirements.", "C1, C4",
          "=IF({ready}=0,\"-\",IF({pref}=\"Auto\",\"INFO: type chosen by the workbook rule\",IF({idx}=0,"
          "\"FAIL: unknown type\",IF(CHOOSE({idx},{e_B},{e_S},{e_E},{e_D})=\"No\",\"FAIL: \"&{type}&\" does not "
          "meet the SLA, networking or key requirement\",\"PASS\"))))", kind="white")
    g.add("c_fit", "Units within the type maximum", "status", "", "C1",
          "=IF({ready}=0,\"-\",IF({units}<={maxu},\"PASS\",IF({type}=\"Dedicated\",\"WARN: above the purchase limit; "
          "Azure supports up to \"&CF_DED_AZURE&\" CKU by request\",\"FAIL: \"&{units}&\" units exceed the \"&{maxu}&\""
          " maximum for \"&{type})))", kind="white")
    g.add("c_fast", "Enterprise scaling", "status", "Enterprise scales fast up to 10 eCKU, then about 20 minutes per "
                                                    "eCKU.", "C1",
          "=IF({type}<>\"Enterprise\",\"-\",IF({units}>CF_ENT_FAST,\"INFO: above 10 eCKU, scaling is about \"&"
          "CF_ENT_SCALE_MIN&\" minutes per eCKU\",\"PASS\"))", kind="white")
    g.add("c_msgp", "Largest message within Pega's default", "status",
          "Above 5,000,000 bytes needs producer JVM arguments in the Pega Helm chart [P1].", "P1",
          "=IF({envs}=0,\"-\",IF({d_msg}=0,\"WARN: largest message not given\",IF({d_msg}>P_MSG_MAX,"
          "\"WARN: above 5,000,000 bytes; set -Dstream.producer.max.request.size\",\"PASS\")))", kind="white")
    g.add("c_msgc", "Topic message size setting", "status",
          "Pega needs 5,000,000-byte messages; the Confluent topic default is 2,097,164 bytes, so max.message.bytes "
          "must be raised on Pega topics.", "P1, C3",
          "=IF(OR({ready}=0,{idx}=0),\"-\",IF(MAX({d_msg},P_MSG_MAX)>INDEX(CF_MSGMAX,{idx}),\"FAIL: larger than the \""
          "&{type}&\" topic maximum\",\"INFO: set topic max.message.bytes to at least \"&TEXT(MAX({d_msg},P_MSG_MAX),"
          "\"#,##0\")))", kind="white")
    g.add("c_p1", "Partitions against Pega's Table 1 statement", "status",
          "Pega's Table 1 sizing easily supports up to 1,000 partitions; above that Pega says to increase resources.",
          "P1", "=IF({ready}=0,\"-\",IF({p_part}>P_KAFKA_PARTS,\"INFO: \"&TEXT({p_part},\"#,##0\")&\" planned "
                "partitions, above the 1,000 of Pega's Table 1 sizing\",\"PASS\"))", kind="white")
    g.add("c_pm", "Partitions per broker within Pega's table", "status", "Self-managed reference only.", "P5",
          "=IF({ready}=0,\"-\",IF({k_rpb}>INDEX(PM_MAX,5),\"WARN: \"&TEXT({k_rpb},\"#,##0\")&\" per broker is beyond "
          "Pega's table (4,000)\",\"PASS\"))", kind="white")
    g.add("c_byok", "Self-managed keys coverage", "status", "", "C4",
          "=IF(OR({ready}=0,{byok}<>\"Yes\"),\"-\",IF({type}=\"Enterprise\",\"INFO: Enterprise temporary storage uses "
          "Confluent-owned keys\",IF({type}=\"Dedicated\",\"PASS\",\"FAIL: not available on \"&{type})))",
          kind="white")
    g.add("c_grow", "Growth allowance entered", "status", "", "W",
          "=IF({envs}=0,\"-\",IF({growth}=0,\"INFO: no growth allowance\",\"PASS\"))", kind="white")
    g.add("c_share", "Production isolation", "status", "", "-",
          "=IF({envs}=0,\"-\",IF(AND({rank}=4,{envs}>1),\"INFO: Production shares this cluster with other "
          "environments; confirm this is intended\",\"PASS\"))", kind="white")
    chk_last = g.r - 1
    status_format(ws, f"{g.cols[0]}{chk_first}:{g.cols[-1]}{chk_last}")
    define(wb, "KS_CHECKS", f"{q(S_KS)}!${g.cols[0]}${chk_first}:${g.cols[-1]}${chk_last}")
    g.grey_unused()

    g.r += 1
    g.section("J. Why this size: plain-English explanation per cluster")
    for i, c in enumerate(g.cols):
        r = g.r
        rr = lambda k: f"{c}{g.rows[k]}"
        reason = (f'IF({rr("pref")}<>"Auto","The customer selected "&{rr("type")}&". ",'
                  f'IF({rr("type")}="Basic","Basic is enough for a 99.5% SLA on public networking [C1]. ",'
                  f'IF({rr("type")}="Standard","Standard meets the SLA on public networking without self-managed keys [C1]. ",'
                  f'IF({rr("type")}="Enterprise",IF({rr("priv")}="Yes","Private networking is required, which Confluent '
                  f'positions on Enterprise [C1]. ",IF({rr("byok")}="Yes","Self-managed keys on Azure need Enterprise or '
                  f'Dedicated [C4]. ","The load is above Standard\'s 10-eCKU maximum [C1]. ")),'
                  f'"The elastic types cannot hold the planned load within their maximum units [C1]. "))))')
        lim = lambda n: f'TEXT(INDEX({n},{rr("idx")}),"#,##0")'
        detail = (f'IF({rr("driver")}="Partitions",TEXT({rr("p_part")},"#,##0")&" planned leader partitions against "&{lim("CF_PART")}&" per unit",'
                  f'IF({rr("driver")}="Ingress",TEXT({rr("p_in")},"0.0")&" MBps planned ingress against "&{lim("CF_IN")}&" MBps per unit",'
                  f'IF({rr("driver")}="Egress",TEXT({rr("p_out")},"0.0")&" MBps planned egress against "&{lim("CF_OUT")}&" MBps per unit",'
                  f'IF({rr("driver")}="Compacted partitions",TEXT({rr("p_compact")},"#,##0")&" planned compacted partitions against "&{lim("CF_COMPACT")}&" per unit",'
                  f'IF({rr("driver")}="Client connections",TEXT({rr("p_conn")},"#,##0")&" planned connections against "&{lim("CF_CONN")}&" per unit",'
                  f'IF({rr("driver")}="Requests",TEXT({rr("p_req")},"#,##0")&" planned requests per second against "&{lim("CF_REQ")}&" per unit",'
                  f'"the minimum units Confluent requires for a "&TEXT({rr("sla")}*100,"0.0#")&"% SLA"))))))')
        f = (f'=IF({rr("envs")}=0,"{KIDS[i]} is not used.",IF({rr("ready")}=0,{rr("name")}&" needs partition and '
             f'throughput figures on the Kafka Demand sheet.",{rr("name")}&" ("&{rr("envlist")}&"): "&{rr("type")}&" with "&'
             f'{rr("units")}&" "&{rr("unit")}&". "&{reason}&"The size is set by "&{detail}&" [C1]. Planned figures '
             f'add "&TEXT({rr("growth")},"0%")&" growth and keep planned use at or below "&TEXT({rr("util")},"0%")&" of '
             f'capacity; the workbook applies Pega\'s 70% CPU action point to every dimension [P5][W]. Monthly quantities for pricing: "&TEXT({rr("b_in")},"#,##0")&" GiB in, "&TEXT({rr("b_out")},"#,##0")'
             f'&" GiB out and "&TEXT({rr("b_post")},"#,##0")&" GiB billed storage after replication [C2]. For '
             f'comparison, Pega\'s reference for a self-managed cluster is "&{rr("k_n")}&" brokers of "&{rr("k_cpu")}&" '
             f'CPU, "&{rr("k_ram")}&" GB RAM and "&TEXT({rr("k_disk")},"#,##0")&" GB storage [P1][P5]."))')
        ws.cell(r, 2, KIDS[i])
        style(ws.cell(r, 2), "white", bold=True, align=CENTER)
        ws.cell(r, 3, f)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=COL0 + 5)
        style(ws.cell(r, 3), "key")
        ws.cell(r, 3).font = F_BASE
        ws.row_dimensions[r].height = 80
        g.rows[f"why_{KIDS[i]}"] = r
        g.r += 1
    protect(ws)
    return ws, g


# ---------------------------------------------------------------------------------------- Search sizing

def build_search_sizing(wb, ss, sd_letters):
    ws = wb.create_sheet(S_SZ)
    title(ws, "7  Search sizing: managed OpenSearch nodes, storage and SRS",
          "One column per search service. Pega's tables give the starting size; storage, shard and replica rules can "
          "raise the data-node count. Green rows are the results.")
    ext = {f"SS_{k}": (S_SS, r) for k, r in ss.rows.items()}
    g = Grid(ws, SIDS, ext=ext, prefix="S.")
    g.section("A. Environments on this service")
    g.add("envs", "Environments assigned", "count", "In-scope environments assigned to this service.", "-",
          "=COUNTIF(SD_SVC,{id})", fmt="0")
    g.add("envlist", "Environment codes", "text", "", "-",
          lambda c: env_list(S_SD, sd_letters["svc"], sd_letters["env"], f"{c}${ID_ROW}"))
    g.add("name", "Service name", "text", "", "-", "=IF({SS_name}=\"\",{id},{SS_name})")
    g.add("rank", "Most demanding landscape (rank)", "1-4", "", "-", "=SUMPRODUCT(MAX((SD_SVC={id})*SD_RANK))", fmt="0")
    g.add("land", "Most demanding landscape", "text", "", "P2", "=IF({rank}=0,\"-\",INDEX(LANDSCAPES,{rank}))")
    g.add("group", "Pega table rows used", "text", "Prod = 'Production, Stage'; Test = 'Testing, Development'.", "P2",
          "=IF({rank}=0,\"-\",IF({rank}>=3,\"Prod\",\"Test\"))")

    g.section("B. Demand: sum of the environments on this service (from the Search Demand sheet)")
    g.add("gb", "Searchable data today", "GB", "", "-", "=SUMIF(SD_SVC,{id},SD_GB)", fmt="#,##0.0")
    g.add("fgb", "Searchable data at end of period", "GB", "Sum of each environment's grown figure.", "-",
          "=SUMIF(SD_SVC,{id},SD_FGB)", fmt="#,##0.0")
    g.add("tmp", "Temporary extra data", "GB", "", "-", "=SUMIF(SD_SVC,{id},SD_TMP)", fmt="#,##0.0")
    g.add("docs", "Documents at end of period", "count", "", "-", "=SUMIF(SD_SVC,{id},SD_FDOCS)", fmt="#,##0")
    g.add("idx", "Indexes", "count", "", "-", "=SUMIF(SD_SVC,{id},SD_IDX)", fmt="#,##0")
    g.add("prim", "Primary shards", "count", "", "-", "=SUMIF(SD_SVC,{id},SD_PRIM)", fmt="#,##0")
    g.add("n_ex", "Environments with Example inputs", "count", "", "-",
          "=SUMPRODUCT((SD_SVC={id})*(SD_STATUS=\"Example\"))", fmt="0")
    g.add("n_nostat", "Environments without an input status", "count", "", "-",
          "=SUMPRODUCT((SD_SVC={id})*(SD_STATUS=\"\"))", fmt="0")

    g.add("ready", "Demand entered", "1/0", "1 when the service has environments and a searchable data figure; "
                                            "results are shown only then.", "-",
          "=IF(AND({envs}>0,{gb}>0),1,0)", fmt="0")

    g.section("C. Requirements from the Search Services sheet")
    g.add("ha", "High availability", "Yes/No", "", "P2", "=IF({SS_ha}=\"\",\"Yes\",{SS_ha})")
    g.add("zones", "Availability zones", "count", "", "O1", "=IF({SS_zones}=\"\",3,{SS_zones})", fmt="0")
    g.add("rep", "Replicas per primary", "count", "", "O2", "=IF({SS_rep}=\"\",1,{SS_rep})", fmt="0")
    g.add("ovh", "Provider storage overhead", "%", "", "A1", "=IF({SS_ovh}=\"\",A_SVC_OVH,{SS_ovh})", fmt="0%")
    g.add("linux", "Operating system reserve", "%", "", "A1", "=IF({SS_linux}=\"\",A_LINUX,{SS_linux})", fmt="0%")
    g.add("iovh", "Indexing overhead", "%", "", "A1", "=IF({SS_idxovh}=\"\",A_INDEX_OVH,{SS_idxovh})", fmt="0%")

    g.section("D. Pega sizing profile")
    g.add("prof_auto", "Profile by the workbook rule", "text",
          "Large when searchable data at the end of the period reaches the 2 TB of Pega's large example, otherwise "
          "Default (W). Document count is reported in the checks.", "P2, W",
          "=IF({envs}=0,\"-\",IF({fgb}>=P_SEARCH_LARGE_GB,\"Large\",\"Default\"))")
    g.add("prof", "Profile used", "text", "The override on the Search Services sheet, if set.", "P2",
          "=IF({envs}=0,\"-\",IF({ready}=0,\"Inputs needed\",IF(OR({SS_prof}=\"\",{SS_prof}=\"Auto\"),{prof_auto},"
          "{SS_prof})))", kind="key")
    lk = lambda role, field, grp="{group}": (
        "=IF({envs}=0,0,IFERROR(INDEX(PS_%s,MATCH({prof}&\"|%s|\"&%s,PS_KEY,0)),0))" % (field, role, grp))
    g.add("p_dn", "Pega data nodes", "count", "From the Pega table for the profile and landscape.", "P2",
          lk("Data", "N"), fmt="0")
    g.add("p_dcpu", "Pega data node CPU", "vCPU", "", "P2", lk("Data", "CPU"), fmt="0")
    g.add("p_dram", "Pega data node RAM", "GB", "", "P2", lk("Data", "RAM"), fmt="0")
    g.add("p_ddisk", "Pega data node storage", "GB", "", "P2", lk("Data", "DISK"), fmt="#,##0")
    g.add("p_mn", "Pega cluster manager nodes", "count", "0 means Pega lists none for this landscape.", "P2",
          lk("Master", "N"), fmt="0")
    g.add("p_mcpu", "Cluster manager CPU", "vCPU", "Production row size, used whenever managers are needed.", "P2",
          lk("Master", "CPU", '"Prod"'), fmt="0")
    g.add("p_mram", "Cluster manager RAM", "GB", "", "P2", lk("Master", "RAM", '"Prod"'), fmt="0")
    g.add("p_srs", "Pega SRS pods", "count", "SRS runs on AKS, not at the search provider.", "P2", lk("SRS", "N"), fmt="0")

    g.section("E. Data node size used")
    g.add("cpu", "Data node vCPU", "vCPU", "Override if set, otherwise the Pega profile.", "P2",
          "=IF({SS_cpu}<>\"\",{SS_cpu},{p_dcpu})", fmt="0", kind="key")
    g.add("ram", "Data node RAM", "GB", "Override if set, otherwise the Pega profile.", "P2",
          "=IF({SS_ram}<>\"\",{SS_ram},{p_dram})", fmt="0", kind="key")
    g.add("disk", "Data node disk", "GB", "Override if set, otherwise the Pega profile.", "P2",
          "=IF({SS_disk}<>\"\",{SS_disk},{p_ddisk})", fmt="#,##0", kind="key")
    g.add("heap", "Java heap per data node", "GB", "Half of RAM.", "O4", "={ram}*O_HEAP", fmt="0.0")

    g.section("F. Storage and shards")
    g.add("src", "Source data to hold", "GB", "Searchable data at end of period + temporary extra data.", "-",
          "={fgb}+{tmp}", fmt="#,##0.0")
    g.add("need", "Minimum storage", "GB",
          "Source x (1 + replicas) x (1 + indexing overhead) / (1 - OS reserve) / (1 - provider overhead).", "A1",
          "={src}*(1+{rep})*(1+{iovh})/(1-{linux})/(1-{ovh})", fmt="#,##0", kind="key")
    g.add("need_s", "Cross-check: simplified formula", "GB", "Source x (1 + replicas) x 1.45.", "A1",
          "={src}*(1+{rep})*A_SIMPLE", fmt="#,##0")
    g.add("shards", "Total shards", "count", "Primary shards x (1 + replicas).", "O2",
          "={prim}*(1+{rep})", fmt="#,##0")

    g.section("G. Data node count: the largest of the rules below")
    g.add("n_pega", "By Pega profile and HA rule", "nodes",
          "Pega table count; at least 3 when high availability is required.", "P2",
          "=IF({envs}=0,0,IF({ha}=\"Yes\",MAX({p_dn},P_SEARCH_HA),{p_dn}))", fmt="0")
    g.add("n_store", "By storage", "nodes",
          "Minimum storage / (data node disk x 85 % low watermark), rounded up, so that data stays below the point "
          "where OpenSearch stops allocating shards to a node.", "A1, O3",
          "=IF(OR({envs}=0,{disk}=0),0,ROUNDUP({need}/({disk}*O_WM_LOW),0))", fmt="0")
    g.add("n_shard", "By shards per node", "nodes", "Total shards / 1,000, rounded up.", "O3",
          "=ROUNDUP({shards}/O_SHARDS_NODE,0)", fmt="0")
    g.add("n_heap", "By shards per GiB of heap", "nodes", "Total shards / (25 x heap), rounded up.", "A2",
          "=IF({heap}=0,0,ROUNDUP({shards}/(A_SHARDS_HEAP*{heap}),0))", fmt="0")
    g.add("n_rep", "By replica placement", "nodes",
          "Replicas go on different nodes from their primary, so replicas + 1 nodes are needed for a green cluster.",
          "O5, O6", "=IF({envs}=0,0,{rep}+1)", fmt="0")
    g.add("n_min", "Largest of the rules", "nodes", "", "-",
          "=MAX({n_pega},{n_store},{n_shard},{n_heap},{n_rep})", fmt="0")
    g.add("dn", "Recommended data nodes", "nodes",
          "With high availability and more than one zone, rounded up to a multiple of the zone count.", "O1",
          "=IF({ready}=0,0,IF(AND({ha}=\"Yes\",{zones}>1),ROUNDUP({n_min}/{zones},0)*{zones},{n_min}))", fmt="0",
          kind="key")
    g.add("driver", "Rule that sets the count", "text", "First rule, in the order above, equal to the largest.", "-",
          "=IF({ready}=0,\"-\",IF({n_min}={n_pega},\"Pega profile and HA rule\",IF({n_min}={n_store},\"Storage\","
          "IF({n_min}={n_shard},\"Shards per node\",IF({n_min}={n_heap},\"Shards per GiB of heap\","
          "\"Replica placement\")))))", kind="key")
    g.add("alt_n", "Data nodes the other rules give", "nodes",
          "Shown when storage sets the count: the count without the storage rule, rounded to the zones as above.",
          "-", "=IF(OR({envs}=0,{driver}<>\"Storage\"),0,IF(AND({ha}=\"Yes\",{zones}>1),"
               "ROUNDUP(MAX({n_pega},{n_shard},{n_heap},{n_rep})/{zones},0)*{zones},MAX({n_pega},{n_shard},{n_heap},{n_rep})))",
          fmt="0")
    g.add("alt", "Disk per node to hold the data at that count", "GB",
          "Alternative to more nodes: minimum storage / (that count x 85 % low watermark), rounded up. Confirm the "
          "provider offers this disk size.", "A1, O3",
          "=IF({alt_n}=0,0,ROUNDUP({need}/({alt_n}*O_WM_LOW),0))", fmt="#,##0")

    g.section("H. Cluster managers and SRS")
    g.add("mn", "Recommended cluster manager nodes", "nodes",
          "3 for Production and Stage, or whenever high availability is required; otherwise Pega's test figure.",
          "P2, O1", "=IF({ready}=0,0,IF(OR({group}=\"Prod\",{ha}=\"Yes\"),MAX({p_mn},O_MANAGERS),{p_mn}))", fmt="0",
          kind="key")
    g.add("mcpu", "Cluster manager vCPU", "vCPU", "", "P2", "=IF({mn}=0,0,{p_mcpu})", fmt="0", kind="key")
    g.add("mram", "Cluster manager RAM", "GB", "", "P2", "=IF({mn}=0,0,{p_mram})", fmt="0", kind="key")
    g.add("srs", "SRS pods on AKS (minimum)", "pods", "3 for Production and Stage or with high availability; SRS "
                                                    "autoscales above this.", "P2",
          "=IF({ready}=0,0,IF(OR({group}=\"Prod\",{ha}=\"Yes\"),MAX({p_srs},P_SEARCH_HA),{p_srs}))", fmt="0", kind="key")
    g.add("srscpu", "SRS pod CPU / RAM", "text", "From the Pega table.", "P2",
          "=IF({srs}=0,\"-\",IFERROR(INDEX(PS_CPU,MATCH({prof}&\"|SRS|\"&{group},PS_KEY,0)),0)&\" CPU / \"&"
          "IFERROR(INDEX(PS_RAM,MATCH({prof}&\"|SRS|\"&{group},PS_KEY,0)),0)&\" GB\")")

    g.section("I. Totals for the provider")
    g.add("t_cpu", "Data node vCPU in total", "vCPU", "", "-", "={dn}*{cpu}", fmt="#,##0")
    g.add("t_ram", "Data node RAM in total", "GB", "", "-", "={dn}*{ram}", fmt="#,##0")
    g.add("t_disk", "Data node disk in total", "GB", "", "-", "={dn}*{disk}", fmt="#,##0", kind="key")
    g.add("t_mcpu", "Cluster manager vCPU in total", "vCPU", "", "-", "={mn}*{mcpu}", fmt="#,##0")
    g.add("t_mram", "Cluster manager RAM in total", "GB", "", "-", "={mn}*{mram}", fmt="#,##0")

    g.section("J. Cross-checks")
    g.add("use", "Disk use at end of period", "%",
          "Source x (1 + replicas) x (1 + indexing overhead) / usable disk, where usable = total disk x (1 - OS "
          "reserve) x (1 - provider overhead).", "A1, O3",
          "=IF({t_disk}=0,0,{src}*(1+{rep})*(1+{iovh})/({t_disk}*(1-{linux})*(1-{ovh})))", fmt="0%")
    g.add("spn", "Shards per data node", "count", "", "O3", "=IF({dn}=0,0,{shards}/{dn})", fmt="#,##0")
    g.add("slim", "Shard limit per data node", "count", "Lower of 1,000 and 25 x heap.", "O3, A2",
          "=MIN(O_SHARDS_NODE,A_SHARDS_HEAP*{heap})", fmt="#,##0")
    g.add("avg", "Average primary shard size", "GB", "Source data x (1 + indexing overhead) / primary shards.", "A2",
          "=IF({prim}=0,0,{fgb}*(1+{iovh})/{prim})", fmt="0.00")
    g.add("aws_cpu", "AWS starting point for demanding workloads: vCPU", "vCPU", "2 vCPU per 100 GiB of storage.",
          "A3", "={need}/100*A_VCPU_100", fmt="#,##0")
    g.add("aws_ram", "AWS starting point for demanding workloads: RAM", "GB", "8 GiB per 100 GiB of storage.", "A3",
          "={need}/100*A_RAM_100", fmt="#,##0")

    g.section("K. Checks (PASS, INFO, WARN, FAIL)")
    chk_first = g.r
    g.add("c_in", "Inputs complete", "status", "", "-",
          "=IF({envs}=0,\"-\",IF({gb}=0,\"FAIL: no searchable data figure\",IF({idx}=0,\"WARN: number of indexes not "
          "given\",IF({n_nostat}>0,\"WARN: input status missing for \"&{n_nostat}&\" environment(s)\",\"PASS\"))))",
          kind="white")
    g.add("c_ex", "No example values", "status", "", "-",
          "=IF({envs}=0,\"-\",IF({n_ex}>0,\"WARN: \"&{n_ex}&\" environment(s) use example values\",\"PASS\"))",
          kind="white")
    g.add("c_ha", "High availability for Production and Stage", "status", "Pega: at least three nodes per service "
                                                                         "for high availability.", "P2",
          "=IF({envs}=0,\"-\",IF(AND({group}=\"Prod\",{ha}<>\"Yes\"),\"WARN: high availability not selected for a "
          "Production or Stage service\",\"PASS\"))", kind="white")
    g.add("c_rep", "Replicas can be allocated", "status", "", "O5, O6, A1",
          "=IF({ready}=0,\"-\",IF({dn}<{rep}+1,\"FAIL: too few data nodes for the replicas (yellow cluster)\","
          "IF({rep}=0,\"WARN: no replica; AWS recommends at least one\",\"PASS\")))", kind="white")
    g.add("c_disk", "Disk use below the low watermark", "status", "Above 85 % OpenSearch stops allocating replicas "
                                                                 "to the node.", "O3",
          "=IF({ready}=0,\"-\",IF({use}>O_WM_LOW,\"FAIL: \"&TEXT({use},\"0%\")&\" is above the 85% low watermark\","
          "\"PASS\"))", kind="white")
    g.add("c_shard", "Shards per node within limits", "status", "", "O3, A2",
          "=IF({ready}=0,\"-\",IF({spn}>{slim},\"FAIL: \"&TEXT({spn},\"#,##0\")&\" shards per node is above \"&"
          "TEXT({slim},\"#,##0\"),\"PASS\"))", kind="white")
    g.add("c_size", "Shard size guideline", "status", "AWS guideline 10 to 50 GiB per shard.", "A2",
          "=IF(OR({ready}=0,{prim}=0),\"-\",IF({avg}>A_SHARD_MAX,\"WARN: average shard \"&TEXT({avg},\"0\")&\" GB; "
          "consider more primary shards\",IF({avg}<A_SHARD_MIN,\"INFO: average shard \"&TEXT({avg},\"0.0\")&\" GB, "
          "below the 10 GiB guideline\",\"PASS\")))", kind="white")
    g.add("c_docs", "Document count against Pega's large example", "status", "", "P2",
          "=IF({ready}=0,\"-\",IF(AND({docs}>=P_SEARCH_LARGE_DOCS,{prof}<>\"Large\"),\"INFO: \"&TEXT({docs},\"#,##0\")&"
          "\" documents, above the ~750,000 of Pega's large example; confirm with indexing and query metrics\","
          "\"PASS\"))", kind="white")
    g.add("c_ver", "Engine version", "status", "Pega best practice is OpenSearch 2.15.", "P2",
          "=IF({envs}=0,\"-\",IF({SS_ver}=\"OpenSearch 2.15\",\"PASS\",IF(OR({SS_ver}=\"OpenSearch 1.3\","
          "{SS_ver}=\"OpenSearch 2.19\"),\"INFO: supported; Pega's best practice is 2.15\",IF({SS_ver}=\"\","
          "\"WARN: version not given\",\"FAIL: not in Pega's supported list\"))))", kind="white")
    g.add("c_zone", "Zones for high availability", "status", "", "O1",
          "=IF({envs}=0,\"-\",IF(AND({ha}=\"Yes\",{zones}<3),\"WARN: OpenSearch advises three zones for cluster "
          "managers\",\"PASS\"))", kind="white")
    g.add("c_t4", "Pega Table 4 figures", "status", "", "P2",
          "=IF({ready}=0,\"-\",IF({prof}=\"Large\",\"INFO: Table 4 RAM total and storage need confirming with Pega "
          "(see Ref_Tables)\",\"PASS\"))", kind="white")
    g.add("c_aws", "Compute against the AWS starting point", "status", "Information only; AWS gives it for "
                                                                      "demanding workloads.", "A3",
          "=IF({ready}=0,\"-\",IF({t_cpu}<{aws_cpu},\"INFO: data-node vCPU below the AWS starting point for demanding "
          "workloads\",\"PASS\"))", kind="white")
    g.add("c_share", "Production isolation", "status", "", "-",
          "=IF({envs}=0,\"-\",IF(AND({rank}=4,{envs}>1),\"INFO: Production shares this service with other "
          "environments; confirm this is intended\",\"PASS\"))", kind="white")
    chk_last = g.r - 1
    status_format(ws, f"{g.cols[0]}{chk_first}:{g.cols[-1]}{chk_last}")
    define(wb, "SZ_CHECKS", f"{q(S_SZ)}!${g.cols[0]}${chk_first}:${g.cols[-1]}${chk_last}")
    g.grey_unused()

    g.r += 1
    g.section("L. Why this size: plain-English explanation per service")
    for i, c in enumerate(g.cols):
        r = g.r
        rr = lambda k: f"{c}{g.rows[k]}"
        prof_reason = (f'IF({rr("prof")}<>{rr("prof_auto")}," (selected by the customer)",IF({rr("prof")}="Large",'
                       f'" (searchable data at or above the 2 TB of Pega\'s large example)",'
                       f'" (searchable data below the 2 TB of Pega\'s large example)"))')
        f = (f'=IF({rr("envs")}=0,"{SIDS[i]} is not used.",IF({rr("ready")}=0,{rr("name")}&" needs a searchable '
             f'data figure on the Search Demand sheet.",{rr("name")}&" ("&{rr("envlist")}&"): "&{rr("dn")}&" data '
             f'nodes of "&{rr("cpu")}&" vCPU, "&{rr("ram")}&" GB RAM and "&TEXT({rr("disk")},"#,##0")&" GB disk, plus "&'
             f'{rr("mn")}&" cluster manager nodes of "&{rr("mcpu")}&" vCPU and "&{rr("mram")}&" GB RAM. Pega profile: "&'
             f'{rr("prof")}&{prof_reason}&" [P2]. The data-node count is set by this rule: "&{rr("driver")}&"."&IF({rr("alt")}>0,'
             f'" Alternatively, "&{rr("alt_n")}&" data nodes with at least "&TEXT({rr("alt")},"#,##0")&" GB disk each would hold the data.","")&" '
             f'Storage for "&TEXT({rr("src")},"#,##0")&" GB of source data with "&{rr("rep")}&" replica(s) is "&'
             f'TEXT({rr("need")},"#,##0")&" GB [A1]; disk use at the end of the period is "&TEXT({rr("use")},"0%")&'
             f'" against the 85% low watermark [O3]. "&TEXT({rr("shards")},"#,##0")&" shards give "&'
             f'TEXT({rr("spn")},"#,##0")&" per node against a limit of "&TEXT({rr("slim")},"#,##0")&" [O3][A2]. '
             f'SRS on AKS: at least "&{rr("srs")}&" pods [P2]."))')
        ws.cell(r, 2, SIDS[i])
        style(ws.cell(r, 2), "white", bold=True, align=CENTER)
        ws.cell(r, 3, f)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=COL0 + 5)
        style(ws.cell(r, 3), "key")
        ws.cell(r, 3).font = F_BASE
        ws.row_dimensions[r].height = 80
        g.rows[f"why_{SIDS[i]}"] = r
        g.r += 1
    protect(ws)
    return ws, g


# -------------------------------------------------------------------------------------- provider requests

def link_grid(wb, sheet, heading, sub, ids, rows, ext):
    ws = wb.create_sheet(sheet)
    title(ws, heading, sub)
    g = Grid(ws, ids, ext=ext, prefix="R.")
    for item in rows:
        if item[0] == "section":
            g.section(item[1])
            continue
        key, label, unit, how, ref, formula, fmt = item
        g.add(key, label, unit, how, ref, formula, fmt=fmt, kind="key" if key in ("type", "units", "dn", "cpu", "ram", "disk", "mn") else "calc")
    return ws, g


def static_table(ws, r, heading, header, rows):
    """Fixed table under a grid: first column in B, second in C:D, third in E:K (two columns: C:K)."""
    spans = [(2, 2), (3, 4), (5, COL0 + 5)] if len(header) == 3 else [(2, 2), (3, COL0 + 5)]
    width = {2: 34, 3: 69, 5: 111} if len(header) == 3 else {2: 34, 3: 180}
    ws.cell(r, 2, heading).font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    r += 1
    for (a, b), h in zip(spans, header):
        for col in range(a, b + 1):
            c = ws.cell(r, col, h if col == a else None)
            c.fill, c.font, c.border, c.alignment = FILL["head"], F_HEAD, BOX, CENTER
        if b > a:
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    r += 1
    for row in rows:
        for (a, b), v in zip(spans, row):
            for col in range(a, b + 1):
                style(ws.cell(r, col, v if col == a else None), "white", bold=(a == 2), align=WRAP)
            if b > a:
                ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        fit_height(ws, r, [(str(v), width[a]) for (a, _b), v in zip(spans, row)])
        r += 1
    return r + 1


def build_confluent_request(wb, ks, kc):
    ext = {f"KS_{k}": (S_KS, r) for k, r in ks.rows.items()}
    ext.update({f"KC_{k}": (S_KC, r) for k, r in kc.rows.items()})
    rows = [
        ("section", "A. Cluster"),
        ("envs", "Environments assigned", "count", "", "-", "={KS_envs}", "0"),
        ("name", "Cluster name", "text", "", "-", "={KS_name}", None),
        ("region", "Azure region", "text", "", "-",
         f"=IF({{KS_envs}}=0,\"-\",IF({{KC_region}}<>\"\",{{KC_region}},IF({q(S_SETUP)}!$C$8=\"\",\"(enter region)\",{q(S_SETUP)}!$C$8)))", None),
        ("envs", "Pega environments", "text", "", "-", "={KS_envlist}", None),
        ("type", "Requested cluster type", "text", "", "C1", "={KS_type}", None),
        ("units", "Requested units", "units", "eCKU or CKU as shown below.", "C1", "={KS_units}", "0"),
        ("unit", "Capacity unit", "text", "", "C1", "={KS_unit}", None),
        ("sla", "Required SLA", "%", "", "C1", "={KS_sla}", "0.00%"),
        ("priv", "Private networking (Azure Private Link)", "Yes/No", "", "C1", "={KS_priv}", None),
        ("byok", "Self-managed encryption keys", "Yes/No", "", "C4", "={KS_byok}", None),
        ("mz", "Multi-zone (if Dedicated)", "Yes/No", "", "C1", "={KS_mz}", None),
        ("section", "B. Planned load (includes growth and target utilisation)"),
        ("p_part", "Leader partitions (before replication)", "count", "", "C1", "={KS_p_part}", "#,##0"),
        ("p_compact", "Compacted partitions", "count", "", "C1", "={KS_p_compact}", "#,##0"),
        ("p_in", "Peak ingress", "MBps", "", "C1", "={KS_p_in}", "0.00"),
        ("p_out", "Peak egress", "MBps", "", "C1", "={KS_p_out}", "0.00"),
        ("p_conn", "Peak client connections", "count", "Blank inputs count as 0; see the checks.", "C1",
         "={KS_p_conn}", "#,##0"),
        ("p_req", "Peak requests per second", "req/s", "Blank inputs count as 0; see the checks.", "C1",
         "={KS_p_req}", "#,##0"),
        ("msg", "Largest message", "bytes", "At least Pega's 5,000,000-byte default.", "P1",
         "=MAX({KS_d_msg},P_MSG_MAX)", "#,##0"),
        ("driver", "Dimension that sets the size", "text", "", "C1", "={KS_driver}", None),
        ("section", "C. Monthly quantities for pricing"),
        ("b_uh", "Capacity unit-hours (upper bound)", "unit-hours", "", "C2", "={KS_b_uh}", "#,##0"),
        ("b_in", "Ingress", "GiB", "", "C2", "={KS_b_in}", "#,##0"),
        ("b_out", "Egress", "GiB", "", "C2", "={KS_b_out}", "#,##0"),
        ("b_post", "Billed storage after replication (average)", "GiB", "", "C2", "={KS_b_post}", "#,##0"),
        ("b_gbh", "Storage", "GiB-hours", "", "C2", "={KS_b_gbh}", "#,##0"),
        ("section", "D. Input quality"),
        ("ex", "Example inputs still present", "status", "", "-", "={KS_c_ex}", None),
        ("conn", "Connections measured", "status", "", "-", "={KS_c_conn}", None),
        ("req", "Requests measured", "status", "", "-", "={KS_c_req}", None),
    ]
    ws, g = link_grid(wb, S_CR, "8  Confluent Cloud request: send this sheet to Confluent for sizing and price",
                      "Figures link to the Kafka Sizing sheet. Ask Confluent to confirm the cluster type and units "
                      "and to price the quantities in Section C.", KIDS, rows, ext)
    status_format(ws, f"{g.cols[0]}{g.rows['ex']}:{g.cols[-1]}{g.rows['req']}")
    g.grey_unused()
    r = g.r + 1
    r = static_table(ws, r, "E. Kafka settings Pega requires [P1]", ["Setting", "Value", "Note for Confluent Cloud"],
                     R.PEGA_KAFKA_SETTINGS)
    r = static_table(ws, r, "F. Access control Pega needs for its service account [P1]",
                     ["Resource type", "Pattern", "Operation"], R.PEGA_ACLS)
    r = static_table(ws, r, "G. Client and topic facts", ["Item", "Value", "Ref"], [
        ("Kafka client library in Pega Platform '26", "4.0.0 (backward compatible with 3.9.2)", "P6"),
        ("Supported authentication", "SASL/OAUTHBEARER, SASL/PLAIN or SASL/SCRAM", "P1"),
        ("Replication factor", "3 (Confluent default, not editable); set Pega Helm replicationFactor to 3", "C3, P7"),
        ("Partitions per topic", "6 by default for queue processors and other Pega topics", "P3, P4"),
        ("Topic naming", "stream.streamNamePattern prefix per environment, for example pega-prd-{stream.name}", "P7"),
        ("Vendor-specific security features", "Not supported by Pega", "P1"),
    ])
    static_table(ws, r, "H. Questions for Confluent", ["#", "Question"], [
        ("1", "Confirm the cluster type and units for each cluster, and price the unit-hours, ingress, egress and "
              "storage in Section C for the Azure region."),
        ("2", "How should topic max.message.bytes be set to at least 5,000,000 on topics that Pega creates "
              "dynamically, given the default of 2,097,164? [P1][C3]"),
        ("3", "Confirm that unclean leader election is disabled and that topics are not auto-created on the cluster "
              "[P1]."),
        ("4", "Confirm compatibility with the Kafka 4.0.0 client library shipped with Pega Platform '26 [P6]."),
        ("5", "Confirm Azure Private Link and self-managed key support for the requested type, and any unit limit "
              "that applies with Private Link [C1][C4]."),
        ("6", "Confirm how the partition, connection, connection-attempt and request limits are enforced for the "
              "requested type, including during rolling restarts of all Pega pods [C1]."),
        ("7", "Where the partition count drives the size, confirm the per-unit partition limit and any option to "
              "raise it [C1]."),
    ])
    protect(ws)
    return ws


def build_search_request(wb, sz, ss):
    ext = {f"SZ_{k}": (S_SZ, r) for k, r in sz.rows.items()}
    ext.update({f"SS_{k}": (S_SS, r) for k, r in ss.rows.items()})
    rows = [
        ("section", "A. Service"),
        ("nenv", "Environments assigned", "count", "", "-", "={SZ_envs}", "0"),
        ("name", "Service name", "text", "", "-", "={SZ_name}", None),
        ("region", "Azure region", "text", "", "-",
         f"=IF({{SZ_envs}}=0,\"-\",IF({{SS_region}}<>\"\",{{SS_region}},IF({q(S_SETUP)}!$C$8=\"\",\"(enter region)\",{q(S_SETUP)}!$C$8)))", None),
        ("envs", "Pega environments", "text", "", "-", "={SZ_envlist}", None),
        ("ver", "Engine version requested", "text", "Pega best practice.", "P2",
         f"=IF({{SZ_envs}}=0,\"-\",\"OpenSearch 2.15 (1.3 and 2.19 also supported)\")", None),
        ("ha", "High availability", "Yes/No", "", "P2", "={SZ_ha}", None),
        ("zones", "Availability zones", "count", "", "O1", "={SZ_zones}", "0"),
        ("priv", "Private connectivity", "Yes/No", "", "-", "=IF({SZ_envs}=0,\"-\",{SS_priv})", None),
        ("cmk", "Customer-managed key", "Yes/No", "", "-", "=IF({SZ_envs}=0,\"-\",{SS_cmk})", None),
        ("section", "B. Nodes"),
        ("dn", "Data nodes", "nodes", "", "P2", "={SZ_dn}", "0"),
        ("cpu", "Data node vCPU", "vCPU", "", "P2", "={SZ_cpu}", "0"),
        ("ram", "Data node RAM", "GB", "", "P2", "={SZ_ram}", "0"),
        ("disk", "Data node disk", "GB", "", "P2", "={SZ_disk}", "#,##0"),
        ("mn", "Dedicated cluster manager nodes", "nodes", "", "P2, O1", "={SZ_mn}", "0"),
        ("mcpu", "Cluster manager vCPU", "vCPU", "", "P2", "={SZ_mcpu}", "0"),
        ("mram", "Cluster manager RAM", "GB", "", "P2", "={SZ_mram}", "0"),
        ("section", "C. Data"),
        ("src", "Source data to hold at end of period", "GB", "", "-", "={SZ_src}", "#,##0"),
        ("need", "Minimum storage including replicas and overheads", "GB", "", "A1", "={SZ_need}", "#,##0"),
        ("rep", "Replicas per primary", "count", "", "O2", "={SZ_rep}", "0"),
        ("shards", "Total shards", "count", "", "O2", "={SZ_shards}", "#,##0"),
        ("docs", "Documents at end of period", "count", "", "-", "={SZ_docs}", "#,##0"),
        ("driver", "Rule that sets the data-node count", "text", "", "-", "={SZ_driver}", None),
        ("section", "D. Input quality"),
        ("ex", "Example inputs still present", "status", "", "-", "={SZ_c_ex}", None),
    ]
    ws, g = link_grid(wb, S_OR, "9  Search request: send this sheet to the managed OpenSearch provider",
                      "Figures link to the Search Sizing sheet. Ask the provider to confirm the node sizes, apply its "
                      "own storage overhead and price the service.", SIDS, rows, ext)
    status_format(ws, f"{g.cols[0]}{g.rows['ex']}:{g.cols[-1]}{g.rows['ex']}")
    g.grey_unused("nenv")
    r = g.r + 1
    r = static_table(ws, r, "E. Cluster settings SRS requires [P2]", ["Setting", "Value", "Note"], [
        ("action.auto_create_index", "false", "SRS sets this itself if its user has manage cluster privileges; "
                                              "otherwise set it manually."),
        ("action.destructive_requires_name", "false", "SRS deletes indexes with a pattern."),
        ("Image", "search-n-reporting-service-os", "Supports cloud and self-managed OpenSearch."),
        ("Traffic encryption", "TLS between Pega and SRS", "Pega best practice: service mesh or load balancer with TLS."),
    ])
    static_table(ws, r, "F. Questions for the provider", ["#", "Question"], [
        ("1", "Price the nodes in Section B in the Azure region, with the zones and options requested."),
        ("2", "What storage overhead do you reserve per node? Replace the 20 % AWS figure on the Search Services "
              "sheet with yours [A1]."),
        ("3", "Can the two SRS cluster settings in Section E be applied, and can the SRS user have manage cluster "
              "privileges? [P2]"),
        ("4", "Which OpenSearch versions do you offer, and what is your support timeline for 2.15? [P2]"),
        ("5", "Confirm dedicated cluster manager nodes across three zones and shard allocation awareness by zone [O1]."),
        ("6", "Confirm that the service will not be shared with non-Pega workloads [P2]."),
        ("7", "Which indexing and query metrics can you export so the size can be adjusted after go-live? [P2]"),
    ])
    protect(ws)
    return ws


# ------------------------------------------------------------------------------------------- checks sheet

def build_checks(wb, ks, sz):
    ws = wb.create_sheet(S_CK)
    title(ws, "10  Checks: summary of every check in the workbook",
          "Resolve every FAIL before sending the request sheets. Review every WARN. INFO items are facts the "
          "customer and the provider should know.")
    widths(ws, [22, 10, 13, 11, 11, 11, 11, 32])
    header_row(ws, 4, ["Kafka cluster / search service", "ID", "FAIL", "WARN", "INFO", "PASS", "", "Overall"])
    r = 5
    for g, ids, sheet in [(ks, KIDS, S_KS), (sz, SIDS, S_SZ)]:
        first = min(v for k, v in g.rows.items() if k.startswith("c_"))
        last = max(v for k, v in g.rows.items() if k.startswith("c_"))
        for i, ident in enumerate(ids):
            col = get_column_letter(COL0 + i)
            rng = f"{q(sheet)}!{col}{first}:{col}{last}"
            style(ws.cell(r, 1, f"={q(sheet)}!{col}{g.rows['name']}"), "calc")
            style(ws.cell(r, 2, ident), "calc", align=CENTER)
            for j, word in enumerate(["FAIL", "WARN", "INFO", "PASS"]):
                style(ws.cell(r, 3 + j, f'=SUMPRODUCT(--(LEFT({rng},4)="{word}"))'), "calc", "0", align=CENTER)
            style(ws.cell(r, 7), "calc")
            style(ws.cell(r, 8, f'=IF({q(sheet)}!{col}{g.rows["envs"]}=0,"Not used",IF(C{r}>0,"FAIL: resolve before '
                                f'sending",IF(D{r}>0,"WARN: review",IF(E{r}>0,"INFO only","PASS"))))'), "white", bold=True)
            r += 1
    status_format(ws, f"H5:H{r - 1}")
    r += 1
    ws.cell(r, 1, "Environment input checks").font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    r += 1
    header_row(ws, r, ["Environment", "In scope", "Landscape", "Kafka cluster", "Search service",
                       "Kafka input status", "Search input status", "Overall"], height=32)
    r += 1
    first = r
    for e in range(N_ENV):
        sr, dr = ENV_ROW0 + e, DEM_ROW0 + e
        S = q(S_SETUP)
        style(ws.cell(r, 1, f'=IF({S}!B{sr}="","(empty row)",{S}!B{sr})'), "calc")
        style(ws.cell(r, 2, f'=IF({S}!E{sr}="","-",{S}!E{sr})'), "calc", align=CENTER)
        style(ws.cell(r, 3, f'=IF({S}!D{sr}="","missing",{S}!D{sr})'), "calc", align=CENTER)
        style(ws.cell(r, 4, f'=IF({S}!F{sr}="","missing",{S}!F{sr})'), "calc", align=CENTER)
        style(ws.cell(r, 5, f'=IF({S}!G{sr}="","missing",{S}!G{sr})'), "calc", align=CENTER)
        style(ws.cell(r, 6, f"=IF({q(S_KD)}!E{dr}=\"\",\"missing\",{q(S_KD)}!E{dr})"), "calc", align=CENTER)
        style(ws.cell(r, 7, f"=IF({q(S_SD)}!E{dr}=\"\",\"missing\",{q(S_SD)}!E{dr})"), "calc", align=CENTER)
        style(ws.cell(r, 8, f'=IF({S}!B{sr}="","-",IF({S}!E{sr}<>"Yes","INFO: out of scope",IF(OR({S}!D{sr}="",'
                            f'{S}!F{sr}="",{S}!G{sr}=""),"FAIL: landscape, cluster or service missing",IF(COUNTIF('
                            f'ENV_CODE,{S}!B{sr})>1,"FAIL: duplicate environment code",IF(OR({q(S_KD)}!E{dr}="",'
                            f'{q(S_SD)}!E{dr}=""),"WARN: input status not set","PASS")))))'), "white", bold=True)
        r += 1
    status_format(ws, f"H{first}:H{r - 1}")
    ws.freeze_panes = "A5"
    protect(ws)
    return ws


# ----------------------------------------------------------------------------------- cover, guide, log

def cover_header(ws, row, values):
    for i, v in enumerate(values, 2):
        c = ws.cell(row, i, v)
        c.fill, c.font, c.border, c.alignment = FILL["head"], F_HEAD, BOX, CENTER
    ws.row_dimensions[row].height = 22


def build_cover(wb, ks, sz, example):
    ws = wb.create_sheet(S_COVER, 0)
    ws.sheet_view.showGridLines = False
    widths(ws, [3, 26, 16, 16, 16, 16, 16, 16, 3])
    ws["B2"] = "Pega Platform '26 on Azure AKS"
    ws["B2"].font = Font(name="Calibri", size=12, bold=True, color="595959")
    ws["B3"] = "Kafka and OpenSearch sizing workbook"
    ws["B3"].font = Font(name="Calibri", size=22, bold=True, color=NAVY)
    ws["B4"] = ("Sizes Confluent Cloud clusters and managed OpenSearch services for the Pega '26 externalised Kafka and "
                "search services, from inputs the customer collects, using Pega's published recommendations and the "
                "providers' published limits.")
    ws.merge_cells("B4:H4")
    ws["B4"].alignment = WRAP
    ws["B4"].font = Font(name="Calibri", size=11)
    ws.row_dimensions[4].height = 44
    kind = "WORKED EXAMPLE: illustrative inputs only" if example else "TEMPLATE: enter the customer's inputs"
    ws["B6"] = kind
    ws["B6"].font = Font(name="Calibri", size=12, bold=True, color="9C0006" if example else "006100")
    meta = [("Version", VERSION), ("Date", DATE), ("Customer", f"=IF({q(S_SETUP)}!C4=\"\",\"(enter on Setup)\",{q(S_SETUP)}!C4)"),
            ("Status", f"={q(S_SETUP)}!C11"), ("Evidence checked", R.CHECKED)]
    for i, (k, v) in enumerate(meta):
        ws.cell(7 + i, 2, k).font = F_BOLD
        ws.cell(7 + i, 3, v).font = F_BASE
    r = 13
    ws.cell(r, 2, "Results summary").font = Font(name="Calibri", size=13, bold=True, color=NAVY)
    r += 1
    cover_header(ws, r, ["Kafka cluster", "Environments", "Type", "Units", "Sized by", "Checks"])
    r += 1
    first = r
    for i, ident in enumerate(KIDS):
        col = get_column_letter(COL0 + i)
        vals = [f"={q(S_KS)}!{col}{ks.rows['name']}", f"={q(S_KS)}!{col}{ks.rows['envlist']}",
                f"={q(S_KS)}!{col}{ks.rows['type']}",
                f"=IF({q(S_KS)}!{col}{ks.rows['units']}=0,\"-\",{q(S_KS)}!{col}{ks.rows['units']}&\" \"&{q(S_KS)}!{col}{ks.rows['unit']})",
                f"={q(S_KS)}!{col}{ks.rows['driver']}", f"={q(S_CK)}!H{5 + i}"]
        for j, v in enumerate(vals):
            style(ws.cell(r, 2 + j, v), "key" if j in (2, 3) else "calc", align=CENTER)
        r += 1
    status_format(ws, f"G{first}:G{r - 1}")
    r += 1
    cover_header(ws, r, ["Search service", "Environments", "Data nodes", "Managers", "Sized by", "Checks"])
    r += 1
    first = r
    for i, ident in enumerate(SIDS):
        col = get_column_letter(COL0 + i)
        S = q(S_SZ)
        vals = [f"={S}!{col}{sz.rows['name']}", f"={S}!{col}{sz.rows['envlist']}",
                f"=IF({S}!{col}{sz.rows['dn']}=0,\"-\",{S}!{col}{sz.rows['dn']}&\" x \"&{S}!{col}{sz.rows['cpu']}&\" vCPU / \"&"
                f"{S}!{col}{sz.rows['ram']}&\" GB / \"&{S}!{col}{sz.rows['disk']}&\" GB\")",
                f"=IF({S}!{col}{sz.rows['mn']}=0,\"-\",{S}!{col}{sz.rows['mn']}&\" x \"&{S}!{col}{sz.rows['mcpu']}&\" vCPU / \"&"
                f"{S}!{col}{sz.rows['mram']}&\" GB\")",
                f"={S}!{col}{sz.rows['driver']}", f"={q(S_CK)}!H{11 + i}"]
        for j, v in enumerate(vals):
            style(ws.cell(r, 2 + j, v), "key" if j in (2, 3) else "calc", align=CENTER)
        r += 1
    status_format(ws, f"G{first}:G{r - 1}")
    r += 1
    ws.cell(r, 2, "Colour key").font = Font(name="Calibri", size=13, bold=True, color=NAVY)
    r += 1
    for k, text in [("input", "Input: the customer enters or confirms this value"),
                    ("calc", "Worked out by formula: do not overwrite"),
                    ("key", "Result: carried to the summary and the provider request sheets"),
                    ("ref", "Published reference value (Ref_Data and Ref_Tables)"),
                    ("warn", "WARN: review before sending"), ("fail", "FAIL: must be resolved before sending")]:
        style(ws.cell(r, 2, ""), k)
        ws.cell(r, 3, text).font = F_BASE
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
        r += 1
    r += 1
    ws.cell(r, 2, "Sheets").font = Font(name="Calibri", size=13, bold=True, color=NAVY)
    r += 1
    sheets = [(S_GUIDE, "How to use the workbook and where each input comes from"),
              (S_SETUP, "Programme details; environments and their cluster and service assignment"),
              (S_KC, "Kafka cluster requirements: SLA, networking, keys, growth and utilisation"),
              (S_KD, "Kafka demand per environment: partitions, throughput, retention, connections"),
              (S_SS, "Search service requirements: high availability, zones, replicas, overheads"),
              (S_SD, "Search demand per environment: data, documents, indexes, growth"),
              (S_KS, "Kafka calculation, step by step, with checks and explanation"),
              (S_SZ, "Search calculation, step by step, with checks and explanation"),
              (S_CR, "Request to send to Confluent Cloud"),
              (S_OR, "Request to send to the managed OpenSearch provider"),
              (S_CK, "Summary of all checks"),
              (S_REF, "Published figures used by the formulas, with quoted evidence"),
              (S_TAB, "Published tables used by the formulas: Confluent limits, Pega search and partition tables"),
              (S_REFS, "Source documents with links"),
              (S_LOG, "Version history")]
    for s, text in sheets:
        c = ws.cell(r, 2, s)
        c.hyperlink = f"#'{s}'!A1"
        c.font = F_LINK
        ws.cell(r, 3, text).font = F_BASE
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
        r += 1
    r += 1
    disc = ("Basis and limits. Pega states that its sizing tables are a starting point and that the environment should "
            "be scaled using indexing, query and stream metrics after deployment [P2]. Confluent and the search "
            "provider make the final sizing and price. Where Pega publishes no figure (connections, request rates, "
            "storage formula), the workbook uses measured inputs or the cited provider guidance, and labels workbook "
            "rules W. Figures were checked against the sources on " + R.CHECKED + "; re-check them before a contract.")
    ws.cell(r, 2, disc).alignment = WRAP
    ws.cell(r, 2).font = Font(name="Calibri", size=10, italic=True, color="404040")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 70
    protect(ws)


GUIDE_STEPS = [
    ("1", "Setup", "Enter the programme details. List the environments, choose each one's Pega landscape and assign "
                   "it to a Kafka cluster (K1 to K6) and a search service (S1 to S6)."),
    ("2", "Kafka Clusters", "For each cluster in use, choose the SLA, private networking, self-managed keys, preferred "
                            "type, growth allowance and target utilisation."),
    ("3", "Kafka Demand", "For each environment, enter partitions (measured, or queue processors and topics), peak "
                          "throughput, daily volume, retention, connections, requests and largest message. Set the "
                          "input status to Measured or Estimated."),
    ("4", "Search Services", "For each service in use, choose high availability, zones and replicas, and enter the "
                             "provider's storage overhead when known."),
    ("5", "Search Demand", "For each environment, enter searchable data, documents, indexes, growth and planning "
                           "period."),
    ("6", "Checks", "Resolve every FAIL and review every WARN. Read the explanation at the foot of the two sizing "
                    "sheets."),
    ("7", "Request sheets", "Send the Confluent Request sheet to Confluent and the Search Request sheet to the "
                            "OpenSearch provider, with the questions on each sheet. Record their answers in the Notes "
                            "rows and update the inputs."),
    ("8", "After go-live in DEV", "Replace estimates with measured figures from the first DEV build on the new services "
                                  "and recalculate before sizing the higher environments."),
]

DATA_SOURCES = [
    ("Leader partitions", "Stream service landing page in Pega 8.8 (partitions per topic) [P5]",
     "Confluent topic list for the DEV cluster after the first Pega '26 start-up",
     "Queue processor topics have six partitions by default [P3]."),
    ("Queue processors and topics", "Admin Studio, Queue processors; Dev Studio, Stream data sets",
     "Confluent topic list", "Include the system pulse topic (6 partitions) [P6] and the messaging topics that "
                             "replace Hazelcast from Pega '25 [P8]; these do not exist on 8.8."),
    ("Peak messages and size", "Stream service monitoring in Pega 8.8, or application estimates",
     "Confluent Cloud cluster metrics: received records and received bytes at peak",
     "Use the busiest hour of a normal business day plus any batch window."),
    ("Peak read and fan-out", "Number of consumers per topic", "Confluent Cloud cluster metrics: sent bytes at peak",
     "Egress counts every read by every consumer group."),
    ("Messages or GiB per day", "Application estimates", "Confluent Cloud cluster metrics: received bytes per day",
     "Drives ingress, egress and storage pricing [C2]."),
    ("Retention", "Design decision", "Topic retention.ms", "Confluent topic default is 7 days [C3]."),
    ("Client connections and requests", "Not available from embedded Kafka in a comparable form",
     "Confluent Cloud cluster metrics for the DEV cluster under test load",
     "Pega publishes no figure. Leave blank until measured; the checks show it as not assessed."),
    ("Largest message", "Application design", "Producer errors in Pega logs if the limit is reached",
     "Pega default limit 5,000,000 bytes [P1]."),
    ("Searchable data and documents", "Current 8.8 search index size and document counts",
     "GET _cat/indices?v (pri.store.size and docs.count) after the first SRS build [A1]",
     "Enter primary size only; the workbook adds replicas."),
    ("Indexes and primary shards", "Current search indexes", "GET _cat/indices?v after the first SRS build",
     "OpenSearch default is one primary shard per index [O2]."),
    ("Growth and planning period", "Business plans", "-", "Applied to both services."),
]


def build_guide(wb):
    ws = wb.create_sheet(S_GUIDE, 1)
    title(ws, "Guide: how to use this workbook",
          "Work through the numbered sheets in order. Hover over cells with a red corner for guidance and sources.")
    widths(ws, [6, 26, 46, 46, 50])
    header_row(ws, 4, ["Step", "Sheet", "What to do", "", ""])
    ws.merge_cells("C4:E4")
    r = 5
    for n, s, text in GUIDE_STEPS:
        style(ws.cell(r, 1, n), "white", align=CENTER)
        style(ws.cell(r, 2, s), "white", bold=True)
        style(ws.cell(r, 3, text), "white")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        fit_height(ws, r, [(text, 130)])
        r += 1
    r += 1
    ws.cell(r, 1, "Where each input comes from").font = Font(name="Calibri", size=12, bold=True, color=NAVY)
    r += 1
    header_row(ws, r, ["", "Input", "Pega 8.8 today (embedded services)", "Pega '26 DEV on the new services", "Notes"])
    r += 1
    for row in DATA_SOURCES:
        style(ws.cell(r, 1, ""), "white")
        for j, v in enumerate(row):
            style(ws.cell(r, 2 + j, v), "white", bold=(j == 0))
        fit_height(ws, r, [(row[1], 46), (row[2], 46), (row[3], 50)])
        r += 1
    r += 1
    ws.cell(r, 1, "How the workbook sizes the services").font = Font(name="Calibri", size=12, bold=True, color=NAVY)
    r += 1
    method = [
        ("Kafka", "Demand from all environments on a cluster is added up. Each figure is raised by the growth "
                  "allowance and divided by the target utilisation. For each Confluent type, the units needed are the "
                  "largest of planned demand divided by the per-unit limit for ingress, egress, partitions, compacted "
                  "partitions, connections and requests, and the minimum units for the SLA [C1]. The first type that "
                  "meets the SLA, networking and key requirements and fits within its maximum is recommended, unless "
                  "the customer forces a type. Pega's Table 1 and partitions-per-broker table give a self-managed "
                  "reference for comparison [P1][P5]."),
        ("Search", "Pega's Default or Large profile gives the starting node sizes and counts [P2]. Storage is "
                   "worked out with the AWS formula (replicas, indexing overhead, OS reserve and provider overhead) "
                   "[A1]. The data-node count is the largest of: the Pega profile with the three-node HA rule, "
                   "storage / disk per node, shards / 1,000 [O3], shards / (25 x heap) [A2], and replicas + 1 [O5]. "
                   "With high availability it is rounded up to a multiple of the zone count [O1]."),
        ("Pricing", "Confluent bills capacity units per hour, ingress and egress per GiB and storage per GiB-hour "
                    "after replication [C2]. The workbook gives these quantities per month for Confluent to price. "
                    "The search provider prices the nodes on the Search Request sheet."),
    ]
    for k, text in method:
        style(ws.cell(r, 2, k), "white", bold=True)
        style(ws.cell(r, 3, text), "white")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        fit_height(ws, r, [(text, 130)])
        r += 1
    protect(ws)


def build_log(wb):
    ws = wb.create_sheet(S_LOG)
    title(ws, "Change log", "Record every change to inputs, reference data or formulas.")
    widths(ws, [10, 16, 70, 24])
    header_row(ws, 4, ["Version", "Date", "Change", "By"])
    rows = [(VERSION, DATE, "First issue. Kafka and OpenSearch sizing for Pega Platform '26 on Azure AKS, using Pega, "
                            "Confluent, OpenSearch and AWS sources checked on " + R.CHECKED + ".", "Implementation team")]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            style(ws.cell(5 + i, 1 + j, v), "white")
        fit_height(ws, 5 + i, [(row[2], 70)])
    for i in range(1, 15):
        for j in range(4):
            style(ws.cell(5 + i, 1 + j), "input")


# ---------------------------------------------------------------------------------------------- build

def build(example):
    wb = Workbook()
    wb.remove(wb.active)
    build_setup(wb, example)
    ws_kc, kc = vertical_inputs(wb, S_KC, "2  Kafka clusters: requirements for each Confluent Cloud cluster",
                                "One column per cluster. Only clusters with environments assigned on the Setup sheet "
                                "are sized.", KIDS, KC_ITEMS, KC_EXAMPLE if example else KC_DEFAULT, "KC.")
    build_demand(wb, S_KD, "3  Kafka demand per environment",
                 "One row per environment from the Setup sheet. Enter measured figures where available.",
                 KD_COLS, "KD", kd_example() if example else None,
                 [("env", "rank", "Environment"), ("status", "parts", "Partitions"),
                  ("compact", "out", "Peak throughput"), ("day", "store", "Daily volume and storage"),
                  ("conn", "notes", "Connections, requests and message size")])
    kd_letters = {k: get_column_letter(i + 1) for i, (k, *_r) in enumerate(KD_COLS)}
    ws_ss, ss = vertical_inputs(wb, S_SS, "4  Search services: requirements for each managed OpenSearch service",
                                "One column per service. Only services with environments assigned on the Setup sheet "
                                "are sized.", SIDS, SS_ITEMS, SS_EXAMPLE if example else SS_DEFAULT, "SS.")
    build_demand(wb, S_SD, "5  Search demand per environment",
                 "One row per environment from the Setup sheet. Enter primary (not replicated) sizes.",
                 SD_COLS, "SD", sd_example() if example else None,
                 [("env", "rank", "Environment"), ("status", "tmp", "Inputs"), ("fgb", "notes", "Worked out")])
    sd_letters = {k: get_column_letter(i + 1) for i, (k, *_r) in enumerate(SD_COLS)}
    _, ks = build_kafka_sizing(wb, kc, kd_letters)
    _, sz = build_search_sizing(wb, ss, sd_letters)
    build_confluent_request(wb, ks, kc)
    build_search_request(wb, sz, ss)
    build_checks(wb, ks, sz)
    build_ref(wb)
    build_ref_tables(wb)
    build_references(wb)
    build_log(wb)
    build_cover(wb, ks, sz, example)
    build_guide(wb)
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = {S_COVER: NAVY, S_GUIDE: NAVY, S_SETUP: "BF8F00", S_KC: "BF8F00",
                                        S_KD: "BF8F00", S_SS: "BF8F00", S_SD: "BF8F00", S_KS: "548235",
                                        S_SZ: "548235", S_CR: "2F5597", S_OR: "2F5597", S_CK: "C00000"}.get(ws.title, "7F7F7F")
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 2 if ws.title == S_KD else 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.oddFooter.center.text = "&A  -  page &P of &N"
        ws.oddFooter.left.text = f"Pega '26 Kafka and OpenSearch sizing workbook v{VERSION}"
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    return wb, ks, sz, kc, ss


def main():
    OUT.mkdir(exist_ok=True)
    paths = []
    for example, name in [(False, "Pega_26_Kafka_OpenSearch_Sizing_Template.xlsx"),
                          (True, "Pega_26_Kafka_OpenSearch_Sizing_Worked_Example.xlsx")]:
        wb, *_ = build(example)
        p = OUT / name
        wb.save(p)
        paths.append(p)
        print("wrote", p)
    return paths


if __name__ == "__main__":
    sys.exit(0 if main() else 1)
