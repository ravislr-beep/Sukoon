"""Front matter and Sections 1 to 4."""
from docx.shared import Pt, Cm

from docx_lib import NAVY, GREY, _fix_grid
import envs as E

GEN = "img/generated/"
PUB = "img/public/"
OWN = "Prepared for this implementation. Not a Pega-published diagram."


def cover(b, meta):
    d = b.doc
    for _ in range(5):
        d.add_paragraph()
    p = d.add_paragraph()
    r = p.add_run("CUSTOMER ENGINEERING RUNBOOK")
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = GREY
    p = d.add_paragraph()
    r = p.add_run("Pega Platform 26.1.1 on Azure Kubernetes Service")
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = NAVY
    p = d.add_paragraph()
    r = p.add_run("Externalized Kafka on Confluent Cloud and search on managed OpenSearch")
    r.font.size = Pt(15)
    r.font.color.rgb = NAVY
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("Design, Helm configuration, testing and operations for six environments built from a cloned "
                  "and upgraded Pega Platform 8.8 database")
    r.font.size = Pt(12)
    r.font.color.rgb = GREY
    for _ in range(6):
        d.add_paragraph()
    rows = [
        ("Document version", meta["version"]),
        ("Status", meta["status"]),
        ("Issue date", meta["date"]),
        ("Classification", "Customer Confidential"),
        ("Target", "Pega Platform 26.1.1 on Azure AKS, Pega Helm charts 4.13.0"),
        ("Source", "Pega Platform 8.8 with embedded Kafka and embedded Elasticsearch"),
        ("Environments", "DEV, SIT, UAT, PERF, PREPROD, PROD"),
    ]
    t = d.add_table(rows=0, cols=2)
    for k, v in rows:
        c = t.add_row().cells
        c[0].width, c[1].width = Cm(4.5), Cm(11.5)
        r = c[0].paragraphs[0].add_run(k)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = NAVY
        r = c[1].paragraphs[0].add_run(v)
        r.font.size = Pt(10)
    _fix_grid(t, [4.5, 11.5])


def document_control(b, meta):
    p = b.doc.add_paragraph(style="Front Heading")
    p.paragraph_format.page_break_before = True
    p.add_run("Document control")
    b.table(["Item", "Detail"], [
        ["Title", "Pega Platform 26.1.1 on Azure AKS: externalized Kafka on Confluent Cloud and search on managed OpenSearch"],
        ["Version", meta["version"]],
        ["Status", meta["status"]],
        ["Issue date", meta["date"]],
        ["Evidence cut-off", "Vendor pages were read between 28 September and 2 October 2026. Section 3 lists the facts used. Appendix F lists every source with its link."],
        ["Supersedes", "Version 2.1 of this runbook, issued earlier on 2 October 2026. Version 2.1 corrected version 2.0, issued on 1 October 2026. Version 2.0 itself replaced version 1.0, which assumed a three-release path through an intermediate Pega release before the customer confirmed the clone-and-upgrade path (Section 1.2)."],
        ["Owner", "Customer platform engineering lead"],
        ["Review cycle", "Before each rehearsal, before the production cutover, and after any change to a vendor page listed in Section 3"],
    ], widths=[3.5, 13], first_col_bold=True, size=9.5)
    p = b.doc.add_paragraph(style="Front Heading")
    p.paragraph_format.space_before = Pt(14)
    p.add_run("Revision history")
    b.table(["Version", "Date", "Change", "Author"], [
        ["1.0", "1 October 2026", "Kafka and search runbook for a three-release upgrade path.", "Implementation team"],
        ["2.0", "1 October 2026", "Rewritten for the clone-and-upgrade path, six environments, shared non-production services, managed OpenSearch and Okta.", "Implementation team"],
        ["2.1", meta["date"], "Expert review: Helm values corrected after rendering with `helm template` (SRS authentication, SRS port 8443, SRS network policy, SRS replicas and resources); availability zone design; Confluent connection, request and partition-creation limits; Okta key rotation and rate limits; OpenSearch sizing formulas; JVM settings; new Sections 15 (performance engineering and sizing) and 18 (observability); environment build waves; Appendix H sizing calculator. Sections 15 to 18 of version 2.0 became Sections 16, 17, 19 and 20.", "Implementation team"],
        ["2.2", meta["date"], "Data review: new Section 9 on the data Pega exchanges with Kafka, SRS, OpenSearch and Okta, its sensitivity, encryption in transit and at rest, personal data controls, residency, retention and erasure, and what happens to data when a transfer fails. New decisions OD-14 to OD-18, Pega Support questions GQ-09 to GQ-11, data tests DT-01 to DT-12, risks RK-17 to RK-21. Sections 9 to 20 of version 2.1 are now Sections 10 to 21.", "Implementation team"],
    ], widths=[2, 2.8, 8.8, 3])
    p = b.doc.add_paragraph(style="Front Heading")
    p.paragraph_format.space_before = Pt(14)
    p.add_run("Approvals")
    b.p("Approval confirms that the approver has reviewed the sections relevant to the role and accepts the design, the "
        "recommendations and the open decisions as stated.")
    b.table(["Role", "Name", "Signature", "Date"], [
        ["Customer business owner", "", "", ""],
        ["Customer enterprise architect", "", "", ""],
        ["Customer security architect", "", "", ""],
        ["Customer platform engineering lead", "", "", ""],
        ["Customer database administration lead", "", "", ""],
        ["Pega Lead System Architect", "", "", ""],
        ["Kafka and search engineering lead", "", "", ""],
    ], widths=[5.5, 4.5, 4, 2.6], zebra=False)


def s1_exec(b):
    b.h1("Executive summary")
    b.h2("Purpose")
    b.p("This runbook tells the customer's IT teams how to design, configure, test and run Kafka and search for Pega "
        "Platform 26.1.1 on Azure Kubernetes Service (AKS). Kafka is provided by Confluent Cloud. Search is provided by the "
        "Pega Search and Reporting Service (SRS) on top of a managed OpenSearch service. Both are configured through the Pega "
        "Helm charts in six environments: DEV, SIT, UAT, PERF, PREPROD and PROD.")
    b.h2("The clone-and-upgrade approach")
    b.p("The customer is moving from Pega Platform 8.8 without upgrading the running 8.8 servers. A new 26.1.1 platform is "
        "built on AKS with Kafka and search outside Pega from the start. The 8.8 database is cloned, and the clone is upgraded "
        "to 26.1.1 by the Pega installer, run as the Helm chart installer job. Pega 26.1.1 then starts on the upgraded clone. "
        "Each lower environment is built the same way, which gives the programme repeated rehearsals. Production uses a final "
        "clone taken after 8.8 is stopped. The 8.8 system and its database stay untouched until the rollback window closes. "
        "{ref:fig_clone} shows the flow.")
    b.figure(GEN + "fig_clone_upgrade_flow.png", "Clone-and-upgrade flow for each environment", OWN, width_cm=8.2, label="clone")
    b.h2("Kafka and search design in brief")
    b.p("{ref:tab_design_brief} summarises the design. Production has its own Kafka cluster and its own OpenSearch service. "
        "The five non-production environments share two sets of services, with isolation enforced per environment.")
    b.table(["Layer", "PROD", "PERF and PREPROD (group NP2)", "DEV, SIT and UAT (group NP1)"], [
        ["Kafka", "Confluent Cloud cluster cc-prd, Enterprise or Dedicated, Azure Private Link", "Shared cluster cc-np2, same type and settings as PROD", "Shared cluster cc-np1, Enterprise with Private Link"],
        ["Isolation in Kafka", "Dedicated cluster", "Topic prefix, service account, prefixed ACLs and client quota per environment", "Same as NP2"],
        ["Search storage", "OpenSearch service os-prd", "Shared service os-np2, same provider and settings as PROD", "Shared service os-np1"],
        ["Isolation in search", "Dedicated service", "SRS per environment, customerDeploymentId per environment, OpenSearch user and index-scoped role per environment", "Same as NP2"],
        ["Pega-to-SRS tokens", "Okta custom authorization server for PROD only", "Shared non-production authorization server, one client per environment", "Same as NP2"],
        ["Secrets", "Key Vault per environment, delivered by the External Secrets Operator", "Same", "Same"],
    ], caption="Kafka and search design by environment group", widths=[2.6, 4.6, 4.7, 4.7], size=8.5, first_col_bold=True, label="design_brief")
    b.h2("Sharing services to reduce cost")
    b.p("Six dedicated Kafka clusters and six OpenSearch services would cost the most and need the most operation. The "
        "recommended model uses three sets of services instead of six. DEV, SIT and UAT share one set. PERF and PREPROD share a "
        "second set that matches production, so load tests and production rehearsals run on production-like services "
        "without affecting functional testing. Section 5 compares this model with the alternatives and gives a cost model "
        "that the customer completes with its own prices.")
    b.h2("Answer to the migration question")
    b.callout("decision", ["No Pega stream topic and no search index is copied from 8.8. The 8.8 system uses embedded Kafka "
              "and embedded Elasticsearch, whose data lives on the 8.8 nodes, not in the database. Pega states that existing "
              "stream data cannot be moved to a new Kafka [R13], and SRS indexes all searchable data after Pega connects to it "
              "[R15].",
              "Before the production clone is taken, intake on 8.8 is held and queues are allowed to empty. Queue items that are "
              "held in the database, such as delayed and broken items, travel with the clone and are handled as listed in "
              "Section 11.3. Application Kafka data sets on the customer's own Kafka clusters are decided one by one (Section 12.3)."])
    b.h2("Questions for Pega Support")
    b.p("Pega documentation does not settle three points that the plan depends on. Section 11.1 lists eleven questions in "
        "full. The three that gate the production plan are these:")
    b.bullets([
        "GQ-01: Pega Support must confirm that a direct upgrade of an 8.8 database to 26.1.1 is supported, and name any required 8.8 patch level.",
        "GQ-03: Pega states that Hazelcast must be removed before updating to '25 or later [R5]. Pega Support must confirm how this applies to a cloned database that is upgraded offline and never started on 8.8 again.",
        "GQ-04: Pega Support must confirm the installer upgrade type for a cloned database (`in-place` or the out-of-place types the Helm chart documents [R23]).",
    ])
    b.h2("Decisions needed")
    b.p("{ref:tab_exec_decisions} lists the decisions that block the first build. Section 21.2 lists all open decisions with "
        "owners and milestones.")
    b.table(["ID", "Decision", "Recommended answer", "Needed by"], [
        ["OD-01", "Pega Support answers to the gating questions", "Raise the case now; plan with the conservative answer until the reply arrives", "M2 DEV build"],
        ["OD-02", "Confluent cluster types", "Enterprise with Private Link for NP1; PROD type for NP2 and PROD after partition measurement", "M2 DEV build"],
        ["OD-03", "OpenSearch provider", "Provider that passes the selection in Section 7.3", "M2 DEV build"],
        ["OD-05", "Okta design", "Dedicated custom authorization server for PROD; one shared non-production server with a per-client guid claim, if proven in DEV", "M2 DEV build"],
        ["OD-07", "Masking approach", "Mask before the upgrade, so the index build uses masked data only", "M2 DEV build"],
    ], caption="Decisions needed before the first build", widths=[1.4, 4.6, 8, 2.6], size=9, label="exec_decisions")
    b.h2("Readiness statement")
    b.p("This document is ready for customer review as a design and runbook. Durations for draining, cloning, upgrading and "
        "index building depend on data volume, so they come from the rehearsals in Section 15.2, not from estimates. Points that "
        "no source settles are marked as tests, Pega Support questions or open decisions. Conflicts between sources are logged "
        "in Section 21.4.")


def s2_scope(b):
    b.h1("Scope, dependencies and reading guide")
    b.h2("Scope")
    b.table(["In scope", "Out of scope"], [
        [["Confluent Cloud and managed OpenSearch design for DEV, SIT, UAT, PERF, PREPROD and PROD.",
          "Configuration in the Pega Helm charts: the `stream` and `pegasearch` sections of the `pega` chart, and the SRS section of the `backingservices` chart.",
          "What the cloned 8.8 database brings that affects Kafka and search, and how it is cleaned up before and after the first 26.1.1 start.",
          "Whether Kafka topics and search indexes are migrated from 8.8, and how 8.8 in-flight work is handled at cutover.",
          "Shared non-production services with isolation between environments, and the cost model.",
          "Okta as the token issuer for Pega-to-SRS authorization.",
          "The data Pega exchanges with Kafka, SRS, OpenSearch and Okta: its sensitivity, protection in transit and at rest, retention, erasure and handling of failed transfers.",
          "Testing, failure scenarios, known issues, troubleshooting and operations for Kafka and search."],
         ["Application functional regression content.",
          "Cassandra and Decision Data Store.",
          "Constellation.",
          "Pega Diagnostic Center beyond the network path it needs.",
          "General AKS build-out not related to Kafka or search.",
          "Commercial terms and prices from Confluent and the OpenSearch provider."]],
    ], caption="Scope of this document", widths=[9.3, 7.3], size=9.5, zebra=False)
    b.h2("Dependencies covered only to the depth Kafka and search need")
    b.p("The database clone, the installer upgrade and Pega's '25 and '26 prerequisites belong to other workstreams. "
        "{ref:tab_deps} names the owner of each, the gate it must pass before Kafka and search work can rely on it, and the source.")
    b.table(["Dependency", "Why it matters for Kafka and search", "Owner", "Gate", "Source"], [
        ["Database clone and masking", "The clone carries 8.8 stream and search settings and queue items (Section 11.3). Masking decides what data the index build uses.", "DBA team", "Clone restored and masked; row counts recorded", "Customer policy"],
        ["Installer upgrade (`installer.upgrade.upgradeType`)", "The upgrade must finish before Pega starts against Kafka and SRS.", "Platform team with Pega LSA", "Installer job completed; log reviewed", "[R23]"],
        ["Hazelcast removal", "From '25, cluster messaging that Hazelcast carried uses Kafka, so a Kafka outage affects more than queue processing.", "Pega LSA", "GQ-03 answered", "[R4, R5, R7]"],
        ["Tomcat 10.1 and Jakarta EE", "Custom JARs that fail to load can stop nodes before Kafka and search checks run.", "Application team", "Custom JARs rebuilt and tested", "[R3, R9]"],
        ["Primary keys on all tables", "Required before an update to '25 or later.", "DBA team", "primaryKeyUtility dry run clean", "[R3, R10]"],
        ["Queue processor changes in '26", "Partitioned, multi-threaded processing; DSS `delayeditems/dataflowbased/threadspernode` must be removed (PEGA0179).", "Pega LSA", "DSS removed on the clone", "[R3]"],
    ], caption="Dependencies and their gates", widths=[3.2, 5.6, 2.4, 3.6, 1.8], size=8.5, label="deps")
    b.h2("Confirmed customer inputs")
    b.table(["Input", "Value", "Effect on this document"], [
        ["Environments", "DEV, SIT, UAT, PERF, PREPROD, PROD", "Six columns in every per-environment table; six prefixes, service accounts, deployment IDs, Okta clients and Key Vaults."],
        ["8.8 Kafka", "Embedded Kafka (Pega stream service inside 8.8)", "Nothing in it can be moved to Confluent Cloud. Queues are emptied on 8.8 before the final clone (Section 12)."],
        ["8.8 search", "Embedded Elasticsearch", "Indexes are on the 8.8 nodes' disks and cannot be carried over. 26.1.1 builds all indexes through SRS (Section 7.8)."],
        ["Token issuer for Pega to SRS", "Okta", "Okta design in Section 7.6."],
        ["OpenSearch provider", "Not chosen yet", "Selection method in Section 7.3; Helm values written so only the endpoint, port, credentials and CA depend on the provider."],
        ["Pega Support answers", "Not received", "Every gating question is open with an owner and a milestone (Section 11.1)."],
    ], caption="Confirmed customer inputs", widths=[3.4, 4.6, 8.6], size=9)
    b.h2("Where each question is answered")
    b.p("The customer asked nineteen questions. {ref:tab_trace} shows where each one is answered, so a reviewer can check "
        "that none is missed.")
    b.table(["#", "Question", "Answered in"], [
        ["1", "Is the direct clone-and-upgrade path from 8.8 to 26.1.1 supported, and on what conditions?", "Section 11.1"],
        ["2", "How is Confluent Cloud designed per environment?", "Sections 5 and 6"],
        ["3", "How is the managed OpenSearch service designed per environment?", "Sections 5 and 7"],
        ["4", "How is each one configured in the Helm charts for the upgrade run and the deploy run?", "Section 10 and Appendix B"],
        ["5", "What in the cloned 8.8 database conflicts with the new configuration, and how is it cleaned up?", "Section 11.3"],
        ["6", "How is a cloned non-production environment kept away from production Kafka, integrations and other environments?", "Sections 5.3, 11.4 and 15.3"],
        ["7", "Must topics or indexes be migrated from 8.8, and how is in-flight work handled at cutover?", "Sections 12 and 14"],
        ["8", "How do several non-production environments share one cluster and one OpenSearch service safely?", "Sections 5.3 and 5.4"],
        ["9", "How much does sharing save, and what risk does it add?", "Section 5.5"],
        ["10", "How is it tested?", "Section 15"],
        ["11", "Which failures are tested, and how?", "Section 15.6"],
        ["12", "What known issues should be avoided?", "Section 17"],
        ["13", "How is it run day to day, including refresh and retirement?", "Section 20"],
        ["14", "Why is performance testing needed, and how is it done?", "Section 16"],
        ["15", "How is each layer sized, and is there a calculator?", "Sections 16.7 to 16.9 and Appendix H"],
        ["16", "How is the platform monitored, logged and observed?", "Section 19"],
        ["17", "In which order are the environments built, and what is configured in each one?", "Section 13.10"],
        ["18", "What happens when an availability zone fails?", "Section 4.6"],
        ["19", "What data goes to Kafka, SRS and OpenSearch, how sensitive is it, how is it protected in transit and at rest, and what happens when a transfer fails?", "Section 9"],
    ], caption="Traceability from customer questions to sections", widths=[0.8, 11.6, 4.2], size=9, label="trace")
    b.h2("Reading guide by role")
    b.table(["Reader", "Read first", "Then"], [
        ["Business owner, programme manager", "Sections 1, 5.5, 14 and 21", "Sections 15.2 and 16.1"],
        ["Enterprise and security architects", "Sections 4, 5, 7.6, 8 and 9", "Sections 19 and 21"],
        ["Pega Lead System Architect", "Sections 11 and 12", "Sections 14 to 17"],
        ["Database administrators", "Sections 9.1, 11.1, 11.3 to 11.5 and 13.7", "Section 14"],
        ["Platform, Kafka and search engineers", "Sections 6 to 10 and 13", "Sections 16, 18, 19 and 20, Appendices A to C and H"],
        ["Test lead, performance lead", "Sections 15 and 16", "Sections 9.8 and 17"],
        ["Operations and monitoring teams", "Section 19", "Sections 18 and 20"],
    ], caption="Reading guide by role", widths=[5, 6.5, 5.1], size=9.5)
    b.h2("Conventions and naming")
    b.bullets([
        "**Must** marks a vendor requirement or a fixed decision. **Should** marks a recommendation that the customer can change through the decision process in Section 21.2.",
        "Values in angle brackets, such as `<bootstrap-host>`, are environment-specific. Each one is listed in Appendix A.",
        "References in square brackets, such as [R11], point to the numbered sources in Appendix F.",
        "Figures marked \"Prepared for this implementation\" were drawn for this document. Figures from Pega, Pega Academy, Confluent and OpenSearch carry a source line, and Appendix G lists them.",
    ])
    b.p("Every environment uses the same naming pattern, shown in {ref:tab_naming}. Prefixes end with a hyphen, and no "
        "deployment ID is the start of another, so a prefixed ACL or an index pattern for one environment can never match "
        "another environment's topics or indexes.")
    b.table(["Environment", "Code", "Group", "Topic prefix", "customerDeploymentId", "Service account", "Okta client"],
            E.per_env(lambda n, c, g: [n, f"`{c}`", g, f"`{E.prefix(c)}`", f"`{E.deployment_id(c)}`", f"`sa-pega-{c}`", f"`pega-srs-{c}`"]),
            caption="Environment naming", widths=[2.2, 1.3, 1.3, 2.6, 3.4, 2.9, 2.9], size=8.5, first_col_bold=True, label="naming")


def s3_evidence(b):
    b.h1("Evidence baseline")
    b.h2("How the evidence was gathered")
    b.p("Every product statement in this document comes from one of these sources, in this order of authority:")
    b.steps([
        "Pega documentation for Pega Platform '26, and for '25 where the topic belongs to that release.",
        "The Pega Helm charts repository on GitHub, chart version 4.13.0, for configuration keys and defaults.",
        "Pega Academy for architecture illustrations.",
        "Confluent Cloud documentation.",
        "OpenSearch documentation, and the chosen provider's documentation once OD-03 is decided.",
        "Okta developer documentation.",
        "Microsoft Learn and the External Secrets Operator documentation.",
    ])
    b.p("Where two sources disagree, the higher one is followed and the conflict is logged in Section 21.4. Where no source "
        "settles a point, the document says so and turns it into a test, a Pega Support question or an open decision.")
    b.h2("Facts this design depends on")
    b.p("{ref:tab_facts} records the facts as published on the evidence cut-off date. Re-check them before each rehearsal, "
        "because vendors update these pages.")
    b.table(["Topic", "Fact", "Source"], [
        ["Externalized services", "From Pega Platform '24.2, embedded Kafka, Elasticsearch and Cassandra are not supported. Kafka and search must be external.", "[R2]"],
        ["Kubernetes", "AKS is supported. x86-64 nodes only. Database in the same region as the worker nodes. Helm 3.0 or later.", "[R20, R27]"],
        ["Kafka role", "\"All Queue Processors and Job Schedulers in Pega Platform require Kafka.\" Without it, Pega Platform is not fully functional.", "[R11]"],
        ["Kafka client", "Pega '26 ships the Kafka client library 4.0.0, backward compatible with 3.9.2.", "[R3]"],
        ["Kafka settings", "message.max.bytes 5000000; replica.fetch.max.bytes 5000100; unclean.leader.election.enable false; auto.create.topics.enable false.", "[R11]"],
        ["Kafka security", "Generic Kafka security only; vendor-specific features such as IAM roles are not supported.", "[R11]"],
        ["Kafka ACLs", "TOPIC and GROUP on the stream name prefix (ALL, PREFIXED); TRANSACTIONAL_ID * (READ and WRITE, LITERAL); CLUSTER IDEMPOTENT_WRITE.", "[R11]"],
        ["Stream data", "Existing stream data cannot be moved to a new Kafka provider.", "[R13]"],
        ["Helm stream keys", "`stream.saslMechanism` accepts PLAIN, SCRAM-SHA-256 and SCRAM-SHA-512. External secret keys: STREAM_TRUSTSTORE_PASSWORD, STREAM_KEYSTORE_PASSWORD, STREAM_JAAS_CONFIG.", "[R23]"],
        ["SRS role", "SRS replaces embedded search and the legacy plug-in, which are not available from '24.2.", "[R14]"],
        ["SRS matrix", "SRS 1.44.3 or later. Image `search-n-reporting-service-os` supports OpenSearch 1.3, 2.15 and 2.19; best practice OpenSearch 2.15.", "[R14]"],
        ["SRS sharing", "SRS can support more than one Pega deployment; it isolates data with a unique CUSTOMER_DEPLOYMENT_ID. Do not share the search service with non-Pega software.", "[R14]"],
        ["Deployment ID", "The customerDeploymentId is the prefix of all indexes and must equal the `guid` claim when OAuth is used. It defaults to the namespace name.", "[R23]"],
        ["Index build", "After Pega connects to SRS, SRS indexes all searchable data, which requires a downtime period.", "[R15]"],
        ["Search settings", "Index auto-creation must be off (SRS sets it if its user has the manage cluster privilege). destructive_requires_name false.", "[R14]"],
        ["SRS tokens", "OAuth client credentials with private_key_jwt or client_secret_basic; scope pega.search:full; private key in SRS_OAUTH_PRIVATE_KEY.", "[R23]"],
        ["Hazelcast", "Removed in '25. \"Before updating to Pega Platform '25, remove Hazelcast from your deployment.\" Affects embedded Hazelcast and the Clustering Service.", "[R5]"],
        ["Queue processing in '26", "Partitioned, multi-threaded queue processing. Remove DSS delayeditems/dataflowbased/threadspernode (PEGA0179). SystemPulse topic uses 6 partitions.", "[R3]"],
        ["Installer actions", "`install`, `deploy`, `install-deploy`, `upgrade`, `upgrade-deploy`. Upgrade types include `in-place`, `out-of-place-rules`, `out-of-place-data`, `zero-downtime`, `custom`.", "[R23]"],
        ["Confluent quotas", "Client quotas are supported on Enterprise, Freight and Dedicated clusters, applied per service account or identity pool, not per API key.", "[R36]"],
        ["Okta", "Custom scopes and custom claims need a custom authorization server. The org authorization server cannot be customized.", "[R42, R43, R44]"],
        ["Okta keys and limits", "Authorization server signing keys rotate about four times a year; clients look keys up by `kid`. Rate limits apply per org; excess requests receive HTTP 429.", "[R52, R53]"],
        ["Confluent limits", "Per Enterprise eCKU: 18,000 connections, 500 connection attempts per second, 7,500 requests per second; partition creation and deletion 500 per five minutes per cluster; limits enforced from 2026. 99.99 % SLA from 2 eCKU.", "[R31]"],
        ["Pega partitions", "New stream topics get 6 partitions by default (DSS `pyTopicPartitionsCount`); queue processor threads beyond the partition count do no work.", "[R49, R50]"],
        ["Queue processor payload", "A queued message carries the producer's operator, access group and application with the payload. By default a page with a `pzInsKey` is queued as the key only; \"Queue current snapshot of page\" stores the entire page. Failed items move to the broken queue after MaxAttempts.", "[R64, R65]"],
        ["Search content in SRS mode", "Only Pega's predefined properties and the classes and properties listed in Relevant Records and Custom Search Properties are indexed, unless DSS `indexer/srs/indexAllFieldsForFTS` is set.", "[R67, R68]"],
        ["Encryption at rest", "Confluent encrypts all clusters at rest; self-managed keys on Enterprise and Dedicated must be chosen at cluster creation. AKS temp and ephemeral OS disks are encrypted only with encryption at host, set at node pool creation.", "[R73, R74, R78]"],
        ["SRS chart", "`srsStorage.tls` is certificate authentication and cannot be combined with basic authentication; minimum 2 SRS replicas, 3 as best practice.", "[R24]"],
    ], caption="Version and support facts used in this design", widths=[2.8, 12, 1.8], size=8.5, label="facts")


def s4_arch(b):
    b.h1("Architecture")
    b.h2("Pega reference architecture for '26")
    b.p("Pega's Kubernetes architecture for '26 has three layers: the Pega application namespace (web and batch deployments "
        "and an installer job), backing services that Pega provides (SRS and Constellation), and services the client provides "
        "(search, Kafka, Cassandra for Customer Decision Hub only, and the database) [R21]. Hazelcast does not appear, because "
        "'25 removed it.")
    b.figure(PUB + "pega_docs_pega-platform-architecture.png", "Pega Platform Kubernetes architecture for '26",
             "Source: Pega Documentation, \"Understanding the Pega Platform deployment architecture\" [R21]. © Pegasystems Inc. Reproduced with attribution. The database shown is an example.",
             width_cm=14)
    b.p("The 8.8 source system runs Kafka and Elasticsearch inside the Pega cluster. Pega Academy shows that older layout "
        "and the externalized layout side by side [R29, R30]. The comparison explains why nothing in the 8.8 Kafka or search "
        "storage can be moved: it lives on the 8.8 nodes.")
    b.figure(PUB + "pega_academy_three_node.png", "Earlier Pega cluster layout with embedded Kafka and in-cluster Elasticsearch",
             "Source: Pega Academy, \"Cloud deployment architecture\" [R30]. © Pegasystems Inc. Reproduced with attribution. Shown to explain the 8.8 source; this layout is not supported from '24.2.",
             width_cm=8.5)
    b.figure(PUB + "pega_academy_external_services.png", "Externalized services in a Pega deployment",
             "Source: Pega Academy, \"Deployment architecture with external services\" [R29]. © Pegasystems Inc. Reproduced with attribution.",
             width_cm=13)
    b.h2("Target architecture on Azure")
    b.p("{ref:fig_target} shows one environment. The Pega web and batch tiers, the installer job and SRS run in a private AKS "
        "cluster. Confluent Cloud, the OpenSearch service, the upgraded database, Key Vault and the container registry are "
        "reached through private endpoints. Users reach Pega through Application Gateway with Web Application Firewall. "
        "Okta is a public SaaS service, reached through the hub firewall on an allow-list.")
    b.figure(GEN + "fig_target_architecture.png", "Target architecture for one environment on Azure AKS", OWN, width_cm=16.5, label="target")
    b.table(["Component", "Role", "Design choice", "Source"], [
        ["Web tier", "User and API traffic (nodeType WebUser)", "HPA on CPU; pod disruption budget", "[R23]"],
        ["Batch tier", "Queue processors, job schedulers, search indexing, data flows", "No ingress; held at zero replicas for the first start of a clone (Section 11.4)", "[R21, R23]"],
        ["Installer job", "Upgrades the cloned database", "Runs with `global.actions.execute: upgrade` (Section 10.4)", "[R23]"],
        ["SRS", "Search and Reporting Service between Pega and OpenSearch", "One deployment per environment; network policy; OAuth enabled", "[R14, R24]"],
        ["External Secrets Operator", "Copies Key Vault secrets into Kubernetes secrets", "Workload identity to Key Vault", "[R23, R48]"],
        ["Confluent Cloud", "Kafka for the Pega stream service", "Private Link; replication factor 3; prefixed ACLs", "[R11, R31, R32]"],
        ["Managed OpenSearch", "Index storage for SRS", "Version on the SRS matrix; private endpoint", "[R14]"],
        ["Okta", "Issues tokens Pega presents to SRS", "Custom authorization server; one client per environment", "[R42, R43]"],
        ["Pega database", "Upgraded clone of the 8.8 database", "Same region as AKS; private endpoint; TLS", "[R20]"],
    ], caption="Target architecture components", widths=[3.2, 4.4, 6.8, 2.2], size=8.5)
    b.h2("Logical flows")
    b.p("{ref:fig_flows} shows the five flows that matter for Kafka and search. Flow 1 is the user request. Flow 2 is queue "
        "processing: producers on any tier write to Kafka, and the batch tier consumes. Flow 3 is search: the batch tier indexes "
        "changed records through SRS, and the web tier sends queries through SRS. Flow 4 is the token path to Okta. Flow 5 is "
        "telemetry.")
    b.figure(GEN + "fig_logical_flows.png", "Logical flows between Pega, Kafka, SRS, OpenSearch and Okta", OWN, width_cm=11, label="flows")
    b.h2("Network and DNS")
    b.p("Every private endpoint needs a DNS record that resolves to its private IP from inside AKS [R47]. Confluent Private "
        "Link needs the wildcard and zonal records that Confluent lists for the cluster's network [R32]. A missing zonal record "
        "is the most common reason a Kafka client reaches the bootstrap server and then fails on broker connections. "
        "{ref:fig_dns} shows the name resolution and traffic paths, and {ref:tab_flows} lists the flows the firewall and "
        "network policies must allow.")
    b.figure(GEN + "fig_network_dns.png", "Name resolution and traffic paths for Kafka, OpenSearch and Okta", OWN, width_cm=12.5, label="dns")
    b.table(["From", "To", "Port", "Protocol", "Purpose"], [
        ["Pega pods (web, batch)", "Confluent private endpoints", "9092", "SASL_SSL", "Stream service [R32]"],
        ["Pega pods", "SRS service in srs-<env>", "8443", "HTTPS with bearer token", "Indexing and search (Section 7.7)"],
        ["SRS pods", "OpenSearch private endpoint", "443 or the provider port", "HTTPS", "Index storage and queries"],
        ["Pega pods and installer job", "Database private endpoint", "Engine port", "TLS", "Rules and data"],
        ["Pega pods", "Okta `<okta-domain>`", "443", "HTTPS through the firewall", "Token request"],
        ["SRS pods", "Okta `<okta-domain>`", "443", "HTTPS through the firewall", "Key set download"],
        ["External Secrets Operator", "Key Vault private endpoint", "443", "HTTPS", "Secret synchronization"],
        ["Any pod in a non-production namespace", "Production Kafka, production OpenSearch, production integrations", "Any", "Any", "Denied by firewall and network policy (Section 5.3)"],
    ], caption="Network flows to allow and to deny", widths=[3.4, 3.8, 2.6, 3, 3.8], size=8.5, label="flows")
    b.h2("Message path inside Pega")
    b.p("Each queue processor writes to a Kafka topic named from the stream name pattern, for example "
        "`pega-sit-<stream name>`. The batch tier reads each topic through a consumer group under the same prefix, and one "
        "partition is read by one consumer at a time. Items that are not yet due wait in the database as delayed items. Items "
        "that fail after their retries become broken items, also held in the database. This split matters for the cloned "
        "database: the topics are empty in 26.1.1, but the database-held items are not (Section 11.3). {ref:fig_msg} shows the path.")
    b.figure(GEN + "fig_message_path.png", "Message path from producer to batch tier", OWN, width_cm=16, label="msg")
    b.callout("note", "Pega '26 moves queue processing to a partitioned, multi-threaded model [R3]. The number of partitions per "
              "topic therefore limits how many batch threads can work on one queue processor at once. Measure partitions after the "
              "first start (Section 6.7) rather than estimating them.")
    b.h2("Availability zone resilience")
    b.p("PERF, PREPROD and PROD are built so that the loss of one Azure availability zone does not stop Pega, Kafka or search. "
        "{ref:fig_zones} shows the layout. DEV, SIT and UAT may run in fewer zones to save cost, but their shared Kafka "
        "cluster and OpenSearch service follow the same rules because they carry three environments.")
    b.figure(GEN + "fig_zone_resilience.png", "Spreading each layer across three availability zones", OWN, width_cm=13, label="zones")
    b.table(["Layer", "Zone design", "Setting", "Source"], [
        ["AKS node pools", "Each user node pool spans zones 1, 2 and 3. The zones of a node pool cannot be changed after it is created, so this is decided before the first build", "`--zones 1 2 3` and `--enable-encryption-at-host` at node pool creation (Section 9.5)", "[R59, R78]"],
        ["Pega web and batch pods", "Spread evenly across zones; the scheduler still places pods if a zone is down", "`tier[*].topologySpreadConstraints` on `topology.kubernetes.io/zone`, `whenUnsatisfiable: ScheduleAnyway` (Appendix B)", "[R59, R61]"],
        ["SRS pods", "Preferred anti-affinity across zones", "`srsRuntime.affinity` (Section 7.7)", "[R24]"],
        ["Node pool capacity", "Two zones must hold the full peak load, so each pool is sized at 1.5 times the peak node count", "Sizing calculator (Appendix H)", "This design"],
        ["Pod disruption budgets", "Keep at least one pod per tier during drains and upgrades", "`tier[*].pdb.minAvailable: 1`", "[R23]"],
        ["Confluent Cloud", "Enterprise: the 99.99 % SLA needs at least 2 eCKU. Dedicated: the cluster must be created multi-zone, which needs at least 2 CKU and cannot be changed later", "Cluster creation (CC-1); one private endpoint per zone", "[R31, R32]"],
        ["OpenSearch", "Three cluster-manager nodes and data nodes in three zones, one replica, zone awareness on", "Provider setting (Section 7.4)", "[R37]"],
        ["Database, Key Vault, gateway, firewall", "Zone-redundant service tiers", "Customer landing zone standards", "Customer"],
    ], caption="Availability zone design by layer", widths=[3, 6.4, 5, 2.2], size=8.5, label="zones")
    b.callout("caution", "Topology spread with `ScheduleAnyway` is a preference. After a zone returns, Kubernetes does not move "
              "running pods back, so the spread stays uneven until the next rollout. Check the spread after FS-21 and after "
              "every zone incident, and restart the tier if one zone holds most of the pods.")
