"""Consistency checks on the built DOCX: cross-references, unresolved labels and wording."""
import re
import sys
from collections import defaultdict

from docx import Document

PATH = "out/Pega_26_AKS_Kafka_Search_Runbook.docx"
BANNED = [r"\bTODO\b", r"\bTBD\b", r"lorem", r"\bdelve", r"\bleverag", r"\bseamless", r"\brobust\b", r"comprehensive",
          r"cutting-edge", r"in today's", r"important to note", r"\bfurthermore\b", r"\bmoreover\b", r"\butiliz",
          r"\u2014", r"game.changer", r"\bholistic", r"\bembark", r"Figure 00", r"Table 00", r"\{ref:", r"\bensures? that\b"]


def texts(doc):
    for p in doc.paragraphs:
        yield p.style.name, p.text
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield "cell", p.text


def main():
    doc = Document(PATH)
    headings = {}
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            m = re.match(r"(?:Appendix )?([A-G]|\d+)((?:\.\d+)*)\.?\s+(.*)", p.text)
            if m:
                headings[m.group(1) + m.group(2)] = m.group(3)
    body = [t for s, t in texts(doc) if not s.startswith("TOC")]
    errors = 0
    refs = defaultdict(int)
    for t in body:
        for m in re.finditer(r"\bSections? ((?:\d+(?:\.\d+)*)(?:(?:, | and | to )\d+(?:\.\d+)*)*)", t):
            for n in re.findall(r"\d+(?:\.\d+)*", m.group(1)):
                refs[n] += 1
        for n in re.findall(r"\bAppendi(?:x|ces) ([A-G])\b", t):
            refs[n] += 1
    for n in sorted(refs, key=lambda x: [int(i) if i.isdigit() else ord(i) for i in x.split(".")]):
        title = headings.get(n)
        if title is None:
            print(f"MISSING  Section {n} ({refs[n]}x)")
            errors += 1
        else:
            print(f"ok       {n:8} {title}")
    defined = set()
    for t in body:
        m = re.match(r"(OD-\d+)$", t.strip())
        if m:
            defined.add(m.group(1))
    used = set(re.findall(r"OD-\d+", " ".join(body)))
    print("OD used but not defined:", sorted(used - defined))
    for t in body:
        for pat in BANNED:
            if re.search(pat, t, re.I):
                print(f"WORDING  {pat}: {t[:140]}")
                errors += 1
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
