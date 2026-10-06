"""Appendices A to H (Appendix H content is in content_h)."""
from docx.shared import Pt

from docx_lib import add_hyperlink
import envs as E
import values_gen as V
import content_h as H


def _per_env(label, fn):
    return [label] + [fn(name, code, grp) for name, code, grp in E.ENVS]


def app_a_inventory(b):
    b.h1("Configuration inventory", appendix="A")
    b.p("This appendix holds every environment-specific value. The names in the first table follow the convention in "
        "Section 2.6 and are fixed once the environment is built. The values in the later tables are recorded during the "
        "build and kept under change control. Secrets are never recorded here; only the Key Vault secret names are.")
    b.h2("Names per environment")
    w = [3.4] + [2.2] * 6
    b.table(["Item"] + E.NAMES, [
        _per_env("Code", lambda n, c, g: f"`{c}`"),
        _per_env("Service group", lambda n, c, g: g),
        _per_env("Pega namespace", lambda n, c, g: f"pega-{c}"),
        _per_env("SRS namespace and release", lambda n, c, g: f"srs-{c}"),
        _per_env("Key Vault", lambda n, c, g: f"kv-pega-{c}"),
        _per_env("Confluent cluster", lambda n, c, g: E.cluster(g)),
        _per_env("Service account", lambda n, c, g: f"sa-pega-{c}"),
        _per_env("Topic and group prefix", lambda n, c, g: E.prefix(c)),
        _per_env("customerDeploymentId and index prefix", lambda n, c, g: E.deployment_id(c)),
        _per_env("OpenSearch service", lambda n, c, g: E.search(g)),
        _per_env("OpenSearch user", lambda n, c, g: f"srs-{c}"),
        _per_env("OpenSearch role", lambda n, c, g: f"pega26-{c}-srs"),
        _per_env("Okta client", lambda n, c, g: f"pega-srs-{c}"),
        _per_env("Web replicas (min to max)", lambda n, c, g: f"{V.SIZES[c][0]} to {V.SIZES[c][1]}"),
        _per_env("Batch replicas (min to max)", lambda n, c, g: f"{V.SIZES[c][2]} to {V.SIZES[c][3]}"),
        _per_env("SRS replicas", lambda n, c, g: str(V.SIZES[c][4])),
    ], caption="Names and sizes per environment", widths=w, size=7.5, first_col_bold=True, label="names")
    b.p("The replica counts are starting values. PERF measurements (M3) set the PROD and PREPROD values.")
    b.h2("Values recorded per service group")
    b.table(["Value", "Where used", "NP1", "NP2", "PRD"], [
        ["Confluent environment and cluster ID (`lkc-...`)", "Commands; Confluent console", "", "", ""],
        ["Cluster type and capacity", "OD-02; Section 6.3", "", "", ""],
        ["Bootstrap host (`<bootstrap-host>`)", "`stream.bootstrapServer`", "", "", ""],
        ["Confluent network ID and private endpoint", "Section 6.4", "", "", ""],
        ["Private DNS zone for Confluent", "Section 4.4", "", "", ""],
        ["Partition limit and measured count", "Section 6.7", "", "", ""],
        ["OpenSearch provider and version", "OD-03; Section 7.2", "", "", ""],
        ["OpenSearch endpoint (`<search-host>`) and port", "`srs.srsStorage.domain`, `port`", "", "", ""],
        ["OpenSearch data nodes, storage and shards", "Section 7.4", "", "", ""],
        ["Disk watermarks (low, high, flood stage)", "Section 7.4; Section 19.5", "", "", ""],
        ["Okta authorization server ID (`<auth-server-id>`)", "Token URL", "", "", ""],
        ["Token endpoint and JWKS URL (`<okta-jwks-url>`)", "`pegasearch.srsAuth.url`; `OAuthPublicKeyURL`", "", "", ""],
        ["Access token lifetime", "Okta access policy rule", "", "", ""],
    ], caption="Values recorded per service group", widths=[5.2, 4.4, 2.3, 2.3, 2.4], size=8)
    b.h2("Values recorded per environment")
    b.table(["Value", "Where used", "Recorded value and date"], [
        ["Confluent API key ID (not the secret)", "`confluent-jaas` in Key Vault", ""],
        ["Client quota (produce and consume)", "Section 5.4", ""],
        ["Okta client ID (`<okta-client-id>`)", "`pegasearch.srsAuth.clientId`", ""],
        ["Okta client key ID (`kid`) and expiry", "Section 8.3", ""],
        ["JDBC URL, driver class and database type", "`global.jdbc`", ""],
        ["Rules and data schema names", "`global.jdbc.rulesSchema`, `dataSchema`", ""],
        ["Pega host name (`<pega-host>`)", "Web tier ingress", ""],
        ["Application Gateway request timeout", "Ingress annotation", ""],
        ["Pega, installer and SRS image digests", "Pipeline record", ""],
        ["Helm chart version", "Pipeline", "4.13.0"],
        ["SRS image tag (`<tag>`)", "`srs.srsRuntime.srsImage`", ""],
        ["Index shard and replica counts after the first build", "Section 7.4", ""],
        ["Full index build time and document counts", "Section 7.8", ""],
        ["Cloned DSS for stream and search (CD-01, CD-03)", "Section 11.3", ""],
    ], caption="Values recorded per environment (one copy per environment)", widths=[6, 5.4, 5.2], size=8.5)


def app_b_values(b):
    b.h1("Complete Helm values files", appendix="B")
    b.p("These files are generated from the same source as the tables in Section 10, for chart version 4.13.0 [R23, R24]. "
        "SIT stands for the NP1 environments and PROD for PROD; the other environments differ only in the code, the names in "
        "{ref:tab_names} and the replica counts. Values in angle brackets come from Appendix A. No secret is written in these "
        "files: each `external_secret_name` and `authSecret` names a Kubernetes secret created by the External Secrets "
        "Operator (Section 8.2). Render every file with `helm template` and review the output before use. Every file below "
        "was rendered with `helm template` against charts 4.13.0 without errors before this version was issued.")
    files = V.all_files()
    for code, name in (("sit", "SIT"), ("prd", "PROD")):
        for run, desc in V.RUNS.items():
            b.h2(f"{name}: pega chart, {desc.split(':')[0].lower()}")
            b.code(files[f"pega-{code}-{run}.yaml"], title=f"pega-{code}-{run}.yaml. {desc}.")
        b.h2(f"{name}: backingservices chart (SRS)")
        b.code(files[f"backingservices-{code}.yaml"], title=f"backingservices-{code}.yaml")
        b.h2(f"{name}: network policies for SRS")
        b.code(files[f"netpol-srs-{code}.yaml"], title=f"netpol-srs-{code}.yaml (applied with `kubectl apply`, Section 7.7)")


def app_c_commands(b):
    b.h1("Command reference", appendix="C")
    b.p("Run these commands from a tooling pod in `pega-<code>`, so they use the same network path, DNS and egress rules as "
        "Pega, unless the command says otherwise. Use a tooling image mirrored into the customer registry. Read secrets from "
        "Key Vault at run time and keep them out of shell history. Check option names against the installed tool versions "
        "before use.")
    b.h2("Kafka and Confluent Cloud")
    b.code('''# K-1 DNS: bootstrap and broker names must resolve to private IPs
nslookup <bootstrap-host>
nslookup <broker-host>

# K-2 TLS: certificate chain and protocol
openssl s_client -connect <bootstrap-host>:9092 -servername <bootstrap-host> </dev/null

# K-3 Authentication and metadata: every broker listed and reachable
kcat -b <bootstrap-host>:9092 -X security.protocol=SASL_SSL -X sasl.mechanisms=PLAIN \\
  -X sasl.username="$API_KEY" -X sasl.password="$API_SECRET" -L

# K-4 Topics under the environment prefix
confluent kafka topic list --cluster <lkc-id> | grep "pega-<code>-"

# K-5 Partition count under the prefix
kcat -b <bootstrap-host>:9092 -X security.protocol=SASL_SSL -X sasl.mechanisms=PLAIN \\
  -X sasl.username="$API_KEY" -X sasl.password="$API_SECRET" -L -J \\
  | jq --arg p "pega-<code>-" \\
    '[.topics[] | select(.topic | startswith($p)) | .partitions | length] | add'

# K-6 Topic configuration and the message size limit
confluent kafka topic describe <topic> --cluster <lkc-id>
confluent kafka topic update <topic> --config max.message.bytes=5000000 \\
  --cluster <lkc-id>

# K-7 ACLs for the environment's service account
confluent kafka acl list --service-account <sa-id> --cluster <lkc-id>''', title="Kafka commands K-1 to K-7")
    b.p("K-8 runs against the customer's own Kafka cluster that 8.8 application data sets read, with that cluster's "
        "client properties. K-9 runs only during a refresh or retirement, after the environment's Pega tiers are at zero.")
    b.code('''# K-8 Committed offsets of an application consumer group (customer cluster)
kafka-consumer-groups.sh --bootstrap-server <customer-bootstrap> \\
  --command-config customer-client.properties --describe --group <group>

# K-9 Delete the environment's topics and consumer groups (refresh or retirement)
confluent kafka topic list --cluster <lkc-id> -o json \\
  | jq -r '.[].name | select(startswith("pega-<code>-"))' > topics.txt
xargs -n1 confluent kafka topic delete --force --cluster <lkc-id> < topics.txt
kafka-consumer-groups.sh --bootstrap-server <bootstrap-host>:9092 \\
  --command-config sa-pega-<code>.properties --list | grep "^pega-<code>-" > groups.txt
xargs -n1 kafka-consumer-groups.sh --bootstrap-server <bootstrap-host>:9092 \\
  --command-config sa-pega-<code>.properties --delete --group < groups.txt''',
           title="Kafka commands K-8 and K-9")
    b.callout("caution", "Review `topics.txt` and `groups.txt` before running the delete lines. Every name must start with "
              "this environment's prefix. The prefixed ACLs stop the service account from deleting another environment's "
              "topics, but the review is still required.")
    b.h2("OpenSearch and SRS")
    b.code('''# S-1 DNS and version: private IP; version 2.15 or 2.19
nslookup <search-host>
curl -s -u "$OS_USER:$OS_PASSWORD" https://<search-host>/

# S-2 Cluster health (green expected)
curl -s -u "$OS_USER:$OS_PASSWORD" "https://<search-host>/_cluster/health?pretty"

# S-3 Settings required by SRS (administrator)
curl -s -u "$OS_ADMIN:$OS_ADMIN_PASSWORD" \\
  "https://<search-host>/_cluster/settings?pretty&include_defaults=false"

# S-4 Indexes under the environment's customerDeploymentId
curl -s -u "$OS_USER:$OS_PASSWORD" "https://<search-host>/_cat/indices/pega26-<code>*?v"

# S-5 SRS from the Pega namespace: refused without a token, accepted with one
SRS=https://srs-<code>.srs-<code>.svc.cluster.local:8443
curl -s --cacert srs-ca.crt "$SRS/health"
curl -s --cacert srs-ca.crt -o /dev/null -w "%{http_code}\\n" "$SRS/"
curl -s --cacert srs-ca.crt -o /dev/null -w "%{http_code}\\n" \\
  -H "Authorization: Bearer $TOKEN" "$SRS/"

# S-6 Delete the environment's indexes (refresh or retirement, as the SRS user)
curl -s -u "$OS_USER:$OS_PASSWORD" \\
  "https://<search-host>/_cat/indices/pega26-<code>*?h=index"
curl -s -u "$OS_USER:$OS_PASSWORD" -X DELETE "https://<search-host>/pega26-<code>*"''', title="Search commands S-1 to S-6")
    b.p("For S-5, `srs-ca.crt` is the customer private CA that signed the SRS server certificate. The health call is the "
        "same one the SRS readiness probe makes, and must return a healthy status. The call without a token must return "
        "401. The call with a token must not return 401 or 403; the exact code depends on the path, which is not part of "
        "the check. S-6 relies on `destructive_requires_name` being false (OS-2) and on the "
        "index-scoped role, which limits the delete to this environment's indexes.")
    b.h2("Okta tokens")
    b.p("O-1 builds the signed client assertion that `private_key_jwt` needs and requests a token, as Pega does [R43, R45]. "
        "The assertion carries the client ID as issuer and subject, the token endpoint as audience, a unique ID and a short "
        "expiry. Run it from the Pega namespace, so the egress path to `<okta-domain>` is tested too. It needs Python with "
        "the PyJWT and cryptography packages in the tooling image.")
    b.code('''# O-1 Token request with private_key_jwt
export TOKEN_URL="https://<okta-domain>/oauth2/<auth-server-id>/v1/token"
export CLIENT_ID="<okta-client-id>"
ASSERTION=$(python3 - <<'EOF'
import os, time, uuid, jwt
now = int(time.time())
claims = {"iss": os.environ["CLIENT_ID"], "sub": os.environ["CLIENT_ID"],
          "aud": os.environ["TOKEN_URL"], "iat": now, "exp": now + 300,
          "jti": str(uuid.uuid4())}
key = open("/tmp/okta-srs-private-key.pem").read()
print(jwt.encode(claims, key, algorithm="RS256", headers={"kid": "<kid>"}))
EOF
)
TOKEN=$(curl -s -X POST "$TOKEN_URL" \\
  -H "Content-Type: application/x-www-form-urlencoded" \\
  -d grant_type=client_credentials -d scope=pega.search:full \\
  -d client_assertion_type=urn:ietf:params:oauth:client-assertion-type:jwt-bearer \\
  -d client_assertion="$ASSERTION" | jq -r .access_token)
rm -f /tmp/okta-srs-private-key.pem

# O-2 Decode the token claims without verifying them.
# Check iss, scp (contains pega.search:full), guid (pega26-<code>) and exp.
# Do not paste production tokens into web tools.
python3 - "$TOKEN" <<'EOF'
import base64, json, sys
payload = sys.argv[1].split(".")[1]
payload += "=" * (-len(payload) % 4)
print(json.dumps(json.loads(base64.urlsafe_b64decode(payload)), indent=2))
EOF''', title="Okta commands O-1 and O-2")
    b.h2("Kubernetes and Helm")
    b.code('''# P-1 Pod status
kubectl get pods -n pega-<code> -o wide
kubectl get pods -n srs-<code> -o wide

# P-2 Pod logs: stream and search errors
kubectl logs -n pega-<code> <pod> --since=30m | grep -Ei "stream|kafka|srs|search"

# P-3 Rolling restart of a tier
kubectl rollout restart deployment/pega-<tier> -n pega-<code>
kubectl rollout status deployment/pega-<tier> -n pega-<code>

# P-4 External secret synchronization (SecretSynced expected)
kubectl get externalsecret -n pega-<code>
kubectl get externalsecret -n srs-<code>

# P-5 Installer job log
kubectl get jobs -n pega-<code>
kubectl logs -n pega-<code> job/<installer-job> -f

# P-6 Render and compare values before any Helm run
helm template pega pega/pega --version 4.13.0 -n pega-<code> \\
  -f pega-<code>-<run>.yaml > rendered.yaml
helm diff upgrade pega pega/pega --version 4.13.0 -n pega-<code> \\
  -f pega-<code>-<run>.yaml''', title="Kubernetes and Helm commands P-1 to P-6")
    b.p("P-6 uses the helm-diff plug-in. If it is not allowed, compare the rendered output with the output of the previous "
        "run instead.")


def _form(b, caption, fields):
    b.table(["Field", "Entry"], [[f, ""] for f in fields], widths=[5.2, 11.4], caption=caption, size=9,
            first_col_bold=True, zebra=False)


def app_d_templates(b):
    b.h1("Evidence and rehearsal templates", appendix="D")
    b.p("Each environment build, refresh, rehearsal and cutover produces an evidence pack, stored with the change record. "
        "The templates fix what each record holds, so rehearsal and production results can be compared line by line.")
    b.h2("Evidence pack contents")
    b.table(["Item", "Content", "Produced by"], [
        ["E-01 Values diff", "P-6 output against the previous run, secrets excluded", "Platform team"],
        ["E-02 Kafka connectivity", "K-1 to K-3 output from a pod in `pega-<code>`", "Kafka engineer"],
        ["E-03 Kafka objects", "K-4 to K-7 output", "Kafka engineer"],
        ["E-04 Search connectivity", "S-1 to S-3 and S-5 output", "Search engineer"],
        ["E-05 Token claims", "O-2 output (claims only; never the token)", "Identity team"],
        ["E-06 Clean-up record", "One line per CD item (template below)", "Pega LSA"],
        ["E-07 Landing pages", "Stream and search landing page screenshots showing the Helm values", "Pega LSA"],
        ["E-08 Index record", "S-4 output, build time and count sheet", "Pega LSA"],
        ["E-09 Test records", "Isolation, functional, performance and failure scenario results", "Test lead"],
        ["E-10 Timings", "Start and end time of each step", "Cutover manager"],
        ["E-11 Data protection", "DP-1 to DP-6 output; DT-01 to DT-12 results; signed indexed property list and snapshot inventory", "Security architect"],
    ], caption="Evidence pack items", widths=[3.6, 9.6, 3.4], size=8.5)
    b.h2("Clean-up record")
    b.table(["CD item", "Found in the clone", "Action taken", "Done by and time", "Check"], [
        [f"CD-{i:02d}", "", "", "", ""] for i in range(1, 17)
    ], caption="Clean-up record (one per clone)", widths=[1.8, 4, 4.6, 3, 3.2], size=8.5, zebra=False)
    b.h2("Check record")
    _form(b, "Check record", ["Check or scenario ID", "Environment and occasion", "Run by and date", "Command or screen",
                              "Expected result", "Actual result", "Result (Pass, Fail, Pass with note)", "Evidence item",
                              "Deviation and action"])
    b.h2("Timed cutover record")
    b.table(["Step", "Planned start", "Actual start", "Actual end", "Duration", "Notes"], [
        [f"C-{i}", "", "", "", "", ""] for i in range(1, 11)
    ], caption="Timed cutover record (rehearsals and production)", widths=[1.6, 2.8, 2.8, 2.8, 2.4, 4.2], size=8.5, zebra=False)
    b.h2("Go/No-Go record")
    b.table(["Gate", "Criteria met", "Evidence items", "Decision owner", "Decision and time"], [
        ["Go/No-Go 1", "", "", "Business owner with platform lead", ""],
        ["First-start gate", "", "", "Cutover manager", ""],
        ["Go/No-Go 2", "", "", "Cutover manager", ""],
        ["Go/No-Go 3", "", "", "Cutover manager with business owner", ""],
    ], caption="Go/No-Go record", widths=[3, 3.4, 3.2, 3.6, 3.4], size=9, zebra=False)
    b.h2("Retirement record")
    b.table(["Item", "Deleted by and date", "Evidence"], [
        ["Helm releases and namespaces", "", ""],
        ["Topics and consumer groups (K-9)", "", ""],
        ["ACLs, API keys and service account", "", ""],
        ["Indexes (S-6), OpenSearch user and role", "", ""],
        ["Okta client and `guid` expression entry", "", ""],
        ["Key Vault secrets and vault", "", ""],
        ["Firewall rules and DNS records", "", ""],
    ], caption="Retirement record", widths=[6.6, 4.4, 5.6], size=9, zebra=False)


GLOSSARY = [
    ("ACL", "Access control list. In Kafka, a rule that allows a principal an operation on a resource such as a topic or consumer group."),
    ("AKS", "Azure Kubernetes Service."),
    ("Authorization server (Okta)", "The Okta component that issues access tokens. A custom authorization server allows custom scopes and claims."),
    ("Clone-and-upgrade", "The path in this document: copy the 8.8 database, upgrade the copy to 26.1.1 with the installer job, and start 26.1.1 on it."),
    ("CodeSet", "A Pega rule that groups Java libraries imported into Pega Platform."),
    ("Confluent Cloud", "Managed Apache Kafka service from Confluent."),
    ("customerDeploymentId", "The ID that SRS uses to separate each Pega environment's data. It becomes the index name prefix and must match the `guid` claim."),
    ("Data set (Kafka)", "A Pega rule that reads from or writes to a Kafka topic for an application integration. It is separate from the stream service."),
    ("DSS", "Dynamic System Setting, a configuration record in the Pega database."),
    ("External Secrets Operator", "A Kubernetes operator that copies secrets from Azure Key Vault into Kubernetes secrets."),
    ("Go/No-Go", "A decision point with written criteria, at which the cutover continues or stops."),
    ("guid", "The token claim that SRS compares with the customerDeploymentId of the environment."),
    ("Hazelcast", "The in-memory clustering library used by Pega Platform up to '24, removed in '25."),
    ("Helm", "The package manager for Kubernetes used to deploy the pega and backingservices charts."),
    ("Index build", "Full reindex of Pega data into SRS from the upgraded database."),
    ("Installer job", "The Kubernetes job the pega chart runs to install or upgrade the Pega database."),
    ("JAAS", "Java Authentication and Authorization Service. The Kafka client reads its SASL credentials from a JAAS configuration string."),
    ("JWKS", "JSON Web Key Set, the public keys that SRS uses to check token signatures."),
    ("Masking", "Replacing personal or sensitive data in a non-production copy of the database."),
    ("OpenSearch", "Open-source search engine. SRS stores Pega indexes in it."),
    ("Private Link", "Azure service that exposes a service through a private IP in the customer VNet."),
    ("private_key_jwt", "Client authentication in which the client signs an assertion with its private key instead of sending a secret."),
    ("Queue processor", "A Pega rule that processes queued items asynchronously through the stream service."),
    ("SASL_SSL", "Kafka security protocol that combines TLS encryption with SASL authentication."),
    ("Service group", "The set of environments that share one Confluent cluster and one OpenSearch service (NP1, NP2) or the dedicated production set (PRD)."),
    ("SRS", "Search and Reporting Service, the Pega backing service between Pega Platform and OpenSearch."),
    ("Stream service", "The Pega service that uses Kafka for queue processors, job schedulers, data flows and, from '25, cluster messaging."),
    ("streamNamePattern", "Helm setting that defines how Pega names its Kafka topics. It carries the environment prefix."),
]


def app_e_glossary(b):
    b.h1("Glossary", appendix="E")
    b.table(["Term", "Meaning"], [[t, m] for t, m in GLOSSARY], widths=[4, 12.6], caption="Glossary of terms", size=9,
            first_col_bold=True)


def app_f_references(b):
    b.h1("References", appendix="F")
    b.p("Sources were read between 28 September and 1 October 2026. Vendor pages change; check each page again before each "
        "rehearsal (RK-12).")
    for key in sorted(b.refs, key=lambda k: int(k[1:])):
        publisher, title, url = b.refs[key]
        para = b.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.left_indent = Pt(34)
        para.paragraph_format.first_line_indent = Pt(-34)
        r = para.add_run(f"[{key}]\t")
        r.bold = True
        r.font.size = Pt(9.5)
        r = para.add_run(f"{publisher}. {title}. ")
        r.font.size = Pt(9.5)
        add_hyperlink(para, url, url, size=8.5)


IMAGE_SOURCES = [
    ("Pega Platform Kubernetes architecture for '26", "Pega Documentation", "R21"),
    ("Earlier Pega cluster layout with embedded Kafka and in-cluster Elasticsearch", "Pega Academy", "R30"),
    ("Externalized services in a Pega deployment", "Pega Academy", "R29"),
    ("The two Kafka use cases in Pega Platform", "Pega Documentation", "R11"),
    ("Azure Private Link between a customer VNet and Confluent Cloud", "Confluent Documentation", "R32"),
    ("DNS records that Confluent lists for Azure Private Link", "Confluent Documentation", "R32"),
    ("Client-managed cloud deployment with SRS as a backing service", "Pega Academy", "R28"),
    ("OpenSearch cluster roles: cluster-manager and data nodes", "OpenSearch Documentation", "R37"),
]


def app_g_images(b):
    b.h1("Image sources and attribution", appendix="G")
    b.p("Figures marked \"Prepared for this implementation\" were drawn for this document from the facts cited in the "
        "surrounding text. They are not Pega-published diagrams. The figures below are reproduced from public vendor pages, "
        "unchanged except for scaling.")
    b.table(["Figure", "Publisher", "Source page"], [[f, p, f"[{r}]"] for f, p, r in IMAGE_SOURCES],
            widths=[9.4, 4.6, 2.6], caption="Third-party images", size=9)
    b.callout("important", "Pega, Pega Academy, Confluent and OpenSearch content is subject to each publisher's terms of use. "
              "Before this document is shared outside the customer and implementation teams, confirm that reuse of these "
              "images is allowed, or replace them with links to the source pages.")


def appendices(b):
    app_a_inventory(b)
    app_b_values(b)
    app_c_commands(b)
    app_d_templates(b)
    app_e_glossary(b)
    app_f_references(b)
    app_g_images(b)
    H.app_h_calculator(b)
