"""Sizing calculator: the formulas used in Section 16 and the Excel workbook of Appendix H."""
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

UNIT_LIMITS = {
    # type: ingress MBps, egress MBps, partitions, connections, requests per second [R31]
    "Enterprise": (60, 180, 3000, 18000, 7500),
    "Dedicated": (60, 180, 4500, 18000, 15000),
}
LIMIT_NAMES = ["ingress", "egress", "partitions", "connections", "requests", "SLA minimum"]
SLA_MIN_UNITS = 2
STORAGE_FACTOR = 1.45
SHARDS_PER_GIB_HEAP = 25
ZONES = 3

KAFKA_EXAMPLE = {
    "NP1": dict(type="Enterprise", ingress=6, egress=18, partitions=2700, headroom=0.3, connections=2400, requests=2500),
    "NP2": dict(type="Enterprise", ingress=10, egress=30, partitions=1800, headroom=0.3, connections=3600, requests=6000),
    "PROD": dict(type="Enterprise", ingress=8, egress=24, partitions=900, headroom=0.3, connections=1800, requests=3000),
}
OS_EXAMPLE = {
    "os-np1": dict(primary_gib=30, replicas=1, growth=0.2, years=3, rebuild_gib=0, indexes=180, primaries=1, disk_gib=100, heap_gib=8),
    "os-np2": dict(primary_gib=80, replicas=1, growth=0.2, years=3, rebuild_gib=0, indexes=120, primaries=1, disk_gib=100, heap_gib=8),
    "os-prd": dict(primary_gib=40, replicas=1, growth=0.2, years=3, rebuild_gib=0, indexes=60, primaries=1, disk_gib=100, heap_gib=8),
}
AKS_EXAMPLE = {
    "PROD Pega pool": dict(pods=12, cpu_req=3, mem_req=12, alloc_cpu=15, alloc_mem=58),
}
OKTA_EXAMPLE = [
    ("DEV", 4, 12), ("SIT", 4, 12), ("UAT", 5, 12), ("PERF", 12, 12), ("PREPROD", 12, 12), ("PROD", 12, 12),
]


def kafka_units(type, ingress, egress, partitions, headroom, connections, requests):
    li, le, lp, lc, lr = UNIT_LIMITS[type]
    planned = math.ceil(partitions * (1 + headroom))
    by = dict(zip(LIMIT_NAMES, [math.ceil(ingress / li), math.ceil(egress / le), math.ceil(planned / lp),
                                math.ceil(connections / lc), math.ceil(requests / lr), SLA_MIN_UNITS]))
    units = max(by.values())
    driver = [k for k, v in by.items() if v == units][-1]
    return dict(type=type, ingress=ingress, egress=egress, partitions=partitions, planned=planned,
                connections=connections, requests=requests, by_limit=by, units=units, driver=driver)


def opensearch(primary_gib, replicas, growth, years, rebuild_gib, indexes, primaries, disk_gib, heap_gib):
    storage = round(primary_gib * (1 + replicas) * STORAGE_FACTOR * (1 + growth) ** years + rebuild_gib, 1)
    shards = indexes * primaries * (1 + replicas)
    n_storage = math.ceil(storage / disk_gib)
    n_shards = math.ceil(shards / (SHARDS_PER_GIB_HEAP * heap_gib))
    nodes = max(n_storage, n_shards, ZONES)
    nodes = math.ceil(nodes / ZONES) * ZONES
    return dict(primary_gib=primary_gib, replicas=replicas, growth=growth, years=years, indexes=indexes,
                primaries=primaries, disk_gib=disk_gib, heap_gib=heap_gib, storage_gib=round(storage, 1),
                shards=shards, nodes_storage=n_storage, nodes_shards=n_shards, nodes=nodes)


def aks(pods, cpu_req, mem_req, alloc_cpu, alloc_mem):
    per_node = min(math.floor(alloc_cpu / cpu_req), math.floor(alloc_mem / mem_req))
    base = math.ceil(pods / per_node)
    nodes = math.ceil(base * 1.5) + 1
    nodes = math.ceil(nodes / ZONES) * ZONES
    return dict(pods=pods, cpu_req=cpu_req, mem_req=mem_req, alloc_cpu=alloc_cpu, alloc_mem=alloc_mem,
                per_node=per_node, base=base, nodes=nodes)


def example():
    return dict(kafka=kafka_units(**KAFKA_EXAMPLE["PROD"]), opensearch=opensearch(**OS_EXAMPLE["os-prd"]),
                aks=aks(**AKS_EXAMPLE["PROD Pega pool"]))


HEAD = PatternFill("solid", fgColor="1F3864")
INPUT = PatternFill("solid", fgColor="FFF2CC")
OUTPUT = PatternFill("solid", fgColor="E2EFDA")
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)


def _header(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font = HEAD, WHITE
        c.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w


def _rows(ws, start, rows, kinds):
    """rows: list of (label, [values per column], note). kinds: 'in' or 'out' per row."""
    for r, ((label, vals, note), kind) in enumerate(zip(rows, kinds), start):
        ws.cell(row=r, column=1, value=label).font = BOLD
        for i, v in enumerate(vals, 2):
            c = ws.cell(row=r, column=i, value=v)
            c.fill = INPUT if kind == "in" else OUTPUT
        ws.cell(row=r, column=2 + len(vals), value=note).alignment = Alignment(wrap_text=True)


def _readme(wb):
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Pega 26.1.1 on Azure AKS: sizing calculator for Kafka, OpenSearch, AKS and Okta", True),
        ("Companion to the runbook, Section 16 and Appendix H.", False),
        ("", False),
        ("Yellow cells are inputs. Green cells are formulas. Change only the yellow cells.", True),
        ("The example inputs are illustrative. They are not measurements of the customer's system. Replace each one with a measured value.", True),
        ("", False),
        ("Sources for the constants:", True),
        ("Confluent Cloud per-unit limits and the 2-unit SLA minimum: Confluent, 'Kafka cluster types in Confluent Cloud' (runbook reference R31).", False),
        ("Pega partitions per topic default of 6: Pega, 'Changing the default number of partitions per topic' (R49).", False),
        ("OpenSearch storage factor 1.45 and 25 shards per GiB of heap: Amazon OpenSearch Service sizing and shard guidance (R54, R55).", False),
        ("AKS zone headroom of 1.5 and one surge node: this design (runbook Section 4.6).", False),
        ("Okta rate limits: enter your org's token endpoint limit from the Okta Admin Console (R52).", False),
        ("", False),
        ("Recheck the constants against the vendor pages before each use; vendors change them.", True),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(bold=bold, size=12 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True)


def _kafka(wb):
    ws = wb.create_sheet("Kafka")
    groups = list(KAFKA_EXAMPLE)
    _header(ws, 1, ["Confluent Cloud capacity units per cluster"] + groups + ["Notes"], [44, 16, 16, 16, 60])
    cols = [chr(66 + i) for i in range(len(groups))]
    inputs = [
        ("Cluster type (Enterprise or Dedicated)", [KAFKA_EXAMPLE[g]["type"] for g in groups], "Per OD-02"),
        ("Peak ingress, MBps", [KAFKA_EXAMPLE[g]["ingress"] for g in groups], "Metrics API, PERF peak, summed over the group"),
        ("Peak egress, MBps", [KAFKA_EXAMPLE[g]["egress"] for g in groups], "Metrics API"),
        ("Measured partitions (sum over the group)", [KAFKA_EXAMPLE[g]["partitions"] for g in groups], "K-5 after first start; default 6 per topic"),
        ("Partition headroom", [KAFKA_EXAMPLE[g]["headroom"] for g in groups], "30 % suggested"),
        ("Peak client connections", [KAFKA_EXAMPLE[g]["connections"] for g in groups], "Metrics API; include a rolling restart"),
        ("Peak requests per second", [KAFKA_EXAMPLE[g]["requests"] for g in groups], "Metrics API"),
    ]
    _rows(ws, 2, inputs, ["in"] * len(inputs))
    dv = DataValidation(type="list", formula1='"Enterprise,Dedicated"', allow_blank=False)
    ws.add_data_validation(dv)
    for c in cols:
        dv.add(f"{c}2")
    ws.cell(row=10, column=1, value="Per-unit limits [R31]").font = BOLD
    lim_rows = [("Ingress MBps per unit", 0), ("Egress MBps per unit", 1), ("Partitions per unit", 2),
                ("Connections per unit", 3), ("Requests per second per unit", 4)]
    for i, (label, k) in enumerate(lim_rows, 11):
        ws.cell(row=i, column=1, value=label)
        for c in cols:
            ws[f"{c}{i}"] = (f'=IF({c}$2="Dedicated",{UNIT_LIMITS["Dedicated"][k]},{UNIT_LIMITS["Enterprise"][k]})')
            ws[f"{c}{i}"].fill = OUTPUT
    ws.cell(row=17, column=1, value="Units needed by each limit").font = BOLD
    calc = [
        ("By ingress", "=ROUNDUP({c}3/{c}11,0)"),
        ("By egress", "=ROUNDUP({c}4/{c}12,0)"),
        ("Planned partitions", "=ROUNDUP({c}5*(1+{c}6),0)"),
        ("By partitions", "=ROUNDUP({c}20/{c}13,0)"),
        ("By connections", "=ROUNDUP({c}7/{c}14,0)"),
        ("By requests", "=ROUNDUP({c}8/{c}15,0)"),
        ("SLA minimum", f"={SLA_MIN_UNITS}"),
    ]
    for i, (label, f) in enumerate(calc, 18):
        ws.cell(row=i, column=1, value=label)
        for c in cols:
            ws[f"{c}{i}"] = f.format(c=c)
            ws[f"{c}{i}"].fill = OUTPUT
    ws.cell(row=26, column=1, value="Capacity units (eCKU or CKU)").font = BOLD
    ws.cell(row=27, column=1, value="Set by").font = BOLD
    for c in cols:
        ws[f"{c}26"] = f"=MAX({c}18,{c}19,{c}21,{c}22,{c}23,{c}24)"
        ws[f"{c}26"].fill, ws[f"{c}26"].font = OUTPUT, BOLD
        ws[f"{c}27"] = (f'=IF({c}26={c}24,"SLA minimum",IF({c}26={c}23,"requests",IF({c}26={c}22,"connections",'
                        f'IF({c}26={c}21,"partitions",IF({c}26={c}19,"egress","ingress")))))')
        ws[f"{c}27"].fill = OUTPUT
    ws.cell(row=28, column=1, value="Check").font = BOLD
    for c in cols:
        ws[f"{c}28"] = f'=IF(AND({c}2="Enterprise",{c}26>32),"Above 32 eCKU: Private Link limit on Azure",IF(AND({c}2="Enterprise",{c}26>10),"Above 10 eCKU: slower scaling; set minimum before peaks","OK"))'
        ws[f"{c}28"].fill = OUTPUT
        ws[f"{c}28"].alignment = Alignment(wrap_text=True)


def _opensearch(wb):
    ws = wb.create_sheet("OpenSearch")
    svcs = list(OS_EXAMPLE)
    _header(ws, 1, ["OpenSearch storage and data nodes per service"] + svcs + ["Notes"], [44, 16, 16, 16, 60])
    cols = [chr(66 + i) for i in range(len(svcs))]
    keys = [("Primary data, GiB", "primary_gib", "Sum of pri.store.size after a full build (S-4), all environments on the service"),
            ("Replicas", "replicas", "At least 1"),
            ("Yearly growth", "growth", "From the workload model"),
            ("Years planned", "years", ""),
            ("Extra peak during a full rebuild, GiB", "rebuild_gib", "Measured in DEV; 0 if none"),
            ("Number of indexes", "indexes", "_cat/indices, all environments on the service"),
            ("Primary shards per index", "primaries", "Set by SRS; read from _cat/indices"),
            ("Usable disk per data node, GiB", "disk_gib", "Provider plan"),
            ("JVM heap per data node, GiB", "heap_gib", "Usually half the node memory")]
    _rows(ws, 2, [(l, [OS_EXAMPLE[s][k] for s in svcs], n) for l, k, n in keys], ["in"] * len(keys))
    calc = [
        ("Storage needed, GiB", "=ROUND({c}2*(1+{c}3)*" + str(STORAGE_FACTOR) + "*(1+{c}4)^{c}5+{c}6,1)", "x 1.45 [R54]"),
        ("Total shards", "={c}7*{c}8*(1+{c}3)", ""),
        ("Data nodes by storage", "=ROUNDUP({c}12/{c}9,0)", ""),
        ("Data nodes by shards", "=ROUNDUP({c}13/(" + str(SHARDS_PER_GIB_HEAP) + "*{c}10),0)", "25 shards per GiB heap [R55]"),
        ("Data nodes (multiple of 3 zones)", "=CEILING(MAX({c}14,{c}15," + str(ZONES) + ")," + str(ZONES) + ")", ""),
        ("Average primary shard size, GiB", "=ROUND({c}2/({c}7*{c}8),2)", "Search workloads: 10 to 30 GiB [R55]; many small Pega indexes fall below this, which is expected"),
    ]
    for i, (label, f, note) in enumerate(calc, 12):
        ws.cell(row=i, column=1, value=label).font = BOLD if i == 16 else Font()
        for c in cols:
            ws[f"{c}{i}"] = f.format(c=c)
            ws[f"{c}{i}"].fill = OUTPUT
        ws.cell(row=i, column=2 + len(cols), value=note).alignment = Alignment(wrap_text=True)


def _aks(wb):
    ws = wb.create_sheet("AKS")
    pools = list(AKS_EXAMPLE)
    _header(ws, 1, ["AKS nodes for the Pega node pool"] + pools + ["Notes"], [44, 20, 60])
    keys = [("Pega pods at HPA maximum (web + batch)", "pods", "Sum of tier HPA maximums"),
            ("CPU request per pod", "cpu_req", "Chart default 3"),
            ("Memory request per pod, GiB", "mem_req", "Chart default 12Gi"),
            ("Allocatable CPU per node", "alloc_cpu", "kubectl describe node; after system reservations"),
            ("Allocatable memory per node, GiB", "alloc_mem", "kubectl describe node")]
    _rows(ws, 2, [(l, [AKS_EXAMPLE[p][k] for p in pools], n) for l, k, n in keys], ["in"] * len(keys))
    calc = [
        ("Pega pods per node", "=MIN(ROUNDDOWN({c}5/{c}3,0),ROUNDDOWN({c}6/{c}4,0))", ""),
        ("Nodes for the load", "=ROUNDUP({c}2/{c}8,0)", ""),
        ("Nodes with zone headroom and surge", "=CEILING(ROUNDUP({c}9*1.5,0)+1," + str(ZONES) + ")", "x 1.5 so two zones carry the peak; +1 surge node; multiple of 3"),
    ]
    for i, (label, f, note) in enumerate(calc, 8):
        ws.cell(row=i, column=1, value=label).font = BOLD if i == 10 else Font()
        ws[f"B{i}"] = f.format(c="B")
        ws[f"B{i}"].fill = OUTPUT
        ws.cell(row=i, column=3, value=note).alignment = Alignment(wrap_text=True)
    ws.cell(row=12, column=1, value="SRS pods and system pods are not included. Place them on another pool or add their requests.").font = Font(italic=True)


def _okta(wb):
    ws = wb.create_sheet("Okta")
    _header(ws, 1, ["Environment", "Pega pods", "Token requests per pod per hour (measured)", "Token requests per minute"], [20, 14, 34, 26])
    for i, (env, pods, rate) in enumerate(OKTA_EXAMPLE, 2):
        ws.cell(row=i, column=1, value=env)
        for col, v in ((2, pods), (3, rate)):
            ws.cell(row=i, column=col, value=v).fill = INPUT
        ws[f"D{i}"] = f"=ROUND(B{i}*C{i}/60,2)"
        ws[f"D{i}"].fill = OUTPUT
    last = 1 + len(OKTA_EXAMPLE)
    ws.cell(row=last + 1, column=1, value="Total, all environments").font = BOLD
    ws[f"D{last + 1}"] = f"=SUM(D2:D{last})"
    ws[f"D{last + 1}"].fill = OUTPUT
    ws.cell(row=last + 2, column=1, value="Org token endpoint limit per minute").font = BOLD
    ws[f"D{last + 2}"].fill = INPUT
    ws.cell(row=last + 2, column=5, value="Enter the value shown in the Okta Admin Console for your org [R52]")
    ws.cell(row=last + 3, column=1, value="Share of the limit used by Pega").font = BOLD
    ws[f"D{last + 3}"] = f'=IF(D{last + 2}="","enter limit",D{last + 1}/D{last + 2})'
    ws[f"D{last + 3}"].fill = OUTPUT


def write(path):
    wb = Workbook()
    _readme(wb)
    _kafka(wb)
    _opensearch(wb)
    _aks(wb)
    _okta(wb)
    for ws in wb.worksheets:
        ws.freeze_panes = "B2" if ws.title != "Read me" else None
    wb.save(path)


if __name__ == "__main__":
    out = Path(__file__).resolve().parent / "out" / "Pega_26_Sizing_Calculator.xlsx"
    write(out)
    for g, v in KAFKA_EXAMPLE.items():
        print(g, kafka_units(**v)["units"], kafka_units(**v)["driver"])
    for s, v in OS_EXAMPLE.items():
        r = opensearch(**v)
        print(s, r["storage_gib"], r["nodes"])
    print(aks(**AKS_EXAMPLE["PROD Pega pool"]))
    print("wrote", out)
